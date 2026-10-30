"""Polish c3-t: the template import names what it replaces, refuses a concentration in the electrolyte
name, lists allowed values once; the extractor pages a cut sheet. Synthetic only.

* A ``correct`` over an extracted value, or over the run's own file-name token (``m800mV`` is -0.8 V
  though the parser reads no value from it), is named on its own line in the preview, compact too:
  never only a "correct N" count.
* "0.01 M HClO4" in Electrolyte is an invalid cell: the name only, the number in its own column.
* The compact preview lists each column's allowed values once, then the invalid cells briefly.
* ``python -m isaac_api.local_import.extract`` cut spreadsheets at 50 rows with no way to read on: a
  cut sheet now says exactly how to page it, ``--sheet/--rows-offset/--rows-limit`` read the rest, and
  the fixed ``limits`` block is left out of the default output.
"""

from __future__ import annotations

import json

import openpyxl
import pytest

from isaac_api.local_import import cli
from isaac_api.local_import import extract as ex
from isaac_api.local_import import mutations as M
from isaac_api.local_import import template as T

from test_local_import_template_adopt import (  # noqa: F401  (fixtures)
    POT, WHO, _col, _import, _row_of, corpus, grouped, grouped_base)


def _export(bdir, tmp_path, name="t.xlsx"):
    out = tmp_path / name
    assert cli.main(["bundle", "template", "export", str(bdir), "--out", str(out)]) == 0
    return out


def _fill(xlsx, out, cells):
    wb = openpyxl.load_workbook(xlsx)
    for (sheet, ref), value in cells.items():
        wb[sheet][ref] = value
    wb.save(out)
    return out


def _text(bdir, xlsx, capsys, *extra):
    capsys.readouterr()
    code = cli.main(["bundle", "template", "import", str(bdir), str(xlsx), "--by", WHO, *extra])
    return code, capsys.readouterr().out


# --- (2) a cell that replaces or contradicts a file reading is named, one line each ---------------


def test_the_preview_names_each_cell_that_replaces_or_contradicts_a_reading_even_compact(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    r5, r7, r1 = _row_of(xlsx, "run-005"), _row_of(xlsx, "run-007"), _row_of(xlsx, "run-001")
    filled = _fill(xlsx, tmp_path / "f.xlsx", {
        ("File List", f"{_col('Potential (V)')}{r5}"): 0.8,   # the name says m800mV: no candidate, a token
        ("File List", f"{_col('pH')}{r7}"): 2,                # the parser read pH12
        ("File List", f"{_col('Potential (V)')}{r1}"): 0.1,   # agrees: a confirm, not named
    })
    code, rep, err = _import(grouped, filled, capsys)
    assert code == 0, err
    named = {(x["cell"], x["run_id"]) for x in rep["replaces_or_contradicts_a_file_reading"]}
    assert named == {(f"'File List'!{_col('Potential (V)')}{r5}", "run-005"), (f"'File List'!{_col('pH')}{r7}", "run-007")}
    pot = next(x for x in rep["replaces_or_contradicts_a_file_reading"] if x["run_id"] == "run-005")
    assert any("m800mV" in r and "-0.8" in r for r in pot["readings"])
    ph = next(x for x in rep["replaces_or_contradicts_a_file_reading"] if x["run_id"] == "run-007")
    assert any("replaces extracted 12" in r for r in ph["readings"]) and any("pH12" in r for r in ph["readings"])

    code, text = _text(grouped, filled, capsys)
    assert code == 0
    lines = [ln for ln in text.splitlines() if ln.startswith("  'File List'!") and ("m800mV" in ln or "pH12" in ln)]
    assert len(lines) == 2
    assert "2 replace or contradict a file reading" in text
    assert not M.load(grouped)["scientist_decisions"]  # a preview


def test_name_token_readings_are_a_comparison_only():
    names = ["O4_17_PtCNPs_5wp_pH2_m800mV", "O4_17_PtCNPs_5wp_pH2_m800mV_dir"]
    assert T.name_token_readings(names, T.POTENTIAL_PATH) == [{"value": -0.8, "token": "m800mV", "name": names[0]}]
    assert [r["value"] for r in T.name_token_readings(["X_1p5V", "X_000mV", "X_10MNaOH_100mMHClO4"], T.POTENTIAL_PATH)] == [1.5, 0.0]
    assert [r["value"] for r in T.name_token_readings(["X_pH12_400mV"], T.PH_PATH)] == [12]
    assert T.name_token_readings(["X_pH12_400mV"], "context.environment") == []


# --- (3) the electrolyte NAME only ---------------------------------------------------------------


@pytest.mark.parametrize("typed, says", [("0.01 M HClO4", "put 0.01 in 'Electrolyte Concentration (M)'"),
                                         ("1M KOH", "put 1 in 'Electrolyte Concentration (M)'"),
                                         ("KOH 100 mM", "100 mM")])
def test_an_electrolyte_name_with_a_concentration_is_an_invalid_cell(grouped, tmp_path, capsys, typed, says):
    xlsx = _export(grouped, tmp_path)
    r1 = _row_of(xlsx, "run-001")
    filled = _fill(xlsx, tmp_path / "e.xlsx", {("File List", f"{_col('Electrolyte')}{r1}"): typed,
                                               ("File List", f"{_col('Electrolyte')}{_row_of(xlsx, 'run-002')}"): "HClO4"})
    code, rep, err = _import(grouped, filled, capsys, "--apply")
    assert code == 0, err
    [bad] = rep["invalid_cells"]
    assert bad["cell"] == f"'File List'!{_col('Electrolyte')}{r1}" and bad["reason"].startswith("name only") and says in bad["reason"]
    assert [d["value"] for d in rep["decisions"]] == ["HClO4"]


# --- (4) compact: allowed values once --------------------------------------------------------------


def test_the_compact_preview_lists_a_columns_allowed_values_once(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    b = M.load(grouped)
    cells = {("File List", f"{_col('Environment')}{_row_of(xlsx, r['run_id'])}"): "aqueous_electrochemistry"
             for r in T.live_rows(b)}
    cells[("File List", f"{_col('Cell Type')}{_row_of(xlsx, 'run-001')}")] = "three_electrode_static_cell"
    filled = _fill(xlsx, tmp_path / "v.xlsx", cells)
    code, text = _text(grouped, filled, capsys)
    assert code == 0
    allowed = ", ".join(T.rule_for_path("context.environment")["allowed"])
    assert text.count(allowed) == 1
    assert "'aqueous_electrochemistry' in 9 cell(s)" in text and "and 6 more" in text
    assert text.count(", ".join(T.rule_for_path("context.electrochemistry.cell_type")["allowed"])) == 1
    code, full = _text(grouped, filled, capsys, "--full")
    assert full.count(allowed) == 9  # --full keeps each cell's whole reason
    assert len(text) < len(full)


# --- (5) the extractor pages a cut sheet -----------------------------------------------------------


def _sheet(tmp_path, rows=78):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "File List"
    ws.append(["Folder Name", "Potential (V)"])
    for i in range(2, rows + 1):
        ws.append([f"S_{i:02d}_dir", i / 10])
    wb.create_sheet("Notes")["A1"] = "only"
    root = tmp_path / "root"
    root.mkdir()
    path = root / "sheet.xlsx"
    wb.save(path)
    return root, path


def _run(capsys, argv):
    capsys.readouterr()
    code = ex.main(argv)
    out = capsys.readouterr()
    return code, (json.loads(out.out) if code == 0 else None), out.err


def test_a_cut_sheet_says_exactly_how_to_read_the_rest_and_paging_reads_it(tmp_path, capsys):
    root, path = _sheet(tmp_path)
    out = tmp_path / "out"
    code, m, err = _run(capsys, [str(path), "--root", str(root), "--out", str(out)])
    assert code == 0, err
    sheet = next(s for s in m["sheets"] if s["name"] == "File List")
    assert sheet["rows_returned"] == "1..50" and sheet["preview_rows_not_returned"] == 28
    rest = sheet["read_rest"]
    assert (rest["sheet"], rest["rows_offset"], rest["rows_limit"], rest["rows_remaining"]) == ("File List", 50, 50, 28)
    assert "--sheet 'File List' --rows-offset 50 --rows-limit 50" in rest["command"]
    assert "NOT read" in " ".join(m["notes"])
    # the fixed limits block is left out unless asked for
    assert "limits" not in m and "--show-limits" in m["limits_shown"]
    code, shown, _ = _run(capsys, [str(path), "--out", str(out), "--show-limits"])
    assert shown["limits"]["max_preview_rows"] == 50

    # the next page: exactly the rows that were cut, and nothing past them is cut now
    code, page, err = _run(capsys, [str(path), "--root", str(root), "--out", str(out), "--sheet", "File List",
                                    "--rows-offset", "50", "--rows-limit", "50"])
    assert code == 0, err
    [sheet] = page["sheets"]
    assert [r["row"] for r in sheet["preview_rows"]] == list(range(51, 79))
    assert sheet["preview_rows"][0]["cells"] == ["S_51_dir", 5.1]
    assert sheet["rows_returned"] == "51..78" and not sheet["preview_rows_not_returned"] and "read_rest" not in sheet
    assert sheet["rows_before_this_page"] == 50 and "header_rows" not in sheet


@pytest.mark.parametrize("argv, why", [(["--sheet", "Nope"], None), (["--rows-offset", "5"], "need --sheet"),
                                       (["--sheet", "File List", "--rows-limit", "100000"], "--rows-limit 1..")])
def test_paging_refuses_what_it_cannot_do(tmp_path, capsys, argv, why):
    root, path = _sheet(tmp_path)
    code, m, err = _run(capsys, [str(path), "--out", str(tmp_path / "out"), *argv])
    if why is None:
        assert code == 0 and m["status"] == ex.STATUS_REFUSED and "no sheet named 'Nope'" in m["reason"]
    else:
        assert code == 2 and why in err
