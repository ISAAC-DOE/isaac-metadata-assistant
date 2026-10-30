"""``context.electrochemistry.control_mode`` is locally answerable (bundle 1.5.0), as the set-point
became at 1.3.0 (DEC-88).

* Its allowed values are the official schema's ``enum``, read at runtime — never a list here.
* RUN scope, like ``potential_scale``: never asked once for the whole beamtime.
* It is never defaulted and never derived: not from a set-point being present, and no set-point
  from it.
* The schema's conditional rules key the set-point requirement on it: a draft whose CONFIRMED
  mode is potentiostatic is incomplete without a composed set-point.
* A 1.4.0 bundle in which a value at this path moves eligibility is STALE, not INVALID.

SYNTHETIC files only (the grouping CLI fixture; the template's two-column fixture).
"""

from __future__ import annotations

import copy
import json
from pathlib import Path

import openpyxl
import pytest

from isaac_api.local_import import bundle as B
from isaac_api.local_import import cli
from isaac_api.local_import import mutations as M
from isaac_api.local_import import questions as Q
from isaac_api.local_import import records as R
from isaac_api.local_import import template as T

from test_local_import_grouping_cli import _folder

BY = "SYNTHETIC TEST"
CM = B.CONTROL_MODE_PATH
POT = "context.electrochemistry.potential_setpoint_V"
SCALE = "context.electrochemistry.potential_scale"
CURRENT = "context.electrochemistry.current_setpoint_mA_cm2"


def _json(b: Path) -> dict:
    return json.loads((b / "bundle.json").read_text())


@pytest.fixture()
def grouped(tmp_path):
    src, b = _folder(tmp_path), tmp_path / "bundle"
    assert cli.main(["import", str(src), "--out", str(b)]) == 0
    M.group(b)
    return b


def _miss(bd, path, run_id):
    return next(m for m in bd["missing_values"] if m["target_field_path"] == path and m["run_id"] == run_id)


def _cand(bd, path, run_id):
    return next(c for c in bd["candidates"] if c["target_field_path"] == path and c["run_id"] == run_id)


def _group_of(bd, run_id):
    key = next(r["run_key"] for r in bd["runs"] if r["run_id"] == run_id)
    return next(g["group_id"] for g in bd["record_groups"] if key in g["member_run_keys"])


def _set_mode(b, run_id, value, *, older_build=False):
    vid = _miss(M.load(b), CM, run_id)["value_id"]
    if older_build:  # off the enum: `decide` refuses it now (contract §25); an older build stored it
        from tests_support_records import refused_then_stored

        assert "not one of the schema's allowed values" in refused_then_stored(b, "missing_value", vid, value)
        return
    M.decide(b, target_ids=[vid], decision="correct", value=value, by=BY, orcid=None)


def _compose(b, run_id):
    bd = M.load(b)
    gid = _group_of(bd, run_id)
    if next(g for g in bd["record_groups"] if g["group_id"] == gid)["status"] != "scientist_confirmed":
        M.decide(b, target_ids=[gid], decision="confirm", value=None, by=BY, orcid=None)
    result = R.build_records(b)
    entry = next(d for d in result["drafts"] if d["group_id"] == gid)
    draft = json.loads((b / entry["draft"]).read_text())
    report = json.loads((b / R.DRAFTS_DIR / f"{gid}.composition.json").read_text())
    return draft, report, entry


# --- writable, run scope, allowed values from the schema -----------------------------------------


def test_control_mode_is_a_run_scope_local_answer_field_new_in_1_5_0():
    assert B.LOCAL_ANSWER_FIELDS[CM] == B.LOCAL_ANSWER_FIELDS[SCALE] == "run"
    assert CM in B._writable_paths() and CM in B.WRITABLE_ADDED_IN["1.5.0"]
    # 1.6.0 added nothing writable; 1.7.0 the current set-point (test_local_import_current_setpoint)
    assert B._added_since("1.4.0") == frozenset({CM, CURRENT}) and B._added_since("1.5.0") == frozenset({CURRENT})
    assert not Q._is_shared(CM) and not Q._is_shared(SCALE)  # never asked once for the whole beamtime


def test_every_run_has_a_missing_value_and_nothing_is_defaulted(grouped):
    bd = M.load(grouped)
    assert not [c for c in bd["candidates"] if c["target_field_path"] == CM]  # no parser states it
    ms = [m for m in bd["missing_values"] if m["target_field_path"] == CM]
    assert {m["run_id"] for m in ms} == {r["run_id"] for r in bd["runs"] if r["run_candidate"]}
    assert all(m["value"] is None and m["value_state"] == B.STATE_MISSING and m["target_write_scope"] == "run" for m in ms)
    assert all(not m["eligibility_for_official_record"]["eligible"] for m in ms)
    assert B.control_mode_by_run(bd) == {}


def test_allowed_values_are_the_schemas_own_enum_in_questions_and_template(grouped, tmp_path):
    vendored = B._schema_node(CM)["enum"]
    target, _ = Q._target_schema()
    assert Q._node(target, CM)["enum"] == vendored  # the pinned portal schema, when installed, agrees
    qs = [q for q in Q.consolidate(M.load(grouped))["questions"] if q["kind"] == "optional_values" and CM in q["fields"]]
    assert qs and all(q["scope"].startswith("group:") for q in qs)
    assert all(q["field_allowed_values"][CM] == vendored for q in qs)
    assert T.field_rules()[CM]["allowed"] == vendored


def test_allowed_values_follow_a_refresh_of_the_pinned_list(monkeypatch):
    """Read at runtime, never transcribed: since audit round 2 the list is the export target's
    (bundle.allowed_values: the pinned vocabulary copy, which for this field equals the schema enum)."""
    pinned = dict(B._pinned_vocabulary())
    pinned[CM] = ("SYNTHETIC_mode",)
    monkeypatch.setattr(B, "_pinned_vocabulary", lambda: pinned)
    assert B.value_problem(CM, "SYNTHETIC_mode") is None
    assert "not one of the schema's allowed values" in B.value_problem(CM, "potentiostatic")


# --- eligibility: on-enum vs off-enum ----------------------------------------------------------


def test_an_enum_value_is_eligible_and_an_off_enum_value_is_confirmed_but_not_eligible(grouped):
    _set_mode(grouped, "run-001", "potentiostatic")
    _set_mode(grouped, "run-002", "Potentiostatic", older_build=True)  # not written exactly as the schema spells it
    bd = M.load(grouped)
    ok, off = _miss(bd, CM, "run-001"), _miss(bd, CM, "run-002")
    assert ok["value_state"] == off["value_state"] == B.STATE_SCIENTIST_CONFIRMED
    assert ok["eligibility_for_official_record"]["eligible"] is True
    assert off["eligibility_for_official_record"]["eligible"] is False
    assert "not one of the schema's allowed values" in off["eligibility_for_official_record"]["reason"]
    assert B.control_mode_by_run(bd) == {"run-001": "potentiostatic"}  # the off-enum answer counts for nothing
    _, report, _ = _compose(grouped, "run-002")
    assert CM not in {f["path"] for f in report["filled"]}
    [listed] = [n for n in report["not_composed"] if n.get("path") == CM]
    assert listed["value_id"] == off["value_id"] and "allowed values" in listed["reason"]
    assert POT not in {m["path"] for m in report["missing_required"]}  # an unfit mode requires nothing


# --- no inference, either way ------------------------------------------------------------------


def test_a_confirmed_setpoint_never_gives_a_control_mode(grouped):
    bd = M.load(grouped)
    M.decide(grouped, target_ids=[_cand(bd, POT, "run-001")["candidate_id"]], decision="confirm", value=None, by=BY, orcid=None)
    M.decide(grouped, target_ids=[_miss(bd, SCALE, "run-001")["value_id"]], decision="correct", value="RHE", by=BY, orcid=None)
    draft, report, _ = _compose(grouped, "run-001")
    assert draft["fields"][POT]["value"] == 0.4
    assert CM not in draft["fields"] and B.control_mode_by_run(M.load(grouped)) == {}
    assert _miss(M.load(grouped), CM, "run-001")["value_state"] == B.STATE_MISSING
    assert CM not in {m["path"] for m in report["missing_required"]}  # not required of a characterization record


def test_a_confirmed_control_mode_never_gives_a_setpoint(grouped):
    _set_mode(grouped, "run-001", "open_circuit")
    draft, report, _ = _compose(grouped, "run-001")
    assert draft["fields"][CM]["value"] == "open_circuit"
    assert POT not in draft["fields"]  # the parsed set-point was never confirmed: nothing fills it
    assert _cand(M.load(grouped), POT, "run-001")["value_state"] == B.STATE_NEEDS_REVIEW
    assert not ({POT, CURRENT} & {m["path"] for m in report["missing_required"]})  # open_circuit requires neither


# --- the set-point a confirmed mode requires ------------------------------------------------------


def test_a_confirmed_potentiostatic_mode_blocks_the_record_until_its_setpoint_composes(grouped):
    _set_mode(grouped, "run-001", "potentiostatic")
    draft, report, entry = _compose(grouped, "run-001")
    assert draft["fields"][CM]["value"] == "potentiostatic"
    assert {"path": POT, "from": "a scientist's answer"} in report["missing_required"]
    assert draft["fields"][POT] == {"value": None, "status": "missing", "evidence": []}
    assert entry["complete"] is False and POT in entry["missing_required"]
    # the set-point confirmed WITHOUT its scale is held back by the scale gate: still blocked
    bd = M.load(grouped)
    M.decide(grouped, target_ids=[_cand(bd, POT, "run-001")["candidate_id"]], decision="confirm", value=None, by=BY, orcid=None)
    _, report, entry = _compose(grouped, "run-001")
    assert any(n.get("path") == POT and n["reason"] == R.SETPOINT_NEEDS_SCALE for n in report["not_composed"])
    assert POT in entry["missing_required"]
    # with its scale the set-point composes and the requirement is met
    M.decide(grouped, target_ids=[_miss(bd, SCALE, "run-001")["value_id"]], decision="correct", value="RHE", by=BY, orcid=None)
    draft, report, entry = _compose(grouped, "run-001")
    assert draft["fields"][POT]["value"] == 0.4 and POT not in entry["missing_required"]


def test_a_confirmed_galvanostatic_mode_requires_the_current_setpoint(grouped):
    _set_mode(grouped, "run-001", "galvanostatic")
    _, report, entry = _compose(grouped, "run-001")
    assert CURRENT in entry["missing_required"] and POT not in entry["missing_required"]


def test_each_run_keeps_its_own_confirmed_mode(grouped):
    bd = M.load(grouped)
    _set_mode(grouped, "run-001", "potentiostatic")
    _set_mode(grouped, "run-002", "open_circuit")
    assert B.control_mode_by_run(M.load(grouped)) == {"run-001": "potentiostatic", "run-002": "open_circuit"}
    assert _group_of(bd, "run-001") != _group_of(bd, "run-002")  # the fixture's groups are one run each
    d1, _, _ = _compose(grouped, "run-001")
    d2, _, _ = _compose(grouped, "run-002")
    assert d1["fields"][CM]["value"] == "potentiostatic" and d2["fields"][CM]["value"] == "open_circuit"


def test_required_record_paths_reads_the_modes_from_the_schemas_rules():
    schema = copy.deepcopy(B._vendored_schema())
    base = B.required_record_paths(schema)
    assert POT not in base and CURRENT not in base
    assert POT in B.required_record_paths(schema, control_mode="potentiostatic")
    assert CURRENT in B.required_record_paths(schema, control_mode="galvanostatic")
    assert B.required_record_paths(schema, control_mode="open_circuit") == base
    # a rule a refreshed schema adds is followed without a code change
    rule = copy.deepcopy(next(r for r in schema["allOf"] if (B._control_mode_rule(r) or ("",))[0] == "potentiostatic"))
    rule["if"]["properties"]["context"]["properties"]["electrochemistry"]["properties"]["control_mode"]["const"] = "potentiodynamic"
    rule["then"]["properties"]["context"]["properties"]["electrochemistry"]["required"] = ["scan_rate_mV_s"]
    schema["allOf"].append(rule)
    assert "context.electrochemistry.scan_rate_mV_s" in B.required_record_paths(schema, control_mode="potentiodynamic")


# --- questions ---------------------------------------------------------------------------------


def test_a_confirmed_potentiostatic_mode_makes_that_groups_setpoint_question_blocking(grouped):
    def setpoint_q(bd, run_id):
        scope = f"group:{_group_of(bd, run_id)}"
        return next(q for q in Q.consolidate(bd)["questions"] if q["target_field_path"] == POT and q["scope"] == scope)

    bd = M.load(grouped)
    before1, before2 = setpoint_q(bd, "run-001"), setpoint_q(bd, "run-002")
    assert not before1["blocking"] and not before2["blocking"]
    _set_mode(grouped, "run-001", "potentiostatic")
    _set_mode(grouped, "run-002", "Potentiostatic", older_build=True)  # off-enum: requires nothing
    bd = M.load(grouped)
    after1, after2 = setpoint_q(bd, "run-001"), setpoint_q(bd, "run-002")
    assert after1["blocking"] and POT in after1["why"] and after1["question_id"] == before1["question_id"]
    assert not after2["blocking"]


def test_the_v1_questions_ask_the_setpoint_of_a_potentiostatic_run_with_no_parsed_potential(tmp_path):
    from test_local_import_reader_ssrl_xdc import make_scan
    from test_local_import_reader_two_column import make_two

    src = tmp_path / "src" / "SYN_beamtime"
    make_scan(src / "SYN_run_A.001")
    make_two(src / "SYN_cell.txt")
    b = tmp_path / "b"
    assert cli.main(["import", str(src), "--out", str(b)]) == 0
    bd = M.load(b)
    rid = next(r["run_id"] for r in bd["runs"] if r["label"] == "SYN_cell")
    pot = _miss(bd, POT, rid)

    def asked(bd):
        return [q for q in bd["questions"] if q["kind"] == "missing_value" and q["target"]["id"] == pot["value_id"]]

    assert asked(bd) == []  # optional while no mode is confirmed
    M.decide(b, target_ids=[_miss(bd, CM, rid)["value_id"]], decision="correct", value="potentiostatic", by=BY, orcid=None)
    [q] = asked(M.load(b))
    assert q["required_by"] == POT


# --- value ids: the new field is numbered last --------------------------------------------------


def test_control_mode_value_ids_come_after_every_older_fields_ids(grouped):
    ms = M.load(grouped)["missing_values"]
    nums = [int(m["value_id"].split("-")[1]) for m in ms]
    older = [n for n, m in zip(nums, ms) if m["target_field_path"] not in (CM, CURRENT)]  # CURRENT: 1.7.0, after CM
    newer = [n for n, m in zip(nums, ms) if m["target_field_path"] == CM]
    assert older == list(range(1, len(older) + 1)) and min(newer) > max(older)


# --- a 1.4.0 bundle --------------------------------------------------------------------------------


UPGRADE_14 = ("upgraded from 1.4.0; fields this build can newly write have no answer slots on runs that lacked a parsed "
              "value — re-import to get them; " + B.CHANGES_ADDED_IN["1.8.0"])  # 1.8.0: the log bounds, disclosed


def test_a_1_4_0_bundle_with_a_confirmed_mode_is_stale_not_invalid_then_rederived(grouped, capsys):
    _set_mode(grouped, "run-001", "potentiostatic")
    data = _json(grouped)
    vid = _miss(data, CM, "run-001")["value_id"]
    data["format_version"] = "1.4.0"
    for m in data["missing_values"]:
        if m["value_id"] == vid:  # what the 1.4.0 build stored: no write path for the field
            m["eligibility_for_official_record"] = {"eligible": False, "reason": "Confirmed, but this build has no write path for its official field, so it stays evidence."}
            data["readiness"]["eligible_value_count"] -= 1
    (grouped / "bundle.json").write_text(json.dumps(data))
    assert B.validate_bundle(_json(grouped))  # the plain validator refuses it
    errors, stale = B.validate_allowing_stale(_json(grouped))
    assert errors == [] and [(s["kind"], s["id"]) for s in stale] == [("missing_value", vid)]
    assert stale[0]["reason"].startswith(f"this build can now write {CM}") and stale[0]["reason"].endswith("eligible now")
    capsys.readouterr()
    assert cli.main(["bundle", "check", str(grouped)]) == cli.EXIT_INVALID
    out = capsys.readouterr().out
    assert out.startswith("STALE:") and vid in out and UPGRADE_14 in out
    result = R.build_records(grouped)
    assert [x["id"] for x in result["eligibility_rederived_now"]] == [vid]
    after = _json(grouped)
    assert after["format_version"] == B.FORMAT_VERSION and B.validate_bundle(after) == []
    [u] = after["format_upgrades"]
    assert (u["from_format_version"], u["newly_writable_paths"], u["statement"]) == ("1.4.0", [CM, CURRENT], UPGRADE_14)


def test_a_1_4_0_bundle_as_that_build_made_it_is_valid_and_upgraded_once_with_no_new_slots(grouped, capsys):
    data = _json(grouped)
    # a 1.4.0 import has no control-mode or current set-point missing values (neither was writable)
    data["missing_values"] = [m for m in data["missing_values"] if m["target_field_path"] not in (CM, CURRENT)]
    data["format_version"] = "1.4.0"
    B.refresh_questions(data)
    B.apply_derivations(data)
    (grouped / "bundle.json").write_text(json.dumps(data))
    capsys.readouterr()
    assert cli.main(["bundle", "check", str(grouped)]) == 0  # nothing at the new path: nothing is stale
    other = next(m for m in data["missing_values"] if m["target_field_path"] == "context.environment")
    M.decide(grouped, target_ids=[other["value_id"]], decision="correct", value="in_situ", by=BY, orcid=None)
    assert capsys.readouterr().err.count(UPGRADE_14) == 1
    after = _json(grouped)
    assert after["format_version"] == B.FORMAT_VERSION and [u["newly_writable_paths"] for u in after["format_upgrades"]] == [[CM, CURRENT]]
    assert not [m for m in after["missing_values"] if m["target_field_path"] in (CM, CURRENT)]  # a save adds no answer slot


# --- the metadata template ---------------------------------------------------------------------------


def _export(b, tmp_path, name="t.xlsx"):
    out = tmp_path / name
    assert cli.main(["bundle", "template", "export", str(b), "--out", str(out)]) == 0
    return out


def _col(header):
    from openpyxl.utils import get_column_letter

    return get_column_letter([c.header for c in T.FILE_COLUMNS].index(header) + 1)


def _row_of(xlsx, run_id):
    wb = openpyxl.load_workbook(xlsx)
    return next(int(v[1]) for v in wb["_isaac"].iter_rows(values_only=True) if v[0] == "File List" and v[2] == run_id)


def _import(b, xlsx, capsys, *extra):
    capsys.readouterr()
    code = cli.main(["bundle", "template", "import", str(b), str(xlsx), "--by", BY, "--json", *extra])
    out = capsys.readouterr().out
    return code, (json.loads(out) if code == 0 else None)


def test_the_template_control_mode_column_takes_the_schema_values_exactly(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    assert wb["File List"][f"{_col('Control Mode')}1"].value == "Control Mode"
    voc = wb["Vocabulary"]
    j = [c.value for c in voc[1]].index("Control Mode") + 1
    assert [voc.cell(row=i, column=j).value for i in range(2, voc.max_row + 1) if voc.cell(row=i, column=j).value] == \
        B._schema_node(CM)["enum"]
    r1, r2 = _row_of(xlsx, "run-001"), _row_of(xlsx, "run-002")
    wb["File List"][f"{_col('Control Mode')}{r1}"] = "potentiostatic"
    wb["File List"][f"{_col('Control Mode')}{r2}"] = "OCV"
    wb.save(tmp_path / "f.xlsx")
    code, rep = _import(grouped, tmp_path / "f.xlsx", capsys, "--apply")
    assert code == 0 and rep["recorded"] == 1
    assert [c["cell"] for c in rep["invalid_cells"]] == [f"'File List'!{_col('Control Mode')}{r2}"]
    bd = M.load(grouped)
    assert _miss(bd, CM, "run-001")["eligibility_for_official_record"]["eligible"]
    assert _miss(bd, CM, "run-002")["value_state"] == B.STATE_MISSING


def test_a_1_1_0_workbook_without_the_column_is_still_read(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    letter = _col("Control Mode")
    wb["File List"][f"{letter}1"] = None  # as 1.1.0 wrote it: no Control Mode column
    marker = wb["_isaac"]
    for row in marker.iter_rows():
        if row[0].value == "format_version":
            row[1].value = "1.1.0"
        if row[0].value == "File List" and row[2].value == "Control Mode":
            for cell in row:
                cell.value = None
    r1 = _row_of(xlsx, "run-001")
    wb["File List"][f"{_col('Electrolyte Concentration (M)')}{r1}"] = 0.1
    wb["File List"][f"{letter}{r1}"] = "potentiostatic"  # past pH in a 1.1.0 workbook: never read
    wb.save(tmp_path / "old.xlsx")
    code, preview = _import(grouped, tmp_path / "old.xlsx", capsys)
    # ...but SAID, in the preview, never dropped silently
    [ignored] = [n for n in preview["recorded_nowhere"] if n["cell"] == f"'File List'!{letter}{r1}"]
    assert ignored["header"] == "Control Mode" and "template '1.1.0'" in ignored["reason"] and "re-export" in ignored["reason"]
    code, rep = _import(grouped, tmp_path / "old.xlsx", capsys, "--apply")
    assert code == 0 and rep["recorded"] == 1
    bd = M.load(grouped)
    assert _miss(bd, "context.electrochemistry.electrolyte.concentration_M", "run-001")["value_state"] == B.STATE_SCIENTIST_CONFIRMED
    assert _miss(bd, CM, "run-001")["value_state"] == B.STATE_MISSING
    # the same workbook claiming 1.2.0 lacks a column 1.2.0 writes: refused, nothing recorded
    wb = openpyxl.load_workbook(tmp_path / "old.xlsx")
    for row in wb["_isaac"].iter_rows():
        if row[0].value == "format_version":
            row[1].value = "1.2.0"
    wb.save(tmp_path / "claims_new.xlsx")
    code, _ = _import(grouped, tmp_path / "claims_new.xlsx", capsys, "--apply")
    assert code == T.EXIT_INPUT_REFUSED


# --- review fixes -----------------------------------------------------------------------------------


def test_field_allowed_values_lists_only_what_eligibility_refuses_outside_of(grouped):
    listed: dict[str, list] = {}
    for q in Q.consolidate(M.load(grouped))["questions"]:
        if q["kind"] == "optional_values":
            listed.update(q["field_allowed_values"])
            for f in q["fields"]:
                allowed = B.allowed_values(f)  # the ONE list decide checks and the export validates against
                if allowed:
                    assert q["field_allowed_values"][f] == allowed
                else:
                    assert f not in q["field_allowed_values"]
    assert CM in listed and SCALE in listed
    for f, values in listed.items():  # an answer outside each list really is never eligible
        assert B.value_problem(f, "SYNTHETIC_not_listed") is not None
        assert all(B.value_problem(f, v) is None for v in values)


def test_an_unreadable_vendored_schema_makes_nothing_eligible(monkeypatch):
    monkeypatch.setattr(B, "_vendored_schema", lambda: None)
    problem = B.value_problem(CM, "bogus")
    assert problem is not None and "could not be read" in problem
    assert B.value_problem(CM, "potentiostatic") is not None
    fit = B.eligibility_for_official_record(B.STATE_SCIENTIST_CONFIRMED, target_field_path=CM, value="bogus",
                                            writable_paths=frozenset({CM}))
    assert fit["eligible"] is False


def _potentiostatic_rule(schema):
    return copy.deepcopy(next(r for r in schema["allOf"] if (B._control_mode_rule(r) or ("",))[0] == "potentiostatic"))


@pytest.mark.parametrize("mutate", [
    pytest.param(lambda r: r["if"]["properties"]["context"]["properties"]["electrochemistry"]["properties"].update(
        control_mode={"enum": ["potentiostatic", "mixed"]}), id="enum-instead-of-const"),
    pytest.param(lambda r: r["if"]["properties"].update(record_domain={"const": "performance"}), id="extra-condition"),
    pytest.param(lambda r: r["if"]["properties"]["context"]["properties"]["electrochemistry"]["properties"].update(
        pH={"const": 7}), id="extra-electrochemistry-condition"),
    pytest.param(lambda r: r.update(then={"required": ["descriptors"]}), id="then-outside-electrochemistry"),
    pytest.param(lambda r: r["then"]["properties"]["context"]["properties"]["electrochemistry"].update(
        properties={"potential_setpoint_V": {"minimum": 0}}), id="then-with-constraints"),
    pytest.param(lambda r: r.update({"else": {"required": ["links"]}}), id="else-branch"),
])
def test_a_control_mode_rule_this_build_cannot_interpret_fails_closed(mutate):
    schema = copy.deepcopy(B._vendored_schema())
    rule = _potentiostatic_rule(schema)
    mutate(rule)
    schema["allOf"].append(rule)
    with pytest.raises(B.SchemaRuleNotUnderstood, match="control_mode"):
        B.required_record_paths(schema)  # whatever mode is asked: an uninterpreted rule is never skipped


def test_the_vendored_and_pinned_schemas_rules_are_understood():
    target, _ = Q._target_schema()
    for schema in (B._vendored_schema(), target):
        modes = {k[0] for r in schema["allOf"] if (k := B._control_mode_rule(r))}
        assert modes == {"potentiostatic", "galvanostatic"}


def _with_mode_candidate(bd, decision=None, value="potentiostatic"):
    bd = copy.deepcopy(bd)
    base = _cand(bd, POT, "run-001")
    cand = {**copy.deepcopy(base), "candidate_id": "cand-9999", "target_field_path": CM, "value": value,
            "observed_literals": [value]}
    bd["candidates"].append(cand)
    if decision is not None:
        bd.setdefault("scientist_decisions", []).append({"decision_id": "d-09999", "target": {"kind": "candidate", "id": "cand-9999"},
                                                       **decision})
    return bd


def test_a_parser_mode_counts_only_once_confirmed_and_as_decided(grouped):
    bd = M.load(grouped)
    assert B.control_mode_by_run(_with_mode_candidate(bd)) == {}  # unconfirmed
    assert B.control_mode_by_run(_with_mode_candidate(bd, {"decision": "reject", "value": None})) == {}
    assert B.control_mode_by_run(_with_mode_candidate(bd, {"decision": "confirm", "value": None})) == {"run-001": "potentiostatic"}
    corrected = _with_mode_candidate(bd, {"decision": "correct", "value": "open_circuit"})
    assert B.control_mode_by_run(corrected) == {"run-001": "open_circuit"}
    off = _with_mode_candidate(bd, {"decision": "confirm", "value": None}, value="Potentiostatic")
    assert B.control_mode_by_run(off) == {}
    # a rejected parser mode does not make the set-point blocking
    rejected = _with_mode_candidate(bd, {"decision": "reject", "value": None})
    scope = f"group:{_group_of(bd, 'run-001')}"
    q = next(q for q in Q.consolidate(rejected)["questions"] if q["target_field_path"] == POT and q["scope"] == scope)
    assert not q["blocking"]


def _two_run_group(bd):
    return {"group_id": "rg-SYNTHETIC-two-runs", "status": "scientist_confirmed",
            "member_run_keys": [r["run_key"] for r in bd["runs"] if r["run_candidate"]]}


def test_two_different_confirmed_modes_in_one_group_are_a_disagreement_and_nothing_is_filled(grouped):
    _set_mode(grouped, "run-001", "potentiostatic")
    _set_mode(grouped, "run-002", "open_circuit")
    bd = M.load(grouped)
    draft, report = R.compose_group(bd, _two_run_group(bd), grouped)
    [dis] = [d for d in report["disagreements"] if d["path"] == CM]
    assert sorted(dis["values"]) == ["open_circuit", "potentiostatic"]
    assert CM not in draft["fields"] and CM not in {f["path"] for f in report["filled"]}
    assert POT not in {m["path"] for m in report["missing_required"]}  # no mode chosen, so none requires anything


def test_one_answered_run_gives_a_merged_group_its_mode(grouped):
    """CURRENT BEHAVIOUR, flagged for domain-owner review (contract doc §20): one run answered
    potentiostatic and one unanswered composes potentiostatic for the group, as for any run-scope field."""
    _set_mode(grouped, "run-001", "potentiostatic")
    bd = M.load(grouped)
    draft, report = R.compose_group(bd, _two_run_group(bd), grouped)
    assert draft["fields"][CM]["value"] == "potentiostatic"
    assert POT in {m["path"] for m in report["missing_required"]}
