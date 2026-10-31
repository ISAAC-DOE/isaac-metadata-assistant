"""Skill polish, cycle 2 (small CLI fixes). SYNTHETIC fixtures only.

* ``bundle show <dir> question <id>`` accepts the consolidated ids ``bundle questions`` prints (and still the
  older fine-grained ones), prints the question as ``bundle questions`` does, records the same exposures, and
  exits 6 for an id that is neither.
* ``summary.md`` and the ``isaac import`` output say an Athena ``.prj`` is read by ``isaac process``; the
  bundle's stored unread_files and counts are unchanged.
"""

from __future__ import annotations

import json
import shutil

import pytest

from isaac_api.local_import import cli
from isaac_api.local_import import mutations as M
from isaac_api.local_import import questions as Q

from test_local_import_records import SEMANTICS, bdir, pristine  # noqa: F401  (fixtures)


def _run(capsys, *argv):
    code = cli.main(list(argv))
    cap = capsys.readouterr()
    return code, cap.out, cap.err


# --- 1. show question takes a consolidated id --------------------------------------------------------


def test_show_question_accepts_every_consolidated_id_bundle_questions_prints(bdir, capsys):
    code, out, _ = _run(capsys, "bundle", "questions", str(bdir), "--limit", "200")
    assert code == 0
    printed = json.loads(out)["questions"]
    assert printed
    for q in printed[:12]:
        code, shown, err = _run(capsys, "bundle", "show", str(bdir), "question", q["question_id"])
        assert code == 0, (q["question_id"], err)
        one = json.loads(shown)
        assert one["question_id"] == q["question_id"] and one["kind"] == q["kind"] and one["text"] == q["text"]
        assert "targets" not in one and "targets_count" in one  # compact, like the page


def test_show_question_still_accepts_a_fine_grained_id(bdir, capsys):
    fine = M.load(bdir)["questions"][0]["question_id"]
    code, shown, _ = _run(capsys, "bundle", "show", str(bdir), "question", fine)
    assert code == 0 and json.loads(shown)["question_id"] == fine


def test_show_question_unknown_id_exits_6(bdir, capsys):
    code, out, err = _run(capsys, "bundle", "show", str(bdir), "question", "Q-does-not-exist")
    assert code == cli.EXIT_NOT_FOUND == 6 and not out and "Q-does-not-exist" in err


def test_show_question_records_the_exposures_bundle_questions_records(bdir, tmp_path, capsys):
    twin = tmp_path / "twin"
    shutil.copytree(bdir, twin)
    payload = Q.consolidate(M.load(bdir))
    with_hints = [q for q in payload["questions"] if any(h.get("candidate_id") for h in q["hints"])]
    assert with_hints, "the fixture must have a question whose hints print parser values"
    q = with_hints[0]
    # `bundle questions` for just that question's page records its hints' candidates ...
    assert M.load(bdir).get("exposures") == []
    index = [x["question_id"] for x in Q.filter_questions(payload, scope=None, blocking_only=False)].index(q["question_id"])
    code, _, _ = _run(capsys, "bundle", "questions", str(bdir), "--offset", str(index), "--limit", "1")
    assert code == 0
    expected = {(e["target"]["kind"], e["target"]["id"]) for e in M.load(bdir)["exposures"]}
    assert expected
    # ... and `show question` records exactly those, before printing
    assert M.load(twin).get("exposures") == []
    code, shown, _ = _run(capsys, "bundle", "show", str(twin), "question", q["question_id"])
    assert code == 0 and json.loads(shown)["hints"]
    assert {(e["target"]["kind"], e["target"]["id"]) for e in M.load(twin)["exposures"]} == expected


def test_show_question_without_parser_values_records_no_exposure(bdir, capsys):
    payload = Q.consolidate(M.load(bdir))
    quiet = next(q for q in payload["questions"] if not any(h.get("candidate_id") for h in q["hints"])
                 and q["kind"] != "resolve_conflict")
    code, _, _ = _run(capsys, "bundle", "show", str(bdir), "question", quiet["question_id"])
    assert code == 0 and not M.load(bdir).get("exposures")


# --- 2. Athena projects are said to be read by processing --------------------------------------------


@pytest.fixture()
def with_prj(tmp_path, capsys):
    src = tmp_path / "beamtime"
    shutil.copytree(SEMANTICS, src)
    for n in range(4):
        (src / f"SYN_analysis_{n}.prj").write_bytes(b"\x1f\x8b SYNTHETIC not a real project")
    out = tmp_path / "b"
    assert cli.main(["import", str(src), "--out", str(out)]) == 0
    return out, capsys.readouterr().out


def test_import_output_says_athena_projects_are_read_by_process(with_prj):
    _, printed = with_prj
    assert "4 Athena projects are read by `isaac process`, not at import" in printed


def test_summary_says_athena_projects_are_read_by_process_and_not_unread(with_prj):
    out, _ = with_prj
    text = (out / "summary.md").read_text(encoding="utf-8")
    assert "4 Athena projects are read by `isaac process`, not at import" in text
    assert "no_reader_for_source_type" not in "\n".join(line for line in text.splitlines() if ".prj" in line)
    assert len(text.encode("utf-8")) <= cli.SUMMARY_TARGET_BYTES


def test_the_stored_bundle_still_lists_the_projects_as_unread(with_prj):
    out, _ = with_prj
    bundle = json.loads((out / "bundle.json").read_text())
    prj = [u for u in bundle["unread_files"] if u["relative_path"].endswith(".prj")]
    assert len(prj) == 4 and {u["reason_code"] for u in prj} == {"no_reader_for_source_type"}
    assert bundle["counts"]["unread_files"] == len(bundle["unread_files"])


def test_no_athena_line_when_there_are_no_projects(imported_plain):
    out, printed = imported_plain
    assert "Athena project" not in printed and "Athena project" not in (out / "summary.md").read_text()


@pytest.fixture()
def imported_plain(tmp_path, capsys):
    out = tmp_path / "plain"
    assert cli.main(["import", str(SEMANTICS), "--out", str(out)]) == 0
    return out, capsys.readouterr().out


# --- 3. the template's facility suggestion shows the facility hints (the optional-values question's) -----


def _campaign_suggestions(xlsx) -> dict[str, str | None]:
    import openpyxl

    sheet = openpyxl.load_workbook(xlsx)["ISAAC Suggestions"]
    rows = list(sheet.iter_rows(values_only=True))
    start = next(n for n, r in enumerate(rows) if r[0] == "Campaign Info")
    return {r[0]: r[1] for r in rows[start + 1:] if r[0]}


@pytest.fixture()
def exported(bdir, tmp_path):
    M.group(bdir)
    out = tmp_path / "t.xlsx"
    assert cli.main(["bundle", "template", "export", str(bdir), "--out", str(out)]) == 0
    return bdir, out


@pytest.mark.parametrize("label", ["Facility Name", "Facility organization (e.g. SLAC / SSRL)", "Beamline/Instrument"])
def test_facility_suggestion_shows_the_folder_name_hint_with_its_source(exported, label):
    _, xlsx = exported
    cell = _campaign_suggestions(xlsx)[label]
    assert cell, f"{label}: a facility hint exists in the bundle (the imported folder's name) but the suggestion is empty"
    assert "hint (folder_name, never an answer): 'multi_operator_corpus'" in cell
    assert "the imported folder's name" in cell and "(source)" not in cell


def test_facility_answer_cells_stay_empty(exported):
    import openpyxl

    _, xlsx = exported
    campaign = openpyxl.load_workbook(xlsx)["Campaign Info"]
    answers = {r[0]: r[1] for r in campaign.iter_rows(values_only=True) if r[0]}
    assert answers["Facility Name"] is None and answers["Beamline/Instrument"] is None
