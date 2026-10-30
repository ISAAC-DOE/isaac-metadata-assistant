"""``isaac bundle template adopt``: the scientist's OWN filled copy of ISAAC's metadata template. Synthetic only.

A scientist who already filled ISAAC's template (no ``_isaac`` sheet, File List keyed by their own
``_dir`` folder names) should not retype it. Adopt writes the bundle's own template with THEIR value in
an answer cell only where no file reading contradicts it; a value a reading contradicts (a dropped
minus sign against ``m800mV``, pH 2 against ``pH12``, a beamline against the folder's AI reading) or
that the field does not take goes to the Disagreements sheet with an empty "Your ruling" column, and
``template import`` records none of them until the scientist rules.
"""

from __future__ import annotations

import json
import shutil

import openpyxl
import pytest

from isaac_api.local_import import bundle as B
from isaac_api.local_import import cli
from isaac_api.local_import import mutations as M
from isaac_api.local_import import template as T

from tests_support_records import SEMANTICS

WHO = "Synthetic Scientist"
POT = "context.electrochemistry.potential_setpoint_V"
PH = "context.electrochemistry.pH"
BEAMLINE = "system.facility.beamline"
THEIRS = "SYN_their_metadata.xlsx"
#: the synthetic corpus renamed so one run's name says -0.8 V (``m800mV``, no parsed value) and one pH 12
RENAMES = {"05_02_ZZ2_base_after400Cycling_filter20_800mV": "05_02_ZZ2_base_after400Cycling_filter20_m800mV",
           "06_03_ZZ3_alk_fresh_f30_900mV": "06_03_ZZ3_alk_pH12_f30_900mV"}
#: their File List, in the shape a filled copy has (renamed columns, extra columns, folder names)
THEIR_HEADERS = ["Folder Name", "Config ID", "Sample Type", "Sample ", "Sample Form", "Environment", "Electrolyte", "pH",
                 "Potential (V)", "Potential Scale", "Temperature ©"]


@pytest.fixture(scope="module")
def corpus(tmp_path_factory):
    src = tmp_path_factory.mktemp("adopt-src") / "SYN_BL99_beamtime"
    shutil.copytree(SEMANTICS, src)
    for p in sorted(src.rglob("*"), key=lambda p: -len(p.parts)):
        for old, new in RENAMES.items():
            if old in p.name:
                p.rename(p.with_name(p.name.replace(old, new)))
    return src


@pytest.fixture(scope="module")
def grouped_base(corpus, tmp_path_factory):
    bdir = tmp_path_factory.mktemp("adopt") / "b"
    assert cli.main(["import", str(corpus), "--out", str(bdir)]) == 0
    M.group(bdir)
    b = M.load(bdir)
    # an AI reading of the beamtime folder's name, and one of THEIR sheet (which is not a file reading)
    findings = {"format": "isaac.reading_findings", "findings": [
        {"field_path": BEAMLINE, "value": "SYN-7", "applies_to": {"scope": "experiment"},
         "source": {"path": ".", "locator": "folder name"}, "quote": "SYN_BL99_beamtime", "strength": "inferred",
         "reasoning": "folder name token"},
        {"field_path": BEAMLINE, "value": "62b", "applies_to": {"scope": "experiment"},
         "source": {"path": THEIRS, "locator": "Campaign Info!B5"}, "quote": "Beamline/Instrument 62b",
         "strength": "stated", "reasoning": "the sheet"},
    ]}
    rep = M.annotate_from_readings(bdir, findings, model_id="synthetic-model", apply=True)
    assert rep["interpretations"] >= 2, rep
    assert {r["run_id"] for r in b["runs"]} >= {"run-001", "run-005", "run-007", "run-009"}
    return bdir


@pytest.fixture()
def grouped(grouped_base, tmp_path):
    bdir = tmp_path / "b"
    shutil.copytree(grouped_base, bdir)
    return bdir


def _stems(bdir):
    return {r["run_id"]: r["stem"] for r in M.load(bdir)["runs"]}


def _their_sheet(bdir, path, *, rows=None, beamline="62b"):
    """A filled copy of ISAAC's metadata template, as a scientist sends it: sheets and Campaign Info
    labels of the template, File List keyed by ``<stem>_dir``."""
    stems = _stems(bdir)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Campaign Info"
    ws["A1"] = "ISAAC Campaign Metadata"
    for r, (label, value) in enumerate([("Facility Name", "Synthetic Light Source"), ("Organization", "Synthetic Lab"),
                                        ("Beamline/Instrument", beamline), ("Start Date", "2099-01-01"),
                                        ("Proposal ID", "X-0")], start=3):
        ws[f"A{r}"], ws[f"B{r}"] = label, value
    wc = wb.create_sheet("Configurations")
    wc.append(["Config ID", "Technique", "Mode", "Spectrometer"])
    wc.append(["Config_A", "HERFD-XAS", "fluorescence", "five analyzers"])
    wf = wb.create_sheet("File List")
    wf.append(THEIR_HEADERS)
    default = {"pH": 13, "Potential (V)": None}
    for run_id in rows or ["run-001", "run-002", "run-003", "run-004", "run-005", "run-006", "run-007", "run-008"]:
        stem = stems[run_id]
        potential = {"run-005": 0.8, "run-007": 0.9}.get(run_id)
        if potential is None:
            potential = next(c["value"] for c in M.load(bdir)["candidates"] if c["run_id"] == run_id and c["target_field_path"] == POT)
        values = {"Folder Name": f"{stem}_dir", "Config ID": "Config_A", "Sample Type": "powder_on_electrode",
                  "Sample ": "ZZ synthetic catalyst", "Sample Form": "powder", "Environment": "aqueous_electrochemistry",
                  "Electrolyte": "0.1 M KOH", "pH": 2 if run_id == "run-007" else default["pH"], "Potential (V)": potential,
                  "Potential Scale": "RHE", "Temperature ©": 25}
        wf.append([values[h] for h in THEIR_HEADERS])
    wf.append(["ZZ_not_a_run_dir", "Config_A", None, "ZZ", "powder"])
    wv = wb.create_sheet("Vocabulary")
    wv.append(["Sample Form"])
    wv.sheet_state = "hidden"
    wb.save(path)
    return path


def _adopt(bdir, theirs, out, capsys, *extra):
    capsys.readouterr()
    code = cli.main(["bundle", "template", "adopt", str(bdir), str(theirs), "--out", str(out), *extra])
    captured = capsys.readouterr()
    return code, captured.out, captured.err


def _adopt_json(bdir, theirs, out, capsys):
    code, text, err = _adopt(bdir, theirs, out, capsys, "--json")
    assert code == 0, err
    return json.loads(text)["adopted"]


def _import(bdir, xlsx, capsys, *extra):
    capsys.readouterr()
    code = cli.main(["bundle", "template", "import", str(bdir), str(xlsx), "--by", WHO, "--json", *extra])
    captured = capsys.readouterr()
    return code, (json.loads(captured.out) if code == 0 else None), captured.err


def _row_of(xlsx, run_id):
    for values in openpyxl.load_workbook(xlsx)["_isaac"].iter_rows(values_only=True):
        if values[0] == "File List" and values[2] == run_id:
            return int(values[1])
    raise AssertionError(run_id)


def _col(header):
    from openpyxl.utils import get_column_letter

    return get_column_letter([c.header for c in T.FILE_COLUMNS].index(header) + 1)


def _miss(b, path, run_id):
    return next(m for m in b["missing_values"] if m["target_field_path"] == path and m["run_id"] == run_id)


def _cand(b, path, run_id):
    return next(c for c in b["candidates"] if c["target_field_path"] == path and c["run_id"] == run_id)


def _d(adopted, header, kind="file_reading", run_id=None):
    hits = [d for d in adopted["disagreements"] if d["header"] == header and d["kind"] == kind
            and (run_id is None or run_id in d["runs"])]
    assert len(hits) == 1, (header, kind, run_id, adopted["disagreements"])
    return hits[0]


# --- what adopt reads, and refuses ----------------------------------------------------------------


def test_adopt_refuses_anything_but_a_filled_copy_of_the_template_and_writes_nothing(grouped, tmp_path, capsys):
    out = tmp_path / "adopted.xlsx"
    other = openpyxl.Workbook()
    other.active["A1"] = "my own sample list"
    other.save(tmp_path / "other.xlsx")
    code, _, err = _adopt(grouped, tmp_path / "other.xlsx", out, capsys)
    assert code == 3 and "not a filled copy of ISAAC's metadata template" in err and not out.exists()

    # ISAAC's blank template: the template's sheets and example values, an empty File List
    blank = openpyxl.Workbook()
    blank.active.title = "Campaign Info"
    blank.active["A1"], blank.active["A3"], blank.active["B3"] = "ISAAC Campaign Metadata", "Facility Name", "Example"
    blank.create_sheet("Configurations").append(["Config ID", "Technique"])
    blank.create_sheet("File List").append(["File Name", "Config ID", "Sample Name"])
    blank.save(tmp_path / "blank.xlsx")
    code, _, err = _adopt(grouped, tmp_path / "blank.xlsx", out, capsys)
    assert code == 3 and "blank metadata template" in err and not out.exists()

    # a template ISAAC exported for a bundle is imported, not adopted
    assert cli.main(["bundle", "template", "export", str(grouped), "--out", str(tmp_path / "exported.xlsx")]) == 0
    code, _, err = _adopt(grouped, tmp_path / "exported.xlsx", out, capsys)
    assert code == 3 and "template import" in err and not out.exists()
    assert not M.load(grouped)["scientist_decisions"]


def test_rows_map_to_runs_by_folder_name_and_unmatched_rows_and_runs_are_listed(grouped, tmp_path, capsys):
    theirs = _their_sheet(grouped, tmp_path / THEIRS)
    a = _adopt_json(grouped, theirs, tmp_path / "adopted.xlsx", capsys)
    assert a["rows_matched"] == 8 and a["runs_matched"] == 8
    assert [r["name"] for r in a["unmatched_rows"]] == ["ZZ_not_a_run_dir"]
    assert [r["run_id"] for r in a["unmatched_runs"]] == ["run-009"]
    # their renamed / extra columns: read under ISAAC's name, or listed as not read; never guessed
    assert {"header": "Sample", "to": "Sample Name", "by_alias": True, "column": "D"} in a["their_columns_mapped"]
    assert {c["header"] for c in a["their_columns_not_read"]} == {"Sample Type", "Temperature ©"}


def test_agreeing_values_are_adopted_and_each_cites_their_cell(grouped, tmp_path, capsys):
    theirs = _their_sheet(grouped, tmp_path / THEIRS)
    out = tmp_path / "adopted.xlsx"
    _adopt_json(grouped, theirs, out, capsys)
    wb = openpyxl.load_workbook(out)
    fl = wb["File List"]
    r1 = _row_of(out, "run-001")
    cell = fl[f"{_col('Potential (V)')}{r1}"]
    assert cell.value == 0.1 and f"{THEIRS} 'File List'!I2" in cell.comment.text  # agrees with the parsed 100mV
    assert fl[f"{_col('Sample Form')}{r1}"].value == "powder"
    assert fl[f"{_col('Sample Name')}{r1}"].value == "ZZ synthetic catalyst"
    assert fl[f"{_col('Config ID')}{r1}"].value == "Config_A"
    assert wb["Configurations"]["A2"].value == "Config_A" and wb["Configurations"]["B2"].value == "HERFD-XAS"
    assert wb["Campaign Info"]["B3"].value == "Synthetic Light Source"
    assert wb["Campaign Info"]["B4"].value == "Synthetic Lab"  # their "Organization" is the facility's
    assert wb.sheetnames == ["Campaign Info", "Configurations", "File List", "Vocabulary", "Disagreements",
                             "ISAAC Suggestions", "_isaac"]


def test_a_value_a_file_reading_contradicts_goes_to_disagreements_and_its_answer_cell_stays_empty(grouped, tmp_path, capsys):
    theirs = _their_sheet(grouped, tmp_path / THEIRS)
    out = tmp_path / "adopted.xlsx"
    a = _adopt_json(grouped, theirs, out, capsys)
    wb = openpyxl.load_workbook(out)
    fl = wb["File List"]

    # the minus sign dropped: the parser reads no value from m800mV, its file name token says -0.8
    minus = _d(a, "Potential (V)", run_id="run-005")
    assert minus["their_value"] == 0.8 and minus["their_cells"] == ["'File List'!I6"]
    assert any("-0.8" in r and "m800mV" in r for r in minus["readings"])
    assert fl[f"{_col('Potential (V)')}{_row_of(out, 'run-005')}"].value is None
    assert "D" in fl[f"{_col('Potential (V)')}{_row_of(out, 'run-005')}"].comment.text

    # pH 2 typed over the file name's pH12 (a parsed value AND the name token)
    ph = _d(a, "pH", run_id="run-007")
    assert ph["their_value"] == 2 and any("12" in r and "extracted" in r for r in ph["readings"])
    assert fl[f"{_col('pH')}{_row_of(out, 'run-007')}"].value is None
    assert fl[f"{_col('pH')}{_row_of(out, 'run-001')}"].value == 13  # nothing states run-001's pH: adopted

    # the beamline against the folder's AI reading; their sheet's own AI reading is not a file's
    bl = _d(a, "Beamline/Instrument")
    assert bl["their_value"] == "62b" and bl["scope"] == "experiment"
    assert [r for r in bl["readings"] if "'SYN-7'" in r] and not [r for r in bl["readings"] if "62b" in r]
    assert wb["Campaign Info"]["B5"].value is None
    assert a["ai_readings_of_their_sheet_not_compared"] >= 1

    # the Disagreements sheet: their value, the file's reading with its source, an EMPTY ruling
    wd = wb["Disagreements"]
    assert [c.value for c in wd[1]] == list(T.DISAGREE_HEADERS)
    row = next(r for r in range(2, wd.max_row + 1) if wd[f"A{r}"].value == minus["id"])
    assert wd[f"D{row}"].value == 0.8 and "m800mV" in wd[f"F{row}"].value and wd[f"H{row}"].value is None


def test_not_allowed_and_name_only_values_are_listed_once_each_with_the_allowed_list_once(grouped, tmp_path, capsys):
    theirs = _their_sheet(grouped, tmp_path / THEIRS)
    out = tmp_path / "adopted.xlsx"
    code, text, err = _adopt(grouped, theirs, out, capsys)
    assert code == 0, err
    a = _adopt_json(grouped, theirs, tmp_path / "again.xlsx", capsys)
    env = _d(a, "Environment", kind="not_allowed")
    assert env["their_value"] == "aqueous_electrochemistry" and len(env["answer_cells"]) == 8
    assert a["allowed_values"]["Environment"] == B.allowed_values("context.environment")
    name = _d(a, "Electrolyte", kind="name_only")
    assert "put 0.1 in 'Electrolyte Concentration (M)'" in name["reason"] and len(name["answer_cells"]) == 8
    fl = openpyxl.load_workbook(out)["File List"]
    assert all(fl[f"{_col(h)}{_row_of(out, 'run-001')}"].value is None for h in ("Environment", "Electrolyte"))
    # the compact summary names every disagreement; each allowed list appears once
    assert text.count(", ".join(B.allowed_values("context.environment"))) == 1
    assert "run-005 Potential (V): yours 0.8" in text and "m800mV" in text
    assert "run-007 pH: yours 2" in text
    assert "Beamline/Instrument: yours '62b'" in text and "SYN-7" in text
    assert "ZZ_not_a_run_dir" in text and "run-009" not in text.split("Runs on none")[0].split("Rows:")[0]
    assert len(text) < 6_000


# --- importing the adopted file -------------------------------------------------------------------


def test_importing_the_adopted_file_records_no_disagreeing_value(grouped, tmp_path, capsys):
    theirs = _their_sheet(grouped, tmp_path / THEIRS)
    out = tmp_path / "adopted.xlsx"
    a = _adopt_json(grouped, theirs, out, capsys)
    code, rep, err = _import(grouped, out, capsys, "--apply")
    assert code == 0, err
    assert rep["recorded"] > 0 and not rep["invalid_cells"] and not rep["clashes"]
    assert {x["id"] for x in rep["awaiting_ruling"]} == {d["id"] for d in a["disagreements"]}
    b = M.load(grouped)
    latest = B._decisions_by_target(b["scientist_decisions"])
    assert _miss(b, POT, "run-005")["value_state"] == B.STATE_MISSING
    assert ("candidate", _cand(b, PH, "run-007")["candidate_id"]) not in latest
    assert not any(k == ("missing_value", m["value_id"]) for k in latest for m in b["missing_values"] if m["target_field_path"] == BEAMLINE)
    assert not [d for d in b["scientist_decisions"] if d.get("value") in ("aqueous_electrochemistry", "0.1 M KOH", "62b")]
    assert ("missing_value", _miss(b, POT, "run-005")["value_id"]) not in latest  # 0.8 is confirmed only where a file says 800mV
    # what WAS recorded cites their sheet's cell, through the adopted file
    pot1 = latest[("candidate", _cand(b, POT, "run-001")["candidate_id"])]
    assert pot1["decision"] == "confirm" and f"adopted from {THEIRS} 'File List'!I2" in pot1["note"]
    assert not rep["replaces_or_contradicts_a_file_reading"]


def test_a_ruling_typed_in_the_disagreements_sheet_is_imported_for_exactly_its_targets(grouped, tmp_path, capsys):
    theirs = _their_sheet(grouped, tmp_path / THEIRS)
    out = tmp_path / "adopted.xlsx"
    a = _adopt_json(grouped, theirs, out, capsys)
    minus, env = _d(a, "Potential (V)", run_id="run-005"), _d(a, "Environment", kind="not_allowed")
    wb = openpyxl.load_workbook(out)
    wd = wb["Disagreements"]
    rows = {wd[f"A{r}"].value: r for r in range(2, wd.max_row + 1)}
    wd[f"H{rows[minus['id']]}"] = -0.8
    wd[f"H{rows[env['id']]}"] = "in_situ"
    wb.save(tmp_path / "ruled.xlsx")
    code, rep, err = _import(grouped, tmp_path / "ruled.xlsx", capsys)
    assert code == 0, err and rep["recorded"] == 0
    assert minus["id"] not in {x["id"] for x in rep["awaiting_ruling"]}
    code, rep, err = _import(grouped, tmp_path / "ruled.xlsx", capsys, "--apply")
    assert code == 0, err
    b = M.load(grouped)
    latest = B._decisions_by_target(b["scientist_decisions"])
    d = latest[("missing_value", _miss(b, POT, "run-005")["value_id"])]
    assert d["value"] == -0.8 and f"'Disagreements'!H{rows[minus['id']]}" in d["note"] and minus["id"] in d["note"]
    envs = {m["run_id"]: latest.get(("missing_value", m["value_id"])) for m in b["missing_values"]
            if m["target_field_path"] == "context.environment"}
    assert {r for r, dd in envs.items() if dd and dd["value"] == "in_situ"} == set(env["runs"])
    assert envs["run-009"] is None  # not on their sheet: the ruling never reaches it


def test_a_ruling_on_a_moved_disagreements_row_is_not_read(grouped, tmp_path, capsys):
    theirs = _their_sheet(grouped, tmp_path / THEIRS)
    out = tmp_path / "adopted.xlsx"
    _adopt_json(grouped, theirs, out, capsys)
    wb = openpyxl.load_workbook(out)
    wd = wb["Disagreements"]
    wd["A2"], wd["H2"] = "D99", "x"
    wb.save(tmp_path / "moved.xlsx")
    code, rep, err = _import(grouped, tmp_path / "moved.xlsx", capsys)
    assert code == 0, err
    assert [r for r in rep["rows_refused"] if "'Disagreements'!H2" in r["cells"]]
    assert not [c for c in rep["decisions"] if "Disagreements" in c["cell"]]
