"""Polish c7-r: review of c7-s (answer slots, the beamline hint, the nearest record). Synthetic only.

* HIGH-1: a beamtime-wide answer that a group's own run-scoped answer contradicts does not reach that group: the
  decide reply and the template import name it as NOT reached (it keeps its own answer and fills nothing).
* MED-2: the minimum's command for a group whose held targets mix a beamtime target and a run's own answer is the
  beamtime question answered for that group (``--only group:<id>``), never ``--target <beamtime target>``.
* MED-3: an answer slot belongs to its run: a merged group fills a field from slots only when every member run is
  covered (by a slot or its own answer of the same value, or the beamtime answer).
* MED-4: a value the nearest record's warnings state is an exposure, wherever the page gets the minimum.
* MED-5: the beamline hint compares the first beamline-shaped token, its separator kept, a facility word ignored.
* LOW-6: a web dry run of ``only=group:<id>`` on a beamtime question previews that group.
* LOW-8: a slot no decision names is dropped on any save; a slot holding no answer is not asked once the beamtime
  question is answered.
"""

from __future__ import annotations

import json
import shutil

import openpyxl
import pytest
from fastapi.testclient import TestClient

from isaac_api import local_bundles as LB
from isaac_api.app import create_app

from isaac_api.local_import import bundle as B
from isaac_api.local_import import cli
from isaac_api.local_import import mutations as M
from isaac_api.local_import import questions as Q
from isaac_api.local_import import records as R
from isaac_api.local_import import template as T

from test_local_import_polish_c7s import TECH, _decide, _latest, _sheet, _technique_question, plain, plain_base  # noqa: F401
from test_local_import_template_adopt import THEIRS, WHO, _adopt_json, _d, _import  # noqa: F401

GROUP = "attribution.produced_by.group"


def _groups(bdir):
    return [g["group_id"] for g in Q.live_groups(M.load(bdir))]


def _composed(bdir, path):
    b = M.load(bdir)
    index = R.composition_index(b)
    out = {}
    for g in Q.live_groups(b):
        draft, rep = R.compose_group(b, g, bdir, index)
        out[g["group_id"]] = ((draft["fields"].get(path) or {}).get("value"), rep)
    return out


def _attribution_question(bdir):
    b = M.load(bdir)
    for q in Q.consolidate(b)["questions"]:
        if q["kind"] not in Q.VALUE_KINDS:
            continue
        for field, _ids in Q._value_paths(q):
            if field == GROUP and "experiment" in (q.get("per_scope_targets") or {}):
                return q["question_id"]
    raise AssertionError("no beamtime-wide attribution question")


# --- MED-5: the beamline normaliser -----------------------------------------------------------------------


@pytest.mark.parametrize("spelled", ["SSRL BL15-2", "BL15-2 (SSRL)", "SSRL 15-2", "BL15–2", "BL 15—2", "bl15.2"])
def test_the_beamline_hint_reads_the_beamline_inside_a_facility_word_and_any_dash(spelled):
    assert T.same_beamline(spelled, "BL152") and T.same_beamline(spelled, "BL15-2")


@pytest.mark.parametrize("a,b", [("BL11-2", "BL1-12"), ("15-2", "1-52"), ("BL6-2c", "BL6-2b"), ("62b", "BL152")])
def test_the_beamline_hint_keeps_the_separator_and_the_letter(a, b):
    assert not T.same_beamline(a, b)


def test_a_beamline_with_no_separator_matches_either_split():
    assert T.same_beamline("BL112", "BL11-2") and T.same_beamline("BL112", "BL1-12")
    assert T.beamline_form("BL15–2") == "15-2" and T.beamline_form("SSRL BL152") == "152"


@pytest.mark.parametrize("spelled", ["SSRL BL15-2", "BL15-2 (SSRL)", "SSRL 15-2", "BL15–2"])
def test_a_facility_spelling_of_the_folders_beamline_is_adopted(plain, tmp_path, capsys, spelled):
    out = tmp_path / "adopted.xlsx"
    a = _adopt_json(plain, _sheet(plain, tmp_path / THEIRS, beamline=spelled), out, capsys)
    assert not [d for d in a["disagreements"] if d["header"] == "Beamline/Instrument"]
    assert openpyxl.load_workbook(out)["Campaign Info"]["B5"].value == spelled


def test_another_split_of_the_same_digits_is_a_disagreement(plain, tmp_path, capsys):
    a = _adopt_json(plain, _sheet(plain, tmp_path / THEIRS, beamline="BL1-52"), tmp_path / "a.xlsx", capsys)
    bl = _d(a, "Beamline/Instrument", kind="hint")
    assert any("'BL15-2'" in r and "compared as '15-2', yours '1-52'" in r for r in bl["readings"])


# --- HIGH-1: a beamtime-wide answer does not reach a group that keeps its own --------------------------------


def _attribution(bdir):
    """``(question id, the --field flag it needs)`` for the question that asks the producing group once for the beamtime."""
    b = M.load(bdir)
    q = Q.question(Q.consolidate(b), _attribution_question(bdir))
    return q["question_id"], (["--field", GROUP] if q["kind"] == "optional_values" else [])


def _only(capsys, bdir, qid, gid, value, *field):
    code, out, err = _decide(capsys, bdir, "--question", qid, *field, "--only", f"group:{gid}", "--decision", "correct",
                             "--value", value)
    assert code == 0, err
    return json.loads(out)


def test_a_beamtime_answer_names_the_group_keeping_its_own_answer_as_not_reached(plain, capsys):
    qid, gid = _technique_question(plain), _groups(plain)[0]
    _only(capsys, plain, qid, gid, "XAS")
    planned = M.decide_question(plain, question_id=qid, decision="correct", value="HERFD-XAS", by=WHO, orcid=None,
                                dry_run=True)
    assert [k["group_id"] for k in planned["reach"]["not_reached"]] == [gid]  # said before anything is written
    code, out, err = _decide(capsys, plain, "--question", qid, "--decision", "correct", "--value", "HERFD-XAS")
    assert code == 0, err
    res = json.loads(out)
    said = res["statement"]
    assert "It reaches every record group" not in said
    assert "It reaches 8 of the beamtime's 9 record groups (8 runs)." in said
    assert (f"NOT reached: group:{gid} keeps its own answer 'XAS', so it fills nothing for system.technique until the two "
            f"agree; answer `--only group:{gid}` to change it") in said
    assert f"--question {qid} --only group:{gid} --decision correct --value <value>" in said
    [kept] = res["reach"]["not_reached"]
    assert kept["group_id"] == gid and kept["own_answers"] == ["XAS"] and res["reach"]["groups_reached"] == 8
    composed = _composed(plain, TECH)
    assert composed[gid][0] is None and sum(1 for v, _ in composed.values() if v == "HERFD-XAS") == 8
    # the group's own answer changed to the beamtime's: the two agree, and the next beamtime-wide reply reaches all
    res = _only(capsys, plain, qid, gid, "HERFD-XAS")
    assert f"It agrees with the beamtime-wide answer to question {qid}." in res["statement"]
    code, out, _ = _decide(capsys, plain, "--question", qid, "--decision", "correct", "--value", "HERFD-XAS")
    assert "It reaches every record group of the beamtime: 9 groups, 9 runs." in json.loads(out)["statement"]


def test_a_group_answer_against_an_answered_beamtime_question_says_the_two_values_plainly(plain, capsys):
    qid, gid = _technique_question(plain), _groups(plain)[0]
    assert _decide(capsys, plain, "--question", qid, "--decision", "correct", "--value", "HERFD-XAS")[0] == 0
    res = _only(capsys, plain, qid, gid, "XAS")
    assert (f"The beamtime-wide answer to question {qid} differs from this one, so this group's record has two values "
            "for technique and fills nothing for it until the two agree.") in res["statement"]


def _campaign_sheet(bdir, tmp_path, label, value):
    xlsx = tmp_path / "t.xlsx"
    assert cli.main(["bundle", "template", "export", str(bdir), "--out", str(xlsx)]) == 0
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Campaign Info"]
    row = next(r for r in range(1, ws.max_row + 1) if label in str(ws[f"A{r}"].value or ""))
    ws[f"B{row}"] = value
    wb.save(tmp_path / "typed.xlsx")
    return tmp_path / "typed.xlsx"


def test_a_campaign_value_names_the_group_keeping_its_own_answer_in_the_preview(plain, tmp_path, capsys):
    qid, field = _attribution(plain)
    gid = _groups(plain)[0]
    _only(capsys, plain, qid, gid, "Lab A", *field)
    typed = _campaign_sheet(plain, tmp_path, "Your group", "Lab B")
    code, rep, err = _import(plain, typed, capsys)
    assert code == 0, err
    [nr] = rep["not_reached"]
    assert (nr["group_id"], nr["own_answers"], nr["value"], nr["target_field_path"]) == (gid, ["Lab A"], "Lab B", GROUP)
    assert f"--question {qid} --only group:{gid}" in nr["command"]
    capsys.readouterr()
    assert cli.main(["bundle", "template", "import", str(plain), str(typed), "--by", WHO]) == 0
    text = capsys.readouterr().out
    assert "Not reached by a beamtime-wide answer" in text and f"group:{gid} keeps its own answer 'Lab A'" in text
    assert "1 group(s) not reached by a beamtime-wide answer" in text
    # --supersede rewrites the group's own answer too: then nothing is held apart
    code, rep, err = _import(plain, typed, capsys, "--supersede")
    assert code == 0 and rep["not_reached"] == []


# --- MED-2: the minimum's command for a group holding a beamtime target and its own answer ----------------------


def test_the_minimum_answers_a_group_holding_two_values_through_its_question_never_the_beamtime_target(plain, capsys):
    qid, gid = _technique_question(plain), _groups(plain)[0]
    _only(capsys, plain, qid, gid, "XAS")
    assert _decide(capsys, plain, "--question", qid, "--decision", "correct", "--value", "HERFD-XAS")[0] == 0
    b = M.load(plain)
    reports: dict = {}
    payload = Q.consolidate(b, compositions_out=reports)
    minimum = Q.minimum_answers_for_one_group(payload, b, scope=f"group:{gid}", bundle_dir=plain, reports=reports)
    [item] = [u for u in minimum["unasked_required"] if u["path"] == TECH]
    beamtime = next(c["candidate_id"] for c in b["candidates"] if c["run_id"] is None and c["target_field_path"] == TECH)
    assert f"--target {beamtime}" not in item["command"]
    assert item["command"].startswith(f"isaac bundle decide <bundle_dir> --question {qid} --only group:{gid} --decision correct")
    assert "without --only, which reaches every group" in item["command"] and item["question_id"] == qid
    assert "the confirmed values for it differ inside this group" in item["reason"]


# --- MED-3: a slot belongs to its run -----------------------------------------------------------------------


def _merged(bdir, a, b):
    """The record group a merge of groups ``a`` and ``b`` makes (their runs together), composed for ``path``."""
    bundle = M.load(bdir)
    by_id = {g["group_id"]: g for g in Q.live_groups(bundle)}
    group = {**by_id[b], "member_run_keys": [*by_id[a]["member_run_keys"], *by_id[b]["member_run_keys"]]}
    draft, rep = R.compose_group(bundle, group, bdir)
    return (draft["fields"].get(GROUP) or {}).get("value"), rep


def test_a_merged_group_fills_nothing_from_one_runs_answer_slot(plain, capsys):
    qid, field = _attribution(plain)
    g0, g1 = _groups(plain)[:2]
    _only(capsys, plain, qid, g0, "Lab A", *field)
    assert _composed(plain, GROUP)[g0][0] == "Lab A"  # its own group still takes it
    value, rep = _merged(plain, g0, g1)
    assert value is None
    [missing] = [m for m in rep["missing_required"] if m["path"] == GROUP]
    assert "belongs to those runs" in missing["reason"] and f"--question {qid} --only group:" in missing["reason"]
    slot = next(m["value_id"] for m in M.load(plain)["missing_values"] if B.is_answer_slot(m) and m["target_field_path"] == GROUP)
    assert any(x.get("value_id") == slot and "1 of this record group's 2 runs only" in x["reason"] for x in rep["not_composed"])
    # every run answered alike: the merged group fills it
    _only(capsys, plain, qid, g1, "Lab A", *field)
    assert _merged(plain, g0, g1)[0] == "Lab A"


def test_a_merged_group_fills_a_slot_value_the_beamtime_answer_gives_the_other_runs(plain, capsys):
    qid, field = _attribution(plain)
    g0, g1 = _groups(plain)[:2]
    _only(capsys, plain, qid, g0, "Lab A", *field)
    code, _, err = _decide(capsys, plain, "--question", qid, *field, "--decision", "correct", "--value", "Lab A")
    assert code == 0, err
    assert _merged(plain, g0, g1)[0] == "Lab A"


# --- LOW-8: dead slots ---------------------------------------------------------------------------------------


def test_a_slot_no_decision_names_is_dropped_on_any_save_and_one_a_decision_names_is_kept(plain, capsys):
    qid, gid = _technique_question(plain), _groups(plain)[0]
    _only(capsys, plain, qid, gid, "XAS")
    b = M.load(plain)
    kept = [m["value_id"] for m in b["missing_values"] if B.is_answer_slot(m)]
    _ids, added = M.run_answer_targets(b, TECH, ["run-002"], question_id=qid, how="a synthetic test")
    assert added and B.is_answer_slot(next(m for m in b["missing_values"] if m["value_id"] == added[0]))
    assert B.answer_slot_question(next(m for m in b["missing_values"] if m["value_id"] == added[0])) == qid
    M.save_bundle(plain, b)
    after = [m["value_id"] for m in M.load(plain)["missing_values"] if B.is_answer_slot(m)]
    assert after == kept


def test_a_deferred_slot_is_not_asked_once_the_beamtime_question_is_answered(plain, capsys):
    qid, gid = _technique_question(plain), _groups(plain)[0]
    _only(capsys, plain, qid, gid, "XAS")
    slot = next(m["value_id"] for m in M.load(plain)["missing_values"] if B.is_answer_slot(m))
    assert _decide(capsys, plain, "--target", slot, "--decision", "defer")[0] == 0
    before = {q["question_id"] for q in Q.consolidate(M.load(plain))["questions"]}
    asks_slot = [q["question_id"] for q in Q.consolidate(M.load(plain))["questions"] if slot in json.dumps(q["targets"])]
    assert asks_slot  # deferred, it is still a question while nothing else answers its run
    assert _decide(capsys, plain, "--question", qid, "--decision", "correct", "--value", "HERFD-XAS")[0] == 0
    b = M.load(plain)
    assert B.validate_bundle(b) == []
    open_now = Q.consolidate(b)["questions"]
    assert not [q for q in open_now if slot in json.dumps(q["targets"])]
    assert {q["question_id"] for q in open_now} <= before  # no question id changed
    answered = Q.consolidate(b, include_answered=True, explain_counter=False)["questions"]
    assert [q["question_id"] for q in answered if slot in json.dumps(q["targets"])] == asks_slot
    assert slot in {m["value_id"] for m in b["missing_values"]}  # the decision log names it: kept
    assert _composed(plain, TECH)[gid][0] == "HERFD-XAS"


# --- MED-4 / LOW-6: the page -----------------------------------------------------------------------------------


LOCAL = dict(base_url="http://127.0.0.1:8000", client=("127.0.0.1", 50123), headers={LB.LOCAL_HEADER: "1"})


def _client(root, monkeypatch, tmp_path):
    monkeypatch.setenv(LB.LOCAL_BUNDLES_ENV, str(root))
    monkeypatch.setenv("ISAAC_UI_WORKSPACE", str(tmp_path / "ws"))
    monkeypatch.delenv("ISAAC_BASE_PATH", raising=False)
    LB._check_memo.clear()
    return TestClient(create_app(), **LOCAL)


def _exposed(bdir, command):
    return {e["target"]["id"] for e in M.load(bdir).get("exposures") or () if e["command"] == command}


def test_a_value_the_nearest_records_warnings_state_is_an_exposure_on_the_progress_page(plain, tmp_path, monkeypatch):
    client = _client(plain.parent, monkeypatch, tmp_path)
    body = client.get(f"/api/local-bundles/{plain.name}/progress").json()
    minimum = body["progress"]["minimum_for_one_group"]
    stated = minimum["values_shown_from"]
    assert stated and any("read as 0.1 V" in w for w in minimum["warnings"])
    assert set(stated) <= _exposed(plain, "bundle progress") and body["progress"]["exposures_recorded"] == len(stated)
    assert body["bundle"]["bundle_version"] == M.bundle_version(plain)  # the version after the exposures
    again = client.get(f"/api/local-bundles/{plain.name}/progress").json()
    assert again["progress"]["exposures_recorded"] == 0
    # the questions page shows the same minimum: the same values are its exposures too
    page = client.get(f"/api/local-bundles/{plain.name}/questions").json()
    assert set(page["minimum_for_one_group"]["values_shown_from"]) <= _exposed(plain, "bundle questions")


def test_the_command_lines_bundle_questions_records_the_minimums_stated_values_too(plain, capsys):
    capsys.readouterr()
    assert cli.main(["bundle", "questions", str(plain)]) == 0
    minimum = json.loads(capsys.readouterr().out)["minimum_for_one_group"]
    assert minimum["values_shown_from"] and set(minimum["values_shown_from"]) <= _exposed(plain, "bundle questions")


def test_a_web_dry_run_of_one_group_of_a_beamtime_question_previews_that_group(plain, tmp_path, monkeypatch):
    client = _client(plain.parent, monkeypatch, tmp_path)
    qid, gid = _technique_question(plain), _groups(plain)[0]
    version = M.bundle_version(plain)
    n = len(M.load(plain)["missing_values"])
    res = client.post(f"/api/local-bundles/{plain.name}/questions/{qid}/answer",
                      json={"decision": "correct", "value": "XAS", "by": WHO, "only": f"group:{gid}", "dry_run": True,
                            "expected_version": version}).json()
    [row] = res["by_scope"]
    assert row["scope"] == f"group:{gid}" and row["answer_slots"] == 1 and len(row["target_ids"]) == 1
    assert res["scopes_total"] == 1 and f"group:{gid} only" in res["would_say"]
    assert M.bundle_version(plain) == version and len(M.load(plain)["missing_values"]) == n


def test_a_merged_group_answered_for_itself_records_every_run_and_then_fills(plain, capsys):
    qid, field = _attribution(plain)
    g0, g1 = _groups(plain)[:2]
    _only(capsys, plain, qid, g0, "Lab A", *field)
    # the groups as a merge leaves them: one live group holding both groups' runs (synthetic, as a regroup makes it)
    b = M.load(plain)
    by_id = {g["group_id"]: g for g in b["record_groups"]}
    merged = {**by_id[g1], "group_id": "rg-zz0000000000merged",
              "member_run_keys": [*by_id[g0]["member_run_keys"], *by_id[g1]["member_run_keys"]]}
    b["record_groups"] = [merged, *[g for g in b["record_groups"] if g["group_id"] not in (g0, g1)]]
    M.save_bundle(plain, b)
    assert _composed(plain, GROUP)[merged["group_id"]][0] is None
    slots_before = [m["value_id"] for m in M.load(plain)["missing_values"] if B.is_answer_slot(m)]
    res = _only(capsys, plain, qid, merged["group_id"], "Lab A", *field)
    assert res["reach"]["runs"] == 2 and res["reach"]["slots_added"] == 1  # g0's run keeps its slot, g1's run gets one
    assert len([m for m in M.load(plain)["missing_values"] if B.is_answer_slot(m)]) == len(slots_before) + 1
    assert _composed(plain, GROUP)[merged["group_id"]][0] == "Lab A"
