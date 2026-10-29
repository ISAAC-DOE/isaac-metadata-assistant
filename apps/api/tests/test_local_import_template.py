"""The ISAAC metadata template round trip (``isaac bundle template export|import``). Synthetic only."""

from __future__ import annotations

import hashlib
import json
import shutil

import openpyxl
import pytest

from isaac_api.local_import import bundle as B
from isaac_api.local_import import cli
from isaac_api.local_import import mutations as M
from isaac_api.local_import import template as T

from tests_support_records import BY, pristine_bundle  # noqa: F401

# ISAAC's template header text (its own format), in its order.
TEMPLATE_FILE_LIST_HEADERS = ["File Name", "Config ID", "Sample Name", "Sample Formula", "Sample Form", "Environment",
                              "Reaction", "Cell Type", "Flow Mode", "Electrolyte", "Potential (V)", "Potential Scale", "Gas Feed"]
# ISAAC's two additions come AFTER the template's own columns
FILE_LIST_HEADERS = [*TEMPLATE_FILE_LIST_HEADERS, "Electrolyte Concentration (M)", "pH"]
CONFIG_HEADERS = ["Config ID", "Technique", "Mode", "Primary Signal Channel", "Normalization Channel", "Detectors",
                  "Energy Calibration", "Beamline Setup", "Notes"]
CAMPAIGN_LABELS = ["Facility Name", "Organization", "Beamline/Instrument", "Start Date", "Proposal ID"]
WHO = "Synthetic Scientist"


@pytest.fixture(scope="module")
def grouped_base(pristine_bundle, tmp_path_factory):
    bdir = tmp_path_factory.mktemp("tgrouped") / "b"
    shutil.copytree(pristine_bundle, bdir)
    M.group(bdir)
    return bdir


@pytest.fixture()
def grouped(grouped_base, tmp_path):
    bdir = tmp_path / "b"
    shutil.copytree(grouped_base, bdir)
    return bdir


def _export(bdir, tmp_path, name="t.xlsx", extra=()):
    out = tmp_path / name
    assert cli.main(["bundle", "template", "export", str(bdir), "--out", str(out), *extra]) == 0
    return out


def _import(bdir, xlsx, capsys, *extra):
    capsys.readouterr()
    code = cli.main(["bundle", "template", "import", str(bdir), str(xlsx), "--by", WHO, "--json", *extra])
    captured = capsys.readouterr()
    return code, (json.loads(captured.out) if code == 0 else None), captured.err


def _digest(bdir):
    return hashlib.sha256((bdir / "bundle.json").read_bytes()).hexdigest()


def _row_of(xlsx, run_id):
    wb = openpyxl.load_workbook(xlsx)
    for values in wb["_isaac"].iter_rows(values_only=True):
        if values[0] == "File List" and values[2] == run_id:
            return int(values[1])
    raise AssertionError(run_id)


def _fill(xlsx, out, cells):
    wb = openpyxl.load_workbook(xlsx)
    for (sheet, ref), value in cells.items():
        wb[sheet][ref] = value
    wb.save(out)
    return out


def _miss(b, path, run_id):
    return next(m for m in b["missing_values"] if m["target_field_path"] == path and m["run_id"] == run_id)


# --- export -------------------------------------------------------------------------------------


def test_export_has_the_template_sheets_headers_and_one_row_per_live_group_run(grouped, tmp_path):
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    assert wb.sheetnames == ["Campaign Info", "Configurations", "File List", "Vocabulary", "ISAAC Suggestions", "_isaac"]
    assert wb["_isaac"].sheet_state == "hidden" and wb["ISAAC Suggestions"].sheet_state == "visible"
    assert [c.value for c in wb["File List"][1]] == FILE_LIST_HEADERS
    assert [c.value for c in wb["ISAAC Suggestions"][1]] == FILE_LIST_HEADERS
    assert [c.value for c in wb["Configurations"][1]] == CONFIG_HEADERS
    assert [wb["Campaign Info"][f"A{r}"].value for r in range(3, 8)] == CAMPAIGN_LABELS

    b = M.load(grouped)
    rows = T.live_rows(b)
    assert len(rows) == 9  # every run of the synthetic corpus is in a proposed group
    fl = wb["File List"]
    assert [fl.cell(row=i, column=1).value for i in range(2, 2 + len(rows))] == [r["file_name"] for r in rows]
    assert fl.max_row == 1 + len(rows)

    # EVERY answer cell is empty: Campaign Info, Configurations (header only), File List B..M
    assert all(wb["Campaign Info"][f"B{r}"].value is None for r in range(3, 8))
    assert all(c.value is None for row in wb["Configurations"].iter_rows(min_row=2) for c in row)
    assert all(c.value is None for row in fl.iter_rows(min_row=2, min_col=2) for c in row)

    # the hidden marker: format, version, bundle, time, row -> run, column -> path
    meta = {v[0]: v[1] for v in wb["_isaac"].iter_rows(values_only=True) if v[0] and v[1] is not None}
    assert meta["format"] == "isaac.metadata_template" and meta["format_version"] == "1.1.0"
    assert meta["bundle_id"] == b["bundle_id"] and meta["exported_utc"].endswith("Z")
    mapped_rows = [v for v in wb["_isaac"].iter_rows(values_only=True) if v[0] == "File List" and isinstance(v[1], int)]
    assert [(v[1], v[2], v[3], v[4]) for v in mapped_rows] == [(i, r["run_id"], r["run_key"], r["group_id"]) for i, r in enumerate(rows, 2)]
    col_paths = {(v[0], v[2]): v[3] for v in wb["_isaac"].iter_rows(values_only=True) if v[0] in ("Campaign Info", "File List", "Configurations") and isinstance(v[1], str)}
    assert col_paths[("Campaign Info", "Facility Name")] == "system.facility.facility_name"
    assert col_paths[("File List", "Environment")] == "context.environment"
    assert col_paths[("Configurations", "Technique")] == "system.technique"
    assert col_paths[("File List", "Gas Feed")] is None


def test_export_vocabulary_comes_from_the_schema_with_dropdowns(grouped, tmp_path):
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    voc = wb["Vocabulary"]
    headers = [c.value for c in voc[1]]
    env_col = headers.index("Environment") + 1
    env = [voc.cell(row=r, column=env_col).value for r in range(2, voc.max_row + 1)]
    assert [v for v in env if v] == list(B._schema_node("context.environment")["enum"])
    dvs = {str(dv.sqref): dv for dv in wb["File List"].data_validations.dataValidation}
    assert dvs["F2:F10"].type == "list" and "Vocabulary" in dvs["F2:F10"].formula1
    assert dvs["K2:K10"].type == "decimal"  # Potential (V) is numeric


def test_export_suggestions_cite_the_source_and_record_an_exposure(grouped, tmp_path):
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    sg = wb["ISAAC Suggestions"]
    sample_col = FILE_LIST_HEADERS.index("Sample Name") + 1
    text = sg.cell(row=2, column=sample_col).value
    assert "literal 'ZZ1'" in text and " @ filename token" in text
    exposures = M.load(grouped).get("exposures") or []
    assert any(e["command"] == "bundle template export" for e in exposures)


def test_no_suggestions_records_no_exposure(grouped, tmp_path):
    before = len(M.load(grouped).get("exposures") or [])
    xlsx = _export(grouped, tmp_path, extra=("--no-suggestions",))
    assert len(M.load(grouped).get("exposures") or []) == before
    sg = openpyxl.load_workbook(xlsx)["ISAAC Suggestions"]
    assert all(c.value is None for row in sg.iter_rows(min_row=2, max_row=10, min_col=2) for c in row)


def test_export_lists_only_runs_of_live_groups_and_refuses_an_ungrouped_bundle(grouped, pristine_bundle, tmp_path):
    b = M.load(grouped)
    g = b["record_groups"][0]
    M.decide(grouped, target_ids=[g["group_id"]], decision="reject", value=None, **BY)
    xlsx = _export(grouped, tmp_path)
    names = [c.value for c in openpyxl.load_workbook(xlsx)["File List"]["A"][1:]]
    assert len(names) == 9 - len(g["member_run_keys"])
    ungrouped = tmp_path / "u"
    shutil.copytree(pristine_bundle, ungrouped)
    assert cli.main(["bundle", "template", "export", str(ungrouped), "--out", str(tmp_path / "u.xlsx")]) == 2
    assert not (tmp_path / "u.xlsx").exists()
    # an existing file is not overwritten without --force
    assert cli.main(["bundle", "template", "export", str(grouped), "--out", str(xlsx)]) == 4


# --- round trip ---------------------------------------------------------------------------------


def test_round_trip_preview_writes_nothing_then_apply_records_and_the_bundle_stays_valid(grouped, tmp_path, capsys):
    tech_run = next(c["run_id"] for c in M.load(grouped)["candidates"] if c["run_id"] and c["target_field_path"] == "system.technique")
    xlsx = _export(grouped, tmp_path)
    r1 = _row_of(xlsx, "run-001")
    filled = _fill(xlsx, tmp_path / "filled.xlsx", {
        ("Campaign Info", "B3"): "Synthetic Light Source",
        ("File List", f"C{r1}"): "ZZ1 synthetic",
        ("File List", f"F{r1}"): "ex_situ",
        ("File List", f"K{r1}"): 0.1,
        ("Configurations", "A2"): "Config_S",
        ("Configurations", "B2"): "HERFD-XAS",
        ("File List", f"B{_row_of(xlsx, tech_run)}"): "Config_S",
    })
    before = _digest(grouped)
    code, rep, _ = _import(grouped, filled, capsys)
    assert code == 0 and rep["applied"] is False and rep["would_record"] > 0
    assert _digest(grouped) == before  # PREVIEW: nothing written
    assert not M.load(grouped)["scientist_decisions"]

    code, rep, _ = _import(grouped, filled, capsys, "--apply")
    assert code == 0 and rep["applied"] and rep["recorded"] == rep["would_record"]
    b = M.load(grouped)
    env = _miss(b, "context.environment", "run-001")
    assert env["value_state"] == B.STATE_SCIENTIST_CONFIRMED
    pot = next(c for c in b["candidates"] if c["run_id"] == "run-001" and c["target_field_path"] == "context.electrochemistry.potential_setpoint_V")
    assert pot["value_state"] == B.STATE_SCIENTIST_CONFIRMED
    d_pot = next(d for d in b["scientist_decisions"] if d["target"]["id"] == pot["candidate_id"])
    assert d_pot["decision"] == "confirm" and d_pot["value"] == 0.1  # cell equals the extracted value
    name = next(c for c in b["candidates"] if c["run_id"] == "run-001" and c["target_field_path"] == "sample.material.name")
    assert next(d for d in b["scientist_decisions"] if d["target"]["id"] == name["candidate_id"])["decision"] == "correct"
    tech = next(c for c in b["candidates"] if c["run_id"] == tech_run and c["target_field_path"] == "system.technique")
    d_tech = next(d for d in b["scientist_decisions"] if d["target"]["id"] == tech["candidate_id"])
    assert d_tech["value"] == "HERFD-XAS" and "'Configurations'!B2" in d_tech["note"] and "'File List'!B" in d_tech["note"]
    assert all(d["decided_by"]["basis"] == "self_reported_local_user" for d in b["scientist_decisions"])
    assert cli.main(["bundle", "check", str(grouped)]) == 0
    # replaying the same workbook records nothing new
    n = len(b["scientist_decisions"])
    code, rep, _ = _import(grouped, filled, capsys, "--apply")
    assert code == 0 and rep["recorded"] == 0 and len(M.load(grouped)["scientist_decisions"]) == n


def test_each_note_cites_the_cell(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    r = _row_of(xlsx, "run-001")
    filled = _fill(xlsx, tmp_path / "cite.xlsx", {("File List", f"F{r}"): "ex_situ"})
    _import(grouped, filled, capsys, "--apply")
    d = M.load(grouped)["scientist_decisions"][-1]
    assert d["note"] == f"template cite.xlsx 'File List'!F{r}"


def test_a_beamtime_cell_fills_every_target_of_that_path(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    filled = _fill(xlsx, tmp_path / "beam.xlsx", {("Campaign Info", "B5"): "ZZ-1"})
    code, rep, _ = _import(grouped, filled, capsys, "--apply")
    b = M.load(grouped)
    targets = [m for m in b["missing_values"] if m["target_field_path"] == "system.facility.beamline"]
    assert len(targets) == 9 and rep["recorded"] == 9
    assert all(m["value_state"] == B.STATE_SCIENTIST_CONFIRMED for m in targets)
    decided = {d["target"]["id"]: d for d in b["scientist_decisions"]}
    assert all(decided[m["value_id"]]["value"] == "ZZ-1" and decided[m["value_id"]]["note"].endswith("'Campaign Info'!B5")
               for m in targets)
    assert len({decided[m["value_id"]]["batch_id"] for m in targets}) == 1


# --- refusals -----------------------------------------------------------------------------------


def test_a_workbook_without_the_isaac_marker_is_refused(grouped, tmp_path, capsys):
    raw = openpyxl.Workbook()
    ws = raw.active
    ws.title = "Campaign Info"
    ws["A3"], ws["B3"] = "Facility Name", "Example Facility"  # an example value, never an answer
    fl = raw.create_sheet("File List")
    fl.append(FILE_LIST_HEADERS)
    raw.save(tmp_path / "raw.xlsx")
    before = _digest(grouped)
    code, _, err = _import(grouped, tmp_path / "raw.xlsx", capsys, "--apply")
    assert code == 3 and "template export" in err and "_isaac" in err
    assert _digest(grouped) == before


def test_a_workbook_for_another_bundle_is_refused(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    for row in wb["_isaac"].iter_rows():
        if row[0].value == "bundle_id":
            row[1].value = "01ANOTHERBUNDLE0000000000"
    wb["Campaign Info"]["B3"] = "Synthetic Light Source"
    wb.save(tmp_path / "other.xlsx")
    before = _digest(grouped)
    code, _, err = _import(grouped, tmp_path / "other.xlsx", capsys, "--apply")
    assert code == 3 and "not this one" in err and "template export" in err
    assert _digest(grouped) == before


def test_an_enum_cell_must_hold_an_allowed_value_exactly(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    r = _row_of(xlsx, "run-001")
    filled = _fill(xlsx, tmp_path / "enum.xlsx", {("File List", f"F{r}"): "Ex_Situ", ("File List", f"H{r}"): "a beaker"})
    code, rep, _ = _import(grouped, filled, capsys, "--apply")
    assert code == 0 and rep["recorded"] == 0
    cells = {i["cell"]: i for i in rep["invalid_cells"]}
    assert set(cells) == {f"'File List'!F{r}", f"'File List'!H{r}"}
    assert "must be written exactly" in cells[f"'File List'!F{r}"]["reason"]
    assert not M.load(grouped)["scientist_decisions"]


def test_a_numeric_cell_must_parse_as_a_number(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    r = _row_of(xlsx, "run-001")
    filled = _fill(xlsx, tmp_path / "num.xlsx", {("File List", f"K{r}"): "about 0.1 V"})
    code, rep, _ = _import(grouped, filled, capsys, "--apply")
    assert rep["recorded"] == 0 and rep["invalid_cells"][0]["cell"] == f"'File List'!K{r}"
    assert "not a number" in rep["invalid_cells"][0]["reason"]
    assert not M.load(grouped)["scientist_decisions"]


def test_a_blank_workbook_records_nothing(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    before = _digest(grouped)
    code, rep, _ = _import(grouped, xlsx, capsys, "--apply")
    assert code == 0 and rep["cells_read"] == 0 and rep["recorded"] == 0 and rep["decisions"] == []
    assert _digest(grouped) == before  # not even rewritten


def test_columns_with_no_official_field_are_reported_with_their_cells(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    r = _row_of(xlsx, "run-001")
    filled = _fill(xlsx, tmp_path / "unmapped.xlsx", {
        ("Campaign Info", "B6"): "2099-01-01", ("File List", f"M{r}"): "synthetic gas",
        ("File List", f"I{r}"): "stagnant", ("Configurations", "A2"): "Config_S", ("Configurations", "F2"): "Synthetic detector",
    })
    code, rep, _ = _import(grouped, filled, capsys, "--apply")
    by_header = {u["header"]: u for u in rep["unmapped_columns"]}
    assert by_header["Start Date"]["reason"].startswith("recorded nowhere: no official field")
    assert by_header["Gas Feed"]["cells"] == [{"cell": f"'File List'!M{r}", "value": "synthetic gas"}]
    assert by_header["Detectors"]["cells"][0]["cell"] == "'Configurations'!F2"
    # an official field this bundle holds no target for is reported too, not dropped
    assert "context.transport.flow_mode" in by_header["Flow Mode"]["reason"]
    assert rep["recorded"] == 0


def test_a_moved_row_is_refused_rather_than_guessed(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    fl = wb["File List"]
    fl["A2"], fl["A3"] = fl["A3"].value, fl["A2"].value  # rows swapped
    fl["F2"] = "ex_situ"
    wb.save(tmp_path / "moved.xlsx")
    code, rep, _ = _import(grouped, tmp_path / "moved.xlsx", capsys, "--apply")
    assert rep["recorded"] == 0 and rep["rows_refused"][0]["row"] == 2


# --- clashes ------------------------------------------------------------------------------------


def test_a_clash_with_an_earlier_decision_is_skipped_unless_supersede(grouped, tmp_path, capsys):
    b = M.load(grouped)
    env = _miss(b, "context.environment", "run-001")
    M.decide(grouped, target_ids=[env["value_id"]], decision="correct", value="in_situ", **BY)
    xlsx = _export(grouped, tmp_path)
    r = _row_of(xlsx, "run-001")
    filled = _fill(xlsx, tmp_path / "clash.xlsx", {("File List", f"F{r}"): "ex_situ"})
    code, rep, _ = _import(grouped, filled, capsys, "--apply")
    assert code == 0 and rep["recorded"] == 0 and len(rep["clashes"]) == 1
    assert rep["clashes"][0]["existing"]["value"] == "in_situ"
    assert B._decisions_by_target(M.load(grouped)["scientist_decisions"])[("missing_value", env["value_id"])]["value"] == "in_situ"

    code, rep, _ = _import(grouped, filled, capsys, "--apply", "--supersede")
    assert code == 0 and rep["recorded"] == 1 and not rep["clashes"]
    latest = B._decisions_by_target(M.load(grouped)["scientist_decisions"])[("missing_value", env["value_id"])]
    assert latest["value"] == "ex_situ" and latest["note"].endswith(f"'File List'!F{r}")
    assert len([d for d in M.load(grouped)["scientist_decisions"] if d["target"]["id"] == env["value_id"]]) == 2  # kept, not overwritten
    assert cli.main(["bundle", "check", str(grouped)]) == 0


# --- review fixes: M3 (unique row identity), M4 (electrolyte split), L10 (campaign reach) ---------


def _col(header):
    from openpyxl.utils import get_column_letter

    return get_column_letter(FILE_LIST_HEADERS.index(header) + 1)


def test_m3_runs_whose_files_share_a_basename_get_distinct_row_names(grouped):
    b = M.load(grouped)
    src = {s["source_id"]: s for s in b["sources"]}
    r1, r2 = b["runs"][0], b["runs"][1]
    src[r1["acquisition_source_id"]]["relative_path"] = "SampleA/scan_001"
    src[r2["acquisition_source_id"]]["relative_path"] = "SampleB/scan_001"
    names = {r["run_id"]: r["file_name"] for r in T.live_rows(b)}
    assert names[r1["run_id"]] == "SampleA/scan_001" and names[r2["run_id"]] == "SampleB/scan_001"
    assert len(set(names.values())) == len(names)


def test_m3_names_that_still_repeat_refuse_the_export(grouped):
    b = M.load(grouped)
    r1, r2 = b["runs"][0], b["runs"][1]
    r2["acquisition_source_id"] = r1["acquisition_source_id"]  # one file, two acquisitions ...
    assert len({r["file_name"] for r in T.live_rows(b)}) == len(T.live_rows(b))  # ... told apart by label
    r2["label"] = r1["label"]  # ... and now not at all
    with pytest.raises(T.TemplateRefused) as exc:
        T.live_rows(b)
    assert exc.value.code == 3 and "could not tell which row is which" in str(exc.value)


def test_m3_the_file_name_is_the_path_and_must_match_exactly(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    b = M.load(grouped)
    src = {s["source_id"]: s["relative_path"] for s in b["sources"]}
    fl = openpyxl.load_workbook(xlsx)["File List"]
    assert fl["A2"].value == src[b["runs"][0]["acquisition_source_id"]]
    wb = openpyxl.load_workbook(xlsx)
    wb["File List"]["A2"] = wb["File List"]["A2"].value + " "  # a trailing space is not the same name
    wb["File List"][f"{_col('Environment')}2"] = "ex_situ"
    wb.save(tmp_path / "space.xlsx")
    code, rep, _ = _import(grouped, tmp_path / "space.xlsx", capsys, "--apply")
    assert rep["recorded"] == 0 and rep["rows_refused"][0]["row"] == 2


def test_m3_a_mapping_that_names_one_file_on_two_rows_is_refused(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    rows = [r for r in wb["_isaac"].iter_rows() if r[0].value == "File List" and isinstance(r[1].value, int)]
    rows[1][5].value = rows[0][5].value
    wb.save(tmp_path / "dup.xlsx")
    code, _, err = _import(grouped, tmp_path / "dup.xlsx", capsys, "--apply")
    assert code == 3 and "cannot be told apart" in err


def test_m3_a_reordered_row_says_where_its_name_came_from(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    fl = wb["File List"]
    fl["A2"], fl["A3"] = fl["A3"].value, fl["A2"].value
    fl[f"{_col('Environment')}2"] = "ex_situ"
    wb.save(tmp_path / "swap.xlsx")
    code, rep, _ = _import(grouped, tmp_path / "swap.xlsx", capsys, "--apply")
    assert rep["recorded"] == 0 and "exported on row 3" in rep["rows_refused"][0]["reason"]


def test_m4_electrolyte_name_concentration_and_ph_are_separate_columns(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    fl = wb["File List"]
    assert [c.value for c in fl[1]][: len(TEMPLATE_FILE_LIST_HEADERS)] == TEMPLATE_FILE_LIST_HEADERS
    assert "NAME only" in fl[f"{_col('Electrolyte')}1"].comment.text
    assert "Added by ISAAC" in fl[f"{_col('pH')}1"].comment.text
    assert any("Electrolyte is the NAME only" in str(c.value) for c in wb["Campaign Info"]["A"])
    paths = {v[2]: v[3] for v in wb["_isaac"].iter_rows(values_only=True) if v[0] == "File List" and isinstance(v[1], str)}
    assert paths["Electrolyte Concentration (M)"] == "context.electrochemistry.electrolyte.concentration_M"
    assert paths["pH"] == "context.electrochemistry.pH"
    r = _row_of(xlsx, "run-001")
    filled = _fill(xlsx, tmp_path / "el.xlsx", {
        ("File List", f"{_col('Electrolyte Concentration (M)')}{r}"): 0.1,
        ("File List", f"{_col('pH')}{r}"): "13",
        ("File List", f"{_col('Electrolyte Concentration (M)')}{_row_of(xlsx, 'run-002')}"): "0.1 M KOH",
    })
    code, rep, _ = _import(grouped, filled, capsys, "--apply")
    b = M.load(grouped)
    conc = _miss(b, "context.electrochemistry.electrolyte.concentration_M", "run-001")
    ph = _miss(b, "context.electrochemistry.pH", "run-001")
    latest = B._decisions_by_target(b["scientist_decisions"])
    assert latest[("missing_value", conc["value_id"])]["value"] == 0.1
    assert latest[("missing_value", ph["value_id"])]["value"] == 13
    assert conc["value_state"] == ph["value_state"] == B.STATE_SCIENTIST_CONFIRMED
    assert [i["header"] for i in rep["invalid_cells"]] == ["Electrolyte Concentration (M)"]  # "0.1 M KOH" is not a number


def test_m4_both_electrolyte_columns_show_the_whole_ai_reading(grouped, tmp_path):
    b = M.load(grouped)
    name = next(c for c in b["candidates"] if c["run_id"] == "run-001" and c["target_field_path"] == "context.electrochemistry.electrolyte.name")
    reading = "0.1 M KOH, synthetic, " + "stated in a long synthetic logbook line " * 5  # longer than any old clip
    M.annotate(grouped, target_id=name["candidate_id"], value=reading, excerpt_locator="line 7", source_id=None,
               model_id="synthetic-model", blind=False, rationale="synthetic")
    xlsx = _export(grouped, tmp_path)
    sg = openpyxl.load_workbook(xlsx)["ISAAC Suggestions"]
    r = _row_of(xlsx, "run-001")
    assert reading in sg[f"{_col('Electrolyte')}{r}"].value
    conc = sg[f"{_col('Electrolyte Concentration (M)')}{r}"].value
    assert reading in conc and "the same reading; split it" in conc


def test_m4_potential_scale_maps_to_its_real_missing_value(grouped, tmp_path, capsys):
    b = M.load(grouped)
    assert _miss(b, "context.electrochemistry.potential_scale", "run-001")
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    status = {v[2]: v[4] for v in wb["_isaac"].iter_rows(values_only=True) if v[0] == "File List" and isinstance(v[1], str)}
    assert status["Potential Scale"] == "mapped"
    voc = wb["Vocabulary"]
    j = [c.value for c in voc[1]].index("Potential Scale") + 1
    assert [voc.cell(row=i, column=j).value for i in range(2, voc.max_row + 1) if voc.cell(row=i, column=j).value] == \
        list(B._schema_node("context.electrochemistry.potential_scale")["enum"])
    r1, r2 = _row_of(xlsx, "run-001"), _row_of(xlsx, "run-002")
    filled = _fill(xlsx, tmp_path / "scale.xlsx", {("File List", f"{_col('Potential Scale')}{r1}"): "RHE",
                                                   ("File List", f"{_col('Potential Scale')}{r2}"): "rhe"})
    code, rep, _ = _import(grouped, filled, capsys, "--apply")
    scale = _miss(M.load(grouped), "context.electrochemistry.potential_scale", "run-001")
    assert scale["value_state"] == B.STATE_SCIENTIST_CONFIRMED and rep["recorded"] == 1
    assert rep["invalid_cells"][0]["cell"] == f"'File List'!{_col('Potential Scale')}{r2}"


POT = "context.electrochemistry.potential_setpoint_V"


def test_potential_v_fills_a_run_with_no_parsed_potential_and_refuses_text(tmp_path, capsys):
    """Potential (V) used to be "recorded nowhere" for a run whose files state no potential:
    the set-point had no missing value. It is answerable now (bundle.LOCAL_ANSWER_FIELDS)."""
    from test_local_import_reader_ssrl_xdc import make_scan
    from test_local_import_reader_two_column import make_two

    src = tmp_path / "src" / "SYN_beamtime"
    make_scan(src / "SYN_run_A.001")
    make_two(src / "SYN_cell.txt")
    bdir = tmp_path / "b"
    assert cli.main(["import", str(src), "--out", str(bdir)]) == 0
    M.group(bdir)
    b = M.load(bdir)
    runs = {r["label"]: r["run_id"] for r in b["runs"]}
    assert not [c for c in b["candidates"] if c["target_field_path"] == POT]
    xlsx = _export(bdir, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    status = {v[2]: v[4] for v in wb["_isaac"].iter_rows(values_only=True) if v[0] == "File List" and isinstance(v[1], str)}
    assert status["Potential (V)"] == "mapped"
    r_cell, r_run = _row_of(xlsx, runs["SYN_cell"]), _row_of(xlsx, runs["SYN_run"])
    filled = _fill(xlsx, tmp_path / "pot.xlsx", {("File List", f"{_col('Potential (V)')}{r_cell}"): -0.8,
                                                 ("File List", f"{_col('Potential (V)')}{r_run}"): "-0.8 V"})
    code, rep, _ = _import(bdir, filled, capsys, "--apply")
    assert code == 0 and rep["recorded"] == 1
    assert [c["cell"] for c in rep["invalid_cells"]] == [f"'File List'!{_col('Potential (V)')}{r_run}"]  # text, not a number
    assert not [u for u in rep["recorded_nowhere"] if u["header"] == "Potential (V)"]
    b = M.load(bdir)
    m = _miss(b, POT, runs["SYN_cell"])
    assert m["value_state"] == B.STATE_SCIENTIST_CONFIRMED and m["eligibility_for_official_record"]["eligible"]
    assert B._decisions_by_target(b["scientist_decisions"])[("missing_value", m["value_id"])]["value"] == -0.8
    assert _miss(b, POT, runs["SYN_run"])["value_state"] == B.STATE_MISSING


def test_a_run_with_a_parsed_potential_has_its_candidate_and_no_missing_value(grouped):
    b = M.load(grouped)
    with_candidate = {c["run_id"] for c in b["candidates"] if c["target_field_path"] == POT}
    assert with_candidate == {r["run_id"] for r in b["runs"] if r["run_candidate"]}
    assert not [m for m in b["missing_values"] if m["target_field_path"] == POT]
    assert all(c["target_write_scope"] == "run" for c in b["candidates"] if c["target_field_path"] == POT)
    # the engine's "this build cannot write it" is about the website; this local path writes it
    assert {c["state_basis"]["not_proposable_summary"] for c in b["candidates"] if c["target_field_path"] == POT} == {
        B.LOCAL_WRITE_PATH_REASON}


def test_l10_a_campaign_answer_reaches_only_the_file_list_runs(grouped, tmp_path, capsys):
    b = M.load(grouped)
    g = b["record_groups"][0]
    M.decide(grouped, target_ids=[g["group_id"]], decision="reject", value=None, **BY)
    outside = set(M.load(grouped)["runs"][i]["run_id"] for i, r in enumerate(M.load(grouped)["runs"])
                  if r["run_key"] in set(g["member_run_keys"]))
    xlsx = _export(grouped, tmp_path)
    filled = _fill(xlsx, tmp_path / "camp.xlsx", {("Campaign Info", "B3"): "Synthetic Light Source"})
    code, rep, _ = _import(grouped, filled, capsys, "--apply")
    b = M.load(grouped)
    decided = {d["target"]["id"] for d in b["scientist_decisions"] if d["target"]["kind"] == "missing_value"}
    targets = [m for m in b["missing_values"] if m["target_field_path"] == "system.facility.facility_name"]
    assert {m["run_id"] for m in targets if m["value_id"] in decided} == {m["run_id"] for m in targets} - outside
    assert not any(m["value_id"] in decided for m in targets if m["run_id"] in outside)
    assert rep["recorded"] == 9 - len(outside)
    nf = rep["campaign_not_filled"][0]
    assert nf["targets"] == len(outside) and set(nf["run_ids"]) == outside
