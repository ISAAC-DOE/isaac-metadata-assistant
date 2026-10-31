"""The metadata template can finish a record (template 1.4.0), and says so briefly. Synthetic only.

* Every field the record composer requires a scientist to supply has a cell: the producer's group
  and organization and the domain on Campaign Info (one answer per beamtime), the reaction's drive
  and catalysis as File List columns (per run). Each maps to the path the questions ask about, and
  a closed list gets its dropdown from ``bundle.allowed_values``.
* Campaign Info's "Organization" was the FACILITY's organization; it says so now, and a 1.3.0
  workbook (which still reads "Organization") imports as before.
* A row's suggestion is that row's own evidence, never another run's reading.
* ``template export`` and ``template import`` print a short summary; the full lists are behind flags.
"""

from __future__ import annotations

import json
import shutil

import openpyxl
import pytest

from isaac_api.local_import import bundle as B
from isaac_api.local_import import cli
from isaac_api.local_import import mutations as M
from isaac_api.local_import import questions as Q
from isaac_api.local_import import template as T

from tests_support_records import pristine_bundle  # noqa: F401

WHO = "Synthetic Scientist"
FACILITY_ORG = "Facility organization (e.g. SLAC / SSRL)"
YOUR_GROUP = "Your group (who produced the data)"
YOUR_ORG = "Your organization (who produced the data)"
DOMAIN = "Domain"


@pytest.fixture(scope="module")
def grouped_base(pristine_bundle, tmp_path_factory):
    bdir = tmp_path_factory.mktemp("tcomplete") / "b"
    shutil.copytree(pristine_bundle, bdir)
    M.group(bdir)
    return bdir


@pytest.fixture()
def grouped(grouped_base, tmp_path):
    bdir = tmp_path / "b"
    shutil.copytree(grouped_base, bdir)
    return bdir


def _export(bdir, tmp_path, name="t.xlsx"):
    out = tmp_path / name
    assert cli.main(["bundle", "template", "export", str(bdir), "--out", str(out)]) == 0
    return out


def _import(bdir, xlsx, capsys, *extra):
    capsys.readouterr()
    code = cli.main(["bundle", "template", "import", str(bdir), str(xlsx), "--by", WHO, "--json", *extra])
    captured = capsys.readouterr()
    return code, (json.loads(captured.out) if code == 0 else None), captured.err


def _row_of(xlsx, run_id):
    wb = openpyxl.load_workbook(xlsx)
    for values in wb["_isaac"].iter_rows(values_only=True):
        if values[0] == "File List" and values[2] == run_id:
            return int(values[1])
    raise AssertionError(run_id)


def _col(header):
    from openpyxl.utils import get_column_letter

    return get_column_letter([c.header for c in T.FILE_COLUMNS].index(header) + 1)


def _campaign_row(header):
    return T.CAMPAIGN_FIRST_ROW + [c.header for c in T.CAMPAIGN_FIELDS].index(header)


def _fill(xlsx, out, cells):
    wb = openpyxl.load_workbook(xlsx)
    for (sheet, ref), value in cells.items():
        wb[sheet][ref] = value
    wb.save(out)
    return out


def _miss(b, path, run_id):
    return next(m for m in b["missing_values"] if m["target_field_path"] == path and m["run_id"] == run_id)


def _mapping(xlsx):
    wb = openpyxl.load_workbook(xlsx)
    return {(v[0], v[2]): (v[1], v[3]) for v in wb["_isaac"].iter_rows(values_only=True)
            if v[0] in ("Campaign Info", "File List", "Configurations") and isinstance(v[1], str)}


# --- (1) every field the composer requires a scientist to supply has a cell ------------------------


def test_the_template_has_a_cell_for_every_field_the_record_needs_from_a_scientist(grouped, tmp_path):
    assert T.FORMAT_VERSION == "1.4.0"
    xlsx = _export(grouped, tmp_path)
    mapping = _mapping(xlsx)
    # beamtime-wide (record scope): Campaign Info, appended after the template's own five rows
    assert mapping[("Campaign Info", YOUR_GROUP)] == ("B8", "attribution.produced_by.group")
    assert mapping[("Campaign Info", YOUR_ORG)] == ("B9", "attribution.produced_by.organization")
    assert mapping[("Campaign Info", DOMAIN)] == ("B10", "system.domain")
    # per run: File List columns, appended after the 1.3.0 ones (cells are read by position)
    headers = [c.header for c in T.FILE_COLUMNS]
    assert headers[-4:] == ["Control Mode", "Current Set-point (mA/cm2)", "Reaction Drive", "Reaction Catalysis"]
    assert T.COLUMNS_ADDED_IN["Reaction Drive"] == T.COLUMNS_ADDED_IN["Reaction Catalysis"] == "1.4.0"
    assert mapping[("File List", "Reaction Drive")] == (_col("Reaction Drive"), "context.reaction.drive")
    assert mapping[("File List", "Reaction Catalysis")] == (_col("Reaction Catalysis"), "context.reaction.catalysis")
    # THE SAME PATHS THE QUESTIONS ASK: every required path of the composer that is a scientist's answer
    for path in (*B.PORTAL_REQUIRED_PATHS, *B.REACTION_REQUIRED_PATHS, "system.domain"):
        assert path in {p for _, p in mapping.values()}, path
    # every new answer cell is EMPTY
    wb = openpyxl.load_workbook(xlsx)
    assert all(wb["Campaign Info"][f"B{r}"].value is None for r in (8, 9, 10))
    fl = wb["File List"]
    for header in ("Reaction Drive", "Reaction Catalysis"):
        assert fl[f"{_col(header)}1"].value == header
        assert all(fl[f"{_col(header)}{r}"].value is None for r in range(2, fl.max_row + 1))


def test_closed_lists_get_dropdowns_from_bundle_allowed_values(grouped, tmp_path):
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    voc = wb["Vocabulary"]
    headers = [c.value for c in voc[1]]
    for header, path in (("Domain", "system.domain"), ("Reaction Drive", "context.reaction.drive"),
                         ("Reaction Catalysis", "context.reaction.catalysis")):
        j = headers.index(header) + 1
        listed = [voc.cell(row=r, column=j).value for r in range(2, voc.max_row + 1)]
        assert [v for v in listed if v is not None] == B.allowed_values(path)
    fl = {str(dv.sqref): dv for dv in wb["File List"].data_validations.dataValidation}
    last = 1 + len(T.live_rows(M.load(grouped)))
    for header in ("Reaction Drive", "Reaction Catalysis"):
        dv = fl[f"{_col(header)}2:{_col(header)}{last}"]
        assert dv.type == "list" and "Vocabulary" in dv.formula1
    camp = {str(dv.sqref): dv for dv in wb["Campaign Info"].data_validations.dataValidation}
    assert camp["B10"].type == "list" and "Vocabulary" in camp["B10"].formula1
    assert "B8" not in camp and "B9" not in camp  # free text: no list closes them


def test_the_new_cells_record_through_the_same_decide_checks(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    r1, r2 = _row_of(xlsx, "run-001"), _row_of(xlsx, "run-002")
    filled = _fill(xlsx, tmp_path / "f.xlsx", {
        ("Campaign Info", "B8"): "Synthetic Group",
        ("Campaign Info", "B9"): "Synthetic University",
        ("Campaign Info", "B10"): "experimental",
        ("File List", f"{_col('Reaction Drive')}{r1}"): "electrochemical",
        ("File List", f"{_col('Reaction Catalysis')}{r1}"): "heterogeneous",
        ("File List", f"{_col('Reaction Drive')}{r2}"): "Electrochemical",  # not written exactly: refused
    })
    code, rep, _ = _import(grouped, filled, capsys, "--apply")
    assert code == 0
    assert [c["cell"] for c in rep["invalid_cells"]] == [f"'File List'!{_col('Reaction Drive')}{r2}"]
    assert "did you mean 'electrochemical'" in rep["invalid_cells"][0]["reason"]
    b = M.load(grouped)
    decided = B._decisions_by_target(b["scientist_decisions"])

    def answer(m):
        return B.final_value(None, decided[("missing_value", m["value_id"])])

    assert answer(_miss(b, "attribution.produced_by.group", None)) == "Synthetic Group"
    assert answer(_miss(b, "attribution.produced_by.organization", None)) == "Synthetic University"
    assert answer(_miss(b, "system.domain", None)) == "experimental"
    assert answer(_miss(b, "context.reaction.drive", "run-001")) == "electrochemical"
    assert answer(_miss(b, "context.reaction.catalysis", "run-001")) == "heterogeneous"
    assert _miss(b, "context.reaction.drive", "run-002")["value_state"] == B.STATE_MISSING
    assert cli.main(["bundle", "check", str(grouped)]) == 0


def test_a_domain_off_the_list_is_refused(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    filled = _fill(xlsx, tmp_path / "d.xlsx", {("Campaign Info", "B10"): "experiment"})
    code, rep, _ = _import(grouped, filled, capsys, "--apply")
    assert code == 0 and rep["recorded"] == 0
    assert [c["cell"] for c in rep["invalid_cells"]] == ["'Campaign Info'!B10"]


# --- (2) the facility's organization is not the producer's -----------------------------------------


def test_campaign_organization_is_labelled_as_the_facility_s(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    ws = openpyxl.load_workbook(xlsx)["Campaign Info"]
    assert ws["A4"].value == FACILITY_ORG and ws["A9"].value == YOUR_ORG
    assert _mapping(xlsx)[("Campaign Info", FACILITY_ORG)] == ("B4", "system.facility.organization")
    filled = _fill(xlsx, tmp_path / "o.xlsx", {("Campaign Info", "B4"): "Synthetic Lab", ("Campaign Info", "B9"): "Synthetic University"})
    code, rep, _ = _import(grouped, filled, capsys, "--apply")
    b = M.load(grouped)
    decided = B._decisions_by_target(b["scientist_decisions"])
    fac = [m for m in b["missing_values"] if m["target_field_path"] == "system.facility.organization"]
    assert fac and all(B.final_value(None, decided[("missing_value", m["value_id"])]) == "Synthetic Lab" for m in fac)
    prod = _miss(b, "attribution.produced_by.organization", None)
    assert B.final_value(None, decided[("missing_value", prod["value_id"])]) == "Synthetic University"


def _as_1_3_0(xlsx, out):
    """The workbook a 1.3.0 build wrote: five Campaign Info rows, "Organization", no reaction columns."""
    wb = openpyxl.load_workbook(xlsx)
    wb["Campaign Info"]["A4"] = "Organization"
    for r in (8, 9, 10):
        wb["Campaign Info"][f"A{r}"] = None
    for header in ("Reaction Drive", "Reaction Catalysis"):
        wb["File List"][f"{_col(header)}1"] = None
    for row in wb["_isaac"].iter_rows():
        if row[0].value == "format_version":
            row[1].value = "1.3.0"
        if row[0].value == "Campaign Info" and row[2].value == FACILITY_ORG:
            row[2].value = "Organization"
        if (row[0].value == "Campaign Info" and row[1].value in ("B8", "B9", "B10")) or (
                row[0].value == "File List" and row[2].value in ("Reaction Drive", "Reaction Catalysis")):
            for cell in row:
                cell.value = None
    wb.save(out)
    return out


def test_a_1_3_0_workbook_with_the_old_organization_label_still_imports(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    old = _as_1_3_0(xlsx, tmp_path / "old.xlsx")
    r1 = _row_of(old, "run-001")
    filled = _fill(old, tmp_path / "old-filled.xlsx", {
        ("Campaign Info", "B4"): "Synthetic Lab",
        ("File List", f"{_col('Reaction Drive')}{r1}"): "electrochemical",  # no such column in 1.3.0: never read
    })
    code, rep, err = _import(grouped, filled, capsys, "--apply")
    assert code == 0, err
    paths = {d["target_field_path"] for d in rep["decisions"]}
    assert paths == {"system.facility.organization"}
    [ignored] = [n for n in rep["recorded_nowhere"] if n["header"] == "Reaction Drive"]
    assert "template '1.3.0'" in ignored["reason"]
    assert _miss(M.load(grouped), "context.reaction.drive", "run-001")["value_state"] == B.STATE_MISSING


def test_the_raw_blank_template_is_still_refused(grouped, tmp_path, capsys):
    raw = openpyxl.Workbook()
    raw.active.title = "Campaign Info"
    raw.active["A4"], raw.active["B4"] = "Organization", "Example University"
    raw.save(tmp_path / "raw.xlsx")
    code, _, err = _import(grouped, tmp_path / "raw.xlsx", capsys, "--apply")
    assert code == 3 and "_isaac" in err and not M.load(grouped)["scientist_decisions"]


# --- (4) a row's suggestion is that row's own evidence ---------------------------------------------


def test_a_row_shows_only_its_own_runs_readings_never_another_runs(grouped, tmp_path, monkeypatch):
    b = M.load(grouped)
    m1 = _miss(b, "context.environment", "run-001")
    own = next(c for c in b["candidates"] if c["run_id"] == "run-001" and c["observed_literals"])
    other = next(c for c in b["candidates"] if c["run_id"] == "run-002" and c["observed_literals"])
    real = Q.consolidate

    def consolidate(bundle, *a, **k):
        payload = real(bundle, *a, **k)
        # the consolidated question for a beamtime-wide path carries readings of OTHER runs as hints
        payload["questions"].insert(0, {"targets": [{"kind": "missing_value", "id": m1["value_id"]}], "hints": [
            {"kind": "related_reading", "candidate_id": other["candidate_id"], "literal": "OTHER-RUN-READING",
             "source_id": None, "locator": "elsewhere"},
            {"kind": "folder_name", "literal": "BEAMTIME-FOLDER", "locator": "the imported folder's name"},
            {"kind": "related_reading", "candidate_id": own["candidate_id"], "literal": "OWN-RUN-READING",
             "source_id": None, "locator": "here"},
        ]})
        return payload

    monkeypatch.setattr(Q, "consolidate", consolidate)
    xlsx = _export(grouped, tmp_path)
    sg = openpyxl.load_workbook(xlsx)["ISAAC Suggestions"]
    env_col = [c.header for c in T.FILE_COLUMNS].index("Environment") + 1
    text = sg.cell(row=_row_of(xlsx, "run-001"), column=env_col).value
    assert "OWN-RUN-READING" in text
    assert "OTHER-RUN-READING" not in text and "BEAMTIME-FOLDER" not in text
    # the answer cell itself stays empty
    fl = openpyxl.load_workbook(xlsx)["File List"]
    assert fl.cell(row=_row_of(xlsx, "run-001"), column=env_col).value is None


def test_a_file_name_value_is_shown_as_that_rows_suggestion_and_its_answer_cell_stays_empty(grouped, tmp_path):
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    col = [c.header for c in T.FILE_COLUMNS].index("Potential (V)") + 1
    for run_id, value in (("run-001", 0.1), ("run-002", 0.2)):
        r = _row_of(xlsx, run_id)
        assert wb["ISAAC Suggestions"].cell(row=r, column=col).value.startswith(f"extracted {value!r}")
        assert wb["File List"].cell(row=r, column=col).value is None


# --- (3) (5) short output; the full lists behind a flag ---------------------------------------------


def test_export_prints_a_short_summary_and_json_keeps_the_columns(grouped, tmp_path, capsys):
    capsys.readouterr()
    out = tmp_path / "s.xlsx"
    assert cli.main(["bundle", "template", "export", str(grouped), "--out", str(out)]) == 0
    text = capsys.readouterr().out
    assert str(out) in text and "9 File List row(s)" in text and "template import" in text
    assert "target_field_path" not in text and len(text.splitlines()) <= 12
    assert cli.main(["bundle", "template", "export", str(grouped), "--out", str(out), "--force", "--json"]) == 0
    full = json.loads(capsys.readouterr().out)
    assert full["written"] == str(out) and any(c["target_field_path"] == "system.domain" for c in full["columns"])


def test_the_import_preview_prints_one_line_per_cell_or_column_and_full_lists_on_request(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    cells = {("Campaign Info", "B5"): "ZZ-1", ("Campaign Info", "B10"): "experimental",
             ("Campaign Info", "B6"): "2099-01-01"}
    for run_id in ("run-001", "run-002", "run-003"):
        cells[("File List", f"F{_row_of(xlsx, run_id)}")] = "ex_situ"
    bad = _row_of(xlsx, "run-004")
    cells[("File List", f"F{bad}")] = "dry"  # not an allowed value: shown in full
    filled = _fill(xlsx, tmp_path / "p.xlsx", cells)
    capsys.readouterr()
    assert cli.main(["bundle", "template", "import", str(grouped), str(filled), "--by", WHO]) == 0
    text = capsys.readouterr().out
    lines = text.splitlines()
    assert any("'Campaign Info'!B5" in ln and "9 decision(s)" in ln for ln in lines)
    assert any("File List / Environment" in ln and "3 decision(s)" in ln and "3 row(s)" in ln for ln in lines)
    assert f"'File List'!F{bad}" in text and "not an allowed value" in text  # invalid cells: in full
    assert "Start Date" in text  # unmapped: one line per column
    assert not any(" missing_value " in ln for ln in lines)  # no per-target line
    assert "--full" in text and len(lines) < 30
    capsys.readouterr()
    assert cli.main(["bundle", "template", "import", str(grouped), str(filled), "--by", WHO, "--full"]) == 0
    full = capsys.readouterr().out
    assert sum(" missing_value " in ln for ln in full.splitlines()) == 9 + 1 + 3
