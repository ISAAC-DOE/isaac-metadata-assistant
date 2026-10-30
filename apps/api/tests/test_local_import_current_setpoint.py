"""``context.electrochemistry.current_setpoint_mA_cm2`` is locally answerable (bundle 1.7.0, DEC-92),
as ``potential_setpoint_V`` became at 1.3.0 and ``control_mode`` at 1.5.0.

Before this, a confirmed ``galvanostatic`` mode made the current set-point required (the schema's own
conditional rule) while nothing could write it, so such a record stayed blocked for ever.

* RUN scope, like the potential set-point; a ``number`` (mA per cm², the unit in the name), checked
  against the vendored schema's ``type`` by eligibility — text such as ``"5 mA"`` is confirmed but
  never composed. The questions say it needs ``--json``.
* NO companion gate (unlike the potential's scale): the schema declares no reference for it and
  requires nothing beside it. Never derived — not from a total current and an electrode area, not
  from the mode, and the mode never from it.
* A 1.6.0 bundle in which a value at this path moves eligibility is STALE, not INVALID.

SYNTHETIC files only (the grouping CLI fixture).
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from isaac_api.local_import import bundle as B
from isaac_api.local_import import cli
from isaac_api.local_import import mutations as M
from isaac_api.local_import import questions as Q
from isaac_api.local_import import records as R
from isaac_api.local_import import template as T

from test_local_import_control_mode import _col, _compose, _export, _group_of, _import, _miss, _row_of, _set_mode
from test_local_import_grouping_cli import _folder

BY = "SYNTHETIC TEST"
CM = B.CONTROL_MODE_PATH
CURRENT = B.CURRENT_SETPOINT_PATH
POT = "context.electrochemistry.potential_setpoint_V"
SCALE = "context.electrochemistry.potential_scale"
COLUMN = "Current Set-point (mA/cm2)"


def _json(b: Path) -> dict:
    return json.loads((b / "bundle.json").read_text())


@pytest.fixture()
def grouped(tmp_path):
    src, b = _folder(tmp_path), tmp_path / "bundle"
    assert cli.main(["import", str(src), "--out", str(b)]) == 0
    M.group(b)
    return b


def _set_current(b, run_id, value, *, older_build=False):
    vid = _miss(M.load(b), CURRENT, run_id)["value_id"]
    if older_build:  # not a number: `decide` refuses it now (contract §25); an older build stored it
        from tests_support_records import refused_then_stored

        assert "must be number" in refused_then_stored(b, "missing_value", vid, value)
        return
    M.decide(b, target_ids=[vid], decision="correct", value=value, by=BY, orcid=None)


# --- writable, run scope, new in 1.7.0, no companion -------------------------------------------------


def test_the_current_setpoint_is_a_run_scope_local_answer_field_new_in_1_7_0():
    assert CURRENT == "context.electrochemistry.current_setpoint_mA_cm2"
    assert B.LOCAL_ANSWER_FIELDS[CURRENT] == B.LOCAL_ANSWER_FIELDS[POT] == "run"
    assert CURRENT in B._writable_paths() and B.WRITABLE_ADDED_IN["1.7.0"] == frozenset({CURRENT})
    assert B._VERSION_ORDER.index("1.7.0") == B._VERSION_ORDER.index(B.FORMAT_VERSION) - 2  # 1.8.0 bounds, 1.9.0 `last`: neither writable
    assert B._added_since("1.6.0") == frozenset({CURRENT}) and B._added_since("1.7.0") == frozenset()
    assert not Q._is_shared(CURRENT)  # never asked once for the whole beamtime
    # DEC-92: no companion gate — the schema declares no reference for a current density
    assert CURRENT not in B.COMPOSES_ONLY_WITH and CURRENT not in B.COMPOSES_ONLY_WITH.values()
    assert B._schema_node(CURRENT) == {"type": "number"}  # the vendored schema says only "a number"


def test_the_bundle_schema_accepts_1_7_0_and_requires_processed_outputs_there(grouped):
    data = _json(grouped)
    assert data["format_version"] == B.FORMAT_VERSION and B.validate_bundle(data) == []
    data["format_version"] = "1.7.0"
    assert B.validate_bundle(data) == []  # a 1.7.0 bundle still validates under the newer schema
    del data["processed_outputs"]
    assert B.validate_bundle(data)  # 1.7.0, like 1.2.0-1.6.0, requires processed_outputs


def test_every_run_has_a_missing_value_and_nothing_is_defaulted(grouped):
    bd = M.load(grouped)
    assert not [c for c in bd["candidates"] if c["target_field_path"] == CURRENT]  # no parser states it
    ms = [m for m in bd["missing_values"] if m["target_field_path"] == CURRENT]
    assert {m["run_id"] for m in ms} == {r["run_id"] for r in bd["runs"] if r["run_candidate"]}
    assert all(m["value"] is None and m["value_state"] == B.STATE_MISSING and m["target_write_scope"] == "run" for m in ms)
    assert all(not m["eligibility_for_official_record"]["eligible"] for m in ms)


def test_current_setpoint_value_ids_come_after_every_older_fields_ids(grouped):
    ms = M.load(grouped)["missing_values"]
    nums = {m["target_field_path"]: [] for m in ms}
    for m in ms:
        nums[m["target_field_path"]].append(int(m["value_id"].split("-")[1]))
    older = sorted(n for p, ns in nums.items() if p not in (CM, CURRENT) for n in ns)
    assert older == list(range(1, len(older) + 1))
    assert max(older) < min(nums[CM]) and max(nums[CM]) < min(nums[CURRENT])  # 1.5.0's, then 1.7.0's


# --- galvanostatic: the requirement and the answer that meets it -------------------------------------


def test_a_galvanostatic_run_with_its_current_setpoint_answered_composes_and_is_not_missing_anything_for_it(grouped):
    _set_mode(grouped, "run-001", "galvanostatic")
    _set_current(grouped, "run-001", -5.0)  # a number, as `decide --value -5 --json` stores it
    m = _miss(M.load(grouped), CURRENT, "run-001")
    assert m["value_state"] == B.STATE_SCIENTIST_CONFIRMED and m["eligibility_for_official_record"]["eligible"] is True
    draft, report, entry = _compose(grouped, "run-001")
    assert draft["fields"][CM]["value"] == "galvanostatic"
    assert draft["fields"][CURRENT]["value"] == -5.0 and draft["fields"][CURRENT]["status"] == "verified"
    assert {"path": CURRENT, "from": [m["value_id"]]} in report["filled"]
    assert CURRENT not in entry["missing_required"] and CURRENT not in {x["path"] for x in report["missing_required"]}
    assert POT not in entry["missing_required"]  # galvanostatic does not require a potential set-point
    # the required-path check itself: what the record carries meets the schema's galvanostatic rule
    required = B.required_record_paths(control_mode=draft["fields"][CM]["value"])
    assert CURRENT in required and draft["fields"][CURRENT]["value"] is not None
    # nothing was derived beside it: no area, no potential, no scale
    assert not [p for p in draft["fields"] if "geometry" in p or "area" in p]
    assert POT not in draft["fields"] and SCALE not in draft["fields"]


def test_a_galvanostatic_run_without_its_current_setpoint_is_incomplete(grouped):
    _set_mode(grouped, "run-001", "galvanostatic")
    draft, report, entry = _compose(grouped, "run-001")
    assert {"path": CURRENT, "from": "a scientist's answer"} in report["missing_required"]
    assert draft["fields"][CURRENT] == {"value": None, "status": "missing", "evidence": []}
    assert entry["complete"] is False and CURRENT in entry["missing_required"]


def test_an_off_type_current_setpoint_is_confirmed_but_not_eligible_and_the_record_stays_incomplete(grouped):
    _set_mode(grouped, "run-001", "galvanostatic")
    _set_current(grouped, "run-001", "5 mA", older_build=True)
    _set_current(grouped, "run-002", "-5", older_build=True)  # a number given WITHOUT --json is text: also not eligible
    bd = M.load(grouped)
    for run_id, value in (("run-001", "5 mA"), ("run-002", "-5")):
        m = _miss(bd, CURRENT, run_id)
        assert m["value"] is None and m["value_state"] == B.STATE_SCIENTIST_CONFIRMED  # confirmed: the answer is kept
        assert m["eligibility_for_official_record"]["eligible"] is False
        assert "must be number; got str" in m["eligibility_for_official_record"]["reason"]
    draft, report, entry = _compose(grouped, "run-001")
    assert draft["fields"][CURRENT]["value"] is None  # the text never reaches the draft
    [listed] = [n for n in report["not_composed"] if n.get("path") == CURRENT]
    assert listed["value_id"] == _miss(bd, CURRENT, "run-001")["value_id"] and "must be number" in listed["reason"]
    assert entry["complete"] is False and CURRENT in entry["missing_required"]


def test_potentiostatic_and_open_circuit_runs_do_not_require_it(grouped):
    _set_mode(grouped, "run-001", "potentiostatic")
    _set_mode(grouped, "run-002", "open_circuit")
    for run_id in ("run-001", "run-002"):
        draft, report, entry = _compose(grouped, run_id)
        assert CURRENT not in entry["missing_required"] and CURRENT not in draft["fields"]
    assert CURRENT not in B.required_record_paths(control_mode="potentiostatic")
    assert CURRENT not in B.required_record_paths()  # no mode confirmed: nothing keyed on it


# --- no companion gate, no inference ------------------------------------------------------------------


def test_the_current_setpoint_composes_without_a_scale_or_an_area_and_gives_no_mode(grouped):
    _set_current(grouped, "run-001", 2.5)
    draft, report, _ = _compose(grouped, "run-001")
    assert draft["fields"][CURRENT]["value"] == 2.5  # no companion gate (DEC-92)
    assert not [n for n in report["not_composed"] if n.get("path") == CURRENT]
    assert CM not in draft["fields"] and B.control_mode_by_run(M.load(grouped)) == {}  # no mode inferred from it
    assert _miss(M.load(grouped), CM, "run-001")["value_state"] == B.STATE_MISSING


def test_a_confirmed_galvanostatic_mode_never_gives_a_current_setpoint(grouped):
    _set_mode(grouped, "run-001", "galvanostatic")
    assert _miss(M.load(grouped), CURRENT, "run-001")["value_state"] == B.STATE_MISSING
    draft, _, _ = _compose(grouped, "run-001")
    assert draft["fields"][CURRENT]["value"] is None


def test_two_different_confirmed_current_setpoints_in_one_group_are_a_disagreement(grouped):
    _set_current(grouped, "run-001", -5.0)
    _set_current(grouped, "run-002", -10.0)
    bd = M.load(grouped)
    group = {"group_id": "rg-SYNTHETIC-two-runs", "status": "scientist_confirmed",
             "member_run_keys": [r["run_key"] for r in bd["runs"] if r["run_candidate"]]}
    draft, report = R.compose_group(bd, group, grouped)
    [dis] = [d for d in report["disagreements"] if d["path"] == CURRENT]
    assert sorted(dis["values"]) == [-10.0, -5.0] and CURRENT not in draft["fields"]


# --- questions ------------------------------------------------------------------------------------------


def test_a_galvanostatic_run_has_its_current_setpoint_asked_as_blocking_with_the_json_hint(grouped):
    def asked(bd, run_id):
        scope = f"group:{_group_of(bd, run_id)}"
        return [q for q in Q.consolidate(bd)["questions"] if q["scope"] == scope
                and (q["target_field_path"] == CURRENT or CURRENT in (q.get("fields") or {}))]

    bd = M.load(grouped)
    [before] = asked(bd, "run-001")
    assert before["kind"] == "optional_values" and not before["blocking"]
    assert "--json" in before["answer_command"] and CURRENT in before["answer_command"]  # it says which fields need --json
    _set_mode(grouped, "run-001", "galvanostatic")
    _set_mode(grouped, "run-002", "potentiostatic")
    bd = M.load(grouped)
    [after] = asked(bd, "run-001")
    assert after["kind"] == "missing_value" and after["blocking"] and CURRENT in after["why"]
    assert "--value <value> --json" in after["answer_command"]
    [other] = asked(bd, "run-002")
    assert not other["blocking"]  # potentiostatic: optional


def test_the_cli_stores_a_number_with_json_and_text_without(grouped, capsys):
    _set_mode(grouped, "run-001", "galvanostatic")
    bd = M.load(grouped)
    qid = next(q["question_id"] for q in Q.consolidate(bd)["questions"]
               if q["target_field_path"] == CURRENT and q["scope"] == f"group:{_group_of(bd, 'run-001')}")
    base = ["bundle", "decide", str(grouped), "--decision", "correct", "--by", BY]
    capsys.readouterr()
    # without --json the answer is the text "-5", which the field refuses: refused, nothing recorded,
    # and the refusal says to give the number with --json (contract §25)
    assert cli.main([*base, "--question", qid, "--value", "-5"]) == 2
    assert "must be number" in (err := capsys.readouterr().err) and "--json" in err
    m = _miss(M.load(grouped), CURRENT, "run-001")
    assert m["value_state"] == B.STATE_MISSING
    assert cli.main([*base, "--question", qid, "--value", "-5", "--json"]) == 0
    m = _miss(M.load(grouped), CURRENT, "run-001")
    assert m["eligibility_for_official_record"]["eligible"] is True
    _, _, entry = _compose(grouped, "run-001")
    assert CURRENT not in entry["missing_required"]


# --- a 1.6.0 bundle --------------------------------------------------------------------------------------


UPGRADE_16 = ("upgraded from 1.6.0; fields this build can newly write have no answer slots on runs that lacked a parsed "
              "value — re-import to get them" + "".join(f"; {c}" for c in B.CHANGES_ADDED_IN.values()))  # 1.8.0, 1.9.0: disclosed


def test_a_1_6_0_bundle_with_a_confirmed_current_setpoint_is_stale_not_invalid_then_rederived(grouped, capsys):
    _set_current(grouped, "run-001", -5.0)
    data = _json(grouped)
    vid = _miss(data, CURRENT, "run-001")["value_id"]
    data["format_version"] = "1.6.0"
    for m in data["missing_values"]:
        if m["value_id"] == vid:  # what the 1.6.0 build stored: no write path for the field
            m["eligibility_for_official_record"] = {"eligible": False, "reason": "Confirmed, but this build has no write path for its official field, so it stays evidence."}
            data["readiness"]["eligible_value_count"] -= 1
    (grouped / "bundle.json").write_text(json.dumps(data))
    assert B.validate_bundle(_json(grouped))  # the plain validator refuses it
    errors, stale = B.validate_allowing_stale(_json(grouped))
    assert errors == [] and [(s["kind"], s["id"]) for s in stale] == [("missing_value", vid)]
    assert stale[0]["reason"].startswith(f"this build can now write {CURRENT}") and stale[0]["reason"].endswith("eligible now")
    assert "composes only beside" not in stale[0]["reason"]  # no companion gate
    capsys.readouterr()
    assert cli.main(["bundle", "check", str(grouped)]) == cli.EXIT_INVALID
    out = capsys.readouterr().out
    assert out.startswith("STALE:") and vid in out and UPGRADE_16 in out
    result = R.build_records(grouped)
    assert [x["id"] for x in result["eligibility_rederived_now"]] == [vid]
    after = _json(grouped)
    assert after["format_version"] == B.FORMAT_VERSION and B.validate_bundle(after) == []
    [u] = after["format_upgrades"]
    assert (u["from_format_version"], u["newly_writable_paths"], u["statement"]) == ("1.6.0", [CURRENT], UPGRADE_16)


def test_a_1_6_0_bundle_as_that_build_made_it_is_valid_and_upgraded_once_with_no_new_slots(grouped, capsys):
    data = _json(grouped)
    data["missing_values"] = [m for m in data["missing_values"] if m["target_field_path"] != CURRENT]
    data["format_version"] = "1.6.0"
    B.refresh_questions(data)
    B.apply_derivations(data)
    (grouped / "bundle.json").write_text(json.dumps(data))
    capsys.readouterr()
    assert cli.main(["bundle", "check", str(grouped)]) == 0
    other = next(m for m in data["missing_values"] if m["target_field_path"] == "context.environment")
    M.decide(grouped, target_ids=[other["value_id"]], decision="correct", value="in_situ", by=BY, orcid=None)
    assert capsys.readouterr().err.count(UPGRADE_16) == 1
    after = _json(grouped)
    assert after["format_version"] == B.FORMAT_VERSION and [u["newly_writable_paths"] for u in after["format_upgrades"]] == [[CURRENT]]
    assert not [m for m in after["missing_values"] if m["target_field_path"] == CURRENT]  # a save adds no answer slot


# --- the metadata template ---------------------------------------------------------------------------------


def test_the_template_current_setpoint_column_takes_a_number_only(grouped, tmp_path, capsys):
    assert T.COLUMNS_ADDED_IN[COLUMN] == "1.3.0"
    headers = [c.header for c in T.FILE_COLUMNS]
    # appended, never inserted: right after Control Mode; only later versions' columns follow it
    assert headers[headers.index(COLUMN) - 1] == "Control Mode"
    assert all(T.COLUMNS_ADDED_IN.get(h, "0") > "1.3.0" for h in headers[headers.index(COLUMN) + 1:])
    assert T.field_rules()[CURRENT]["numeric"] is True
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    assert wb["File List"][f"{_col(COLUMN)}1"].value == COLUMN
    r1, r2 = _row_of(xlsx, "run-001"), _row_of(xlsx, "run-002")
    wb["File List"][f"{_col(COLUMN)}{r1}"] = -5
    wb["File List"][f"{_col(COLUMN)}{r2}"] = "5 mA"
    wb.save(tmp_path / "f.xlsx")
    code, rep = _import(grouped, tmp_path / "f.xlsx", capsys, "--apply")
    assert code == 0 and rep["recorded"] == 1
    assert [c["cell"] for c in rep["invalid_cells"]] == [f"'File List'!{_col(COLUMN)}{r2}"]
    bd = M.load(grouped)
    m1 = _miss(bd, CURRENT, "run-001")
    assert m1["eligibility_for_official_record"]["eligible"] and B.final_value(None, B._decisions_by_target(bd["scientist_decisions"])[("missing_value", m1["value_id"])]) == -5
    assert _miss(bd, CURRENT, "run-002")["value_state"] == B.STATE_MISSING


def test_a_1_2_0_workbook_without_the_column_is_still_read_and_a_cell_past_it_is_said(grouped, tmp_path, capsys):
    xlsx = _export(grouped, tmp_path)
    wb = openpyxl.load_workbook(xlsx)
    letter = _col(COLUMN)
    wb["File List"][f"{letter}1"] = None  # as 1.2.0 wrote it: no current set-point column
    for row in wb["_isaac"].iter_rows():
        if row[0].value == "format_version":
            row[1].value = "1.2.0"
        if row[0].value == "File List" and row[2].value == COLUMN:
            for cell in row:
                cell.value = None
    r1 = _row_of(xlsx, "run-001")
    wb["File List"][f"{_col('Control Mode')}{r1}"] = "galvanostatic"  # a 1.2.0 column: read
    wb["File List"][f"{letter}{r1}"] = -5  # past Control Mode in a 1.2.0 workbook: never read
    wb.save(tmp_path / "old.xlsx")
    code, preview = _import(grouped, tmp_path / "old.xlsx", capsys)
    [ignored] = [n for n in preview["recorded_nowhere"] if n["cell"] == f"'File List'!{letter}{r1}"]
    assert ignored["header"] == COLUMN and "template '1.2.0'" in ignored["reason"] and "re-export" in ignored["reason"]
    code, rep = _import(grouped, tmp_path / "old.xlsx", capsys, "--apply")
    assert code == 0 and rep["recorded"] == 1
    bd = M.load(grouped)
    assert _miss(bd, CM, "run-001")["value_state"] == B.STATE_SCIENTIST_CONFIRMED
    assert _miss(bd, CURRENT, "run-001")["value_state"] == B.STATE_MISSING
    # the same workbook claiming 1.3.0 lacks a column 1.3.0 writes: refused, nothing recorded
    wb = openpyxl.load_workbook(tmp_path / "old.xlsx")
    for row in wb["_isaac"].iter_rows():
        if row[0].value == "format_version":
            row[1].value = "1.3.0"
    wb.save(tmp_path / "claims_new.xlsx")
    code, _ = _import(grouped, tmp_path / "claims_new.xlsx", capsys, "--apply")
    assert code == T.EXIT_INPUT_REFUSED
