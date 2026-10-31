"""Polish c3-r2: the scientist's own filled sheet, after review. Synthetic only.

* H1: an AI reading of a FILE is a file reading at import too. A template ``adopt`` wrote earlier holds a
  value no reading contradicted then; once an AI reading of the run's scan file contradicts it, ``template
  import`` names it under ``replaces_or_contradicts_a_file_reading`` and lists it as awaiting a ruling
  instead of recording it. A cell typed into an ordinary exported template that an AI reading contradicts
  is named too, and still recorded: the scientist typed it. The preview says which.
* L1: a File List row names a run by the run's IMMEDIATE folder (or its stem / file name), never by an
  ancestor folder; every row that reaches more than one run is reported with the count.
* L3: ISAAC's blank template is recognised by its example values, not only by an empty File List: a copy
  that ships example File List rows is refused too.
"""

from __future__ import annotations

import json
import shutil

import openpyxl
import pytest

from isaac_api.local_import import bundle as B
from isaac_api.local_import import cli
from isaac_api.local_import import mutations as M
from isaac_api.local_import import template as T

from tests_support_records import SEMANTICS
from test_local_import_template_adopt import (  # noqa: F401  (fixtures)
    THEIRS, WHO, _adopt, _adopt_json, _col, _d, _import, _miss, _row_of, _their_sheet, _stems, corpus, grouped,
    grouped_base)

FORM = "sample.sample_form"


def _annotate(bdir, run_id, path, value, *, model_id="synthetic-model"):
    """An AI reading of ``run_id``'s own scan file (a FILE reading, not their sheet's)."""
    b = M.load(bdir)
    src = {s["source_id"]: s for s in b["sources"]}
    run = next(r for r in b["runs"] if r["run_id"] == run_id)
    rel = src[run["acquisition_source_id"]]["relative_path"]
    findings = {"format": "isaac.reading_findings", "findings": [
        {"field_path": path, "value": value, "applies_to": {"scope": "runs", "runs": [run_id]},
         "source": {"path": rel, "locator": "header line 3"}, "quote": f"sample: {value}", "strength": "stated",
         "reasoning": "the scan header names the sample form"}]}
    rep = M.annotate_from_readings(bdir, findings, model_id=model_id, apply=True)
    assert rep["interpretations"] >= 1, rep
    return rel


# --- H1: an AI reading of a file contradicts a value adopted before it ---------------------------


def test_an_adopted_value_an_ai_reading_of_a_file_now_contradicts_awaits_a_ruling_and_is_not_recorded(grouped, tmp_path, capsys):
    theirs = _their_sheet(grouped, tmp_path / THEIRS)
    earlier = tmp_path / "adopted.xlsx"
    _adopt_json(grouped, theirs, earlier, capsys)
    r1 = _row_of(earlier, "run-001")
    cell = f"'File List'!{_col('Sample Form')}{r1}"
    assert openpyxl.load_workbook(earlier)["File List"][f"{_col('Sample Form')}{r1}"].value == "powder"

    rel = _annotate(grouped, "run-001", FORM, "pellet")
    fresh = _adopt_json(grouped, theirs, tmp_path / "fresh.xlsx", capsys)
    assert any("pellet" in r and rel in r for r in _d(fresh, "Sample Form", run_id="run-001")["readings"])

    code, rep, err = _import(grouped, earlier, capsys)
    assert code == 0, err
    assert not [d for d in rep["decisions"] if d["target_field_path"] == FORM and d["run_id"] == "run-001"]
    assert [d for d in rep["decisions"] if d["target_field_path"] == FORM and d["run_id"] == "run-002"]  # uncontradicted
    named = [x for x in rep["replaces_or_contradicts_a_file_reading"] if x["cell"] == cell and x["run_id"] == "run-001"]
    assert len(named) == 1 and named[0]["outcome"] == "awaiting_ruling"
    assert any("'pellet'" in r and "AI reading" in r and rel in r for r in named[0]["readings"])
    waiting = [x for x in rep["awaiting_ruling"] if x["cell"] == cell]
    assert len(waiting) == 1 and waiting[0]["run_id"] == "run-001" and waiting[0]["their_value"] == "powder"
    assert "pellet" in waiting[0]["why"] and "template adopt" in waiting[0]["why"]

    capsys.readouterr()
    assert cli.main(["bundle", "template", "import", str(grouped), str(earlier), "--by", WHO]) == 0
    text = capsys.readouterr().out
    block = text.split("Replaces or contradicts a file reading")[1].split("\n\n")[0]
    line = next(ln for ln in block.splitlines() if cell in ln and "run-001" in ln)
    assert "pellet" in line and "not recorded" in line and "awaiting" in line
    awaiting = text.split("Disagreements awaiting your ruling")[1].split("\n\n")[0]
    assert cell in awaiting and "run-001" in awaiting

    code, rep, err = _import(grouped, earlier, capsys, "--apply")
    assert code == 0, err
    b = M.load(grouped)
    latest = B._decisions_by_target(b["scientist_decisions"])
    assert ("missing_value", _miss(b, FORM, "run-001")["value_id"]) not in latest
    assert latest[("missing_value", _miss(b, FORM, "run-002")["value_id"])]["value"] == "powder"


def test_a_typed_cell_an_ai_reading_of_a_file_contradicts_is_named_and_still_recorded(grouped, tmp_path, capsys):
    rel = _annotate(grouped, "run-001", FORM, "pellet")
    xlsx = tmp_path / "t.xlsx"
    assert cli.main(["bundle", "template", "export", str(grouped), "--out", str(xlsx)]) == 0
    r1 = _row_of(xlsx, "run-001")
    wb = openpyxl.load_workbook(xlsx)
    wb["File List"][f"{_col('Sample Form')}{r1}"] = "powder"
    wb.save(tmp_path / "f.xlsx")
    cell = f"'File List'!{_col('Sample Form')}{r1}"

    code, rep, err = _import(grouped, tmp_path / "f.xlsx", capsys)
    assert code == 0, err
    named = [x for x in rep["replaces_or_contradicts_a_file_reading"] if x["cell"] == cell]
    assert len(named) == 1 and named[0]["outcome"] == "recorded_as_typed"
    assert any("'pellet'" in r and rel in r for r in named[0]["readings"])
    assert [d for d in rep["decisions"] if d["cell"] == cell and d["value"] == "powder"]
    assert not rep["awaiting_ruling"]
    capsys.readouterr()
    assert cli.main(["bundle", "template", "import", str(grouped), str(tmp_path / "f.xlsx"), "--by", WHO]) == 0
    line = next(ln for ln in capsys.readouterr().out.splitlines() if ln.startswith(f"  {cell}") and "pellet" in ln)
    assert "you typed it" in line

    code, rep, err = _import(grouped, tmp_path / "f.xlsx", capsys, "--apply")
    assert code == 0, err
    b = M.load(grouped)
    assert B._decisions_by_target(b["scientist_decisions"])[("missing_value", _miss(b, FORM, "run-001")["value_id"])]["value"] == "powder"


def test_an_ai_reading_of_their_sheet_itself_does_not_hold_an_adopted_value_back(grouped, tmp_path, capsys):
    theirs = _their_sheet(grouped, tmp_path / THEIRS)
    earlier = tmp_path / "adopted.xlsx"
    _adopt_json(grouped, theirs, earlier, capsys)
    findings = {"format": "isaac.reading_findings", "findings": [
        {"field_path": FORM, "value": "pellet", "applies_to": {"scope": "runs", "runs": ["run-001"]},
         "source": {"path": THEIRS, "locator": "File List!E2"}, "quote": "pellet", "strength": "stated", "reasoning": "x"}]}
    assert M.annotate_from_readings(grouped, findings, model_id="synthetic-model", apply=True)["interpretations"] >= 1
    code, rep, err = _import(grouped, earlier, capsys)
    assert code == 0, err
    assert [d for d in rep["decisions"] if d["target_field_path"] == FORM and d["run_id"] == "run-001"]
    assert not rep["replaces_or_contradicts_a_file_reading"]


# --- L1: a row names a run by its immediate folder, and a row reaching several runs is reported ---


def _fake_bundle(paths):
    return {"sources": [{"source_id": f"s{i}", "relative_path": p} for i, p in enumerate(paths)],
            "runs": [{"run_id": f"run-{i:03d}", "acquisition_source_id": f"s{i}", "stem": p.rsplit("/", 1)[-1].rsplit(".", 1)[0]}
                     for i, p in enumerate(paths)]}


def test_a_row_matches_a_runs_immediate_folder_or_its_stem_never_an_ancestor_folder():
    b = _fake_bundle(["raw/day1/ZZ_a_100mV", "raw/day1/ZZ_a_100mV_dir/ZZ_a_100mV_calib_1.txt", "raw/day2/ZZ_b_200mV.dat"])
    assert T._match_runs(b, "raw") == []
    assert T._match_runs(b, "ZZ_a_100mV_dir") == ["run-000", "run-001"]  # the scan beside its folder, and the file in it
    assert T._match_runs(b, "day1") == ["run-000"]
    assert T._match_runs(b, "ZZ_b_200mV") == ["run-002"]
    assert T._match_runs(b, "ZZ_b_200mV.dat") == ["run-002"]
    assert T._match_runs(b, "raw/day2/ZZ_b_200mV.dat") == ["run-002"]


@pytest.fixture(scope="module")
def nested(tmp_path_factory):
    """The synthetic corpus one folder deeper (``raw/day1/``), so an ancestor folder can be named."""
    src = tmp_path_factory.mktemp("nested-src") / "SYN_BL99_nested"
    shutil.copytree(SEMANTICS, src / "raw" / "day1")
    bdir = tmp_path_factory.mktemp("nested") / "b"
    assert cli.main(["import", str(src), "--out", str(bdir)]) == 0
    M.group(bdir)
    return bdir


def test_rows_naming_an_ancestor_folder_reach_no_run_and_rows_reaching_several_runs_are_counted(nested, tmp_path, capsys):
    stems = _stems(nested)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Campaign Info"
    ws["A1"], ws["A3"], ws["B3"] = "ISAAC Campaign Metadata", "Facility Name", "Synthetic Light Source"
    wb.create_sheet("Configurations").append(["Config ID", "Technique"])
    wf = wb.create_sheet("File List")
    wf.append(["Folder Name", "Sample Form"])
    wf.append(["raw", "powder"])
    wf.append(["SYN_BL99_nested", "powder"])
    wf.append(["day1", "film"])
    wf.append([f"{stems['run-001']}_dir", "powder"])
    wb.save(tmp_path / "theirs.xlsx")
    code, text, err = _adopt(nested, tmp_path / "theirs.xlsx", tmp_path / "out.xlsx", capsys, "--json")
    assert code == 0, err
    a = json.loads(text)["adopted"]
    assert {r["name"] for r in a["unmatched_rows"]} == {"raw", "SYN_BL99_nested"}
    several = {x["name"]: x for x in a["rows_reaching_several_runs"]}
    assert set(several) == {"day1"}
    assert several["day1"]["runs"] == len(several["day1"]["run_ids"]) > 1 and several["day1"]["row"] == "'File List'!A4"

    code, text, err = _adopt(nested, tmp_path / "theirs.xlsx", tmp_path / "out2.xlsx", capsys)
    assert code == 0, err
    line = next(ln for ln in text.splitlines() if ln.startswith("Your rows that reach more than one run"))
    assert "(1)" in line and f"day1 -> {several['day1']['runs']} runs" in line


# --- L3: ISAAC's blank template with example rows is refused ---------------------------------------


def _blank_copy(path, file_rows):
    """ISAAC's blank template as it ships: its example values, plus (hypothetically) example File List rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Campaign Info"
    ws["A1"] = "ISAAC Campaign Metadata"
    for r, (label, value) in enumerate(T.BLANK_TEMPLATE_CAMPAIGN_EXAMPLES.items(), start=3):
        ws[f"A{r}"], ws[f"B{r}"] = label, value
    ws["A10"] = "Instructions:"
    wc = wb.create_sheet("Configurations")
    headers = list(T.BLANK_TEMPLATE_CONFIG_EXAMPLES[0])
    wc.append(headers + ["Notes"])
    for row in T.BLANK_TEMPLATE_CONFIG_EXAMPLES:
        wc.append([row.get(h) for h in headers])
    wf = wb.create_sheet("File List")
    wf.append(["File Name", "Config ID", "Sample Name", "Sample Form"])
    for row in file_rows:
        wf.append(row)
    wb.save(path)
    return path


def test_the_blank_template_with_example_file_list_rows_is_refused_and_writes_nothing(grouped, tmp_path, capsys):
    stems = _stems(grouped)
    blank = _blank_copy(tmp_path / "ISAAC_Metadata_Template.xlsx",
                        [[f"{stems['run-001']}_dir", "Config_A", "Example sample", "powder"],
                         [f"{stems['run-002']}_dir", "Config_B", "Example sample", "film"]])
    out = tmp_path / "adopted.xlsx"
    code, _, err = _adopt(grouped, blank, out, capsys)
    assert code == 3 and "blank metadata template" in err and "example values" in err and not out.exists()
    assert not M.load(grouped)["scientist_decisions"]

    # their own values in Campaign Info (a real beamtime's): a filled copy, adopted
    wb = openpyxl.load_workbook(blank)
    wb["Campaign Info"]["B5"] = "62b"
    wb.save(tmp_path / "filled.xlsx")
    code, _, err = _adopt(grouped, tmp_path / "filled.xlsx", out, capsys)
    assert code == 0, err and out.exists()


def test_the_blank_templates_start_date_read_back_as_a_date_is_still_the_blank(grouped, tmp_path, capsys):
    from datetime import datetime

    stems = _stems(grouped)
    blank = _blank_copy(tmp_path / "blank.xlsx", [[f"{stems['run-001']}_dir", "Config_A", "Example sample", "powder"]])
    wb = openpyxl.load_workbook(blank)
    wb["Campaign Info"]["B6"] = datetime.fromisoformat(T.BLANK_TEMPLATE_CAMPAIGN_EXAMPLES["Start Date"])
    wb.save(blank)
    code, _, err = _adopt(grouped, blank, tmp_path / "adopted.xlsx", capsys)
    assert code == 3 and "blank metadata template" in err


def test_adopt_json_reports_a_date_cell_of_theirs_as_text(grouped, tmp_path, capsys):
    """A filled sheet's Start Date is a real Excel date: ``adopt --json`` printed nothing and crashed."""
    from datetime import datetime

    theirs = _their_sheet(grouped, tmp_path / THEIRS)
    wb = openpyxl.load_workbook(theirs)
    wb["Campaign Info"]["B6"] = datetime(2099, 5, 16)
    wb.save(theirs)
    a = _adopt_json(grouped, theirs, tmp_path / "adopted.xlsx", capsys)
    start = next(x for x in a["not_adopted"] if x["header"] == "Start Date")
    assert start["value"] == "2099-05-16"


def test_import_json_shows_a_date_typed_in_a_text_cell_as_an_invalid_cell(grouped, tmp_path, capsys):
    from datetime import datetime

    xlsx = tmp_path / "t.xlsx"
    assert cli.main(["bundle", "template", "export", str(grouped), "--out", str(xlsx)]) == 0
    wb = openpyxl.load_workbook(xlsx)
    wb["File List"][f"{_col('Sample Name')}{_row_of(xlsx, 'run-001')}"] = datetime(2099, 5, 16)
    wb.save(xlsx)
    code, rep, err = _import(grouped, xlsx, capsys)
    assert code == 0, err
    assert [x for x in rep["invalid_cells"] if x["value"] == "2099-05-16" and "date" in x["reason"]]
