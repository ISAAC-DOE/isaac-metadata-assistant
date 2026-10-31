"""Polish c7-s: the scientist's own filled sheet, after the real-data walkthrough. Synthetic only.

* H1: a Campaign Info beamline (or facility acronym) that names another one than a facility HINT (the folder's
  name, ``SYN_BL152_beamtime``, or the acquisition system's) is never adopted: it is a Disagreements row (kind
  ``hint``) with the hint and its source. ``BL15-2``, ``BL152`` and ``15-2`` are one beamline. The import preview names a typed one once (recorded:
  the scientist typed it) and holds back an adopted one.
* M1: ``decide --question <Q> --only group:<id>`` on a question asked once for the whole beamtime records that
  group's answer per run (a per-run answer slot where the run has no target of its own); a ruling records for
  exactly the runs its row names the same way; every reply on a beamtime-wide question says what it reaches.
* M2: a Disagreements row lists every reading, the agreeing ones too, each marked.
* M3: an Electrolyte row that also states a concentration takes the concentration in its own ruling column.
* M4: a column adopt does not read whose value is (loosely) one of the field's allowed values is a hint on the row.
* L1-L3: runs vs cells; one cell named once; a reaction says what it makes required.
"""

from __future__ import annotations

import shutil

import openpyxl
import pytest

from isaac_api.local_import import bundle as B
from isaac_api.local_import import cli
from isaac_api.local_import import mutations as M
from isaac_api.local_import import questions as Q
from isaac_api.local_import import records as R
from isaac_api.local_import import template as T

from tests_support_records import SEMANTICS
from test_local_import_polish_c3r2 import _annotate
from test_local_import_template_adopt import (  # noqa: F401  (fixtures)
    BEAMLINE, POT, THEIRS, WHO, _adopt, _adopt_json, _col, _d, _import, _row_of, _their_sheet, corpus, grouped,
    grouped_base)

TECH = "system.technique"
CONC = "context.electrochemistry.electrolyte.concentration_M"
NAME = "context.electrochemistry.electrolyte.name"
RUNS = ["run-001", "run-002", "run-003", "run-004", "run-005", "run-006", "run-007", "run-008"]


@pytest.fixture(scope="module")
def plain_base(tmp_path_factory):
    """The synthetic beamtime with NO AI reading, in a folder named for beamline 15-2 (``SYN_BL152_beamtime``): the
    facility hints are the folder's name and the acquisition system its scan headers match (``SSRL BL15-2 …``)."""
    src = tmp_path_factory.mktemp("c7s-src") / "SYN_BL152_beamtime"
    shutil.copytree(SEMANTICS, src)
    bdir = tmp_path_factory.mktemp("c7s") / "b"
    assert cli.main(["import", str(src), "--out", str(bdir)]) == 0
    M.group(bdir)
    return bdir


@pytest.fixture()
def plain(plain_base, tmp_path):
    bdir = tmp_path / "plain"
    shutil.copytree(plain_base, bdir)
    return bdir


def _sheet(bdir, path, *, beamline="62b", facility=None, technique=None, extra=None):
    """``_their_sheet`` with a few cells changed: ``extra`` = ``{header: {run_id: value}}`` columns appended."""
    _their_sheet(bdir, path, beamline=beamline)
    wb = openpyxl.load_workbook(path)
    if facility is not None:
        wb["Campaign Info"]["B3"] = facility
    if technique is not None:
        wb["Configurations"]["B2"] = technique
    if extra:
        wf = wb["File List"]
        stems = {r["run_id"]: r["stem"] for r in M.load(bdir)["runs"]}
        row_of = {f"{stems[r]}_dir": r for r in stems}
        for header, by_run in extra.items():
            col = wf.max_column + 1
            wf.cell(row=1, column=col, value=header)
            for r in range(2, wf.max_row + 1):
                rid = row_of.get(wf.cell(row=r, column=1).value)
                if rid in by_run:
                    wf.cell(row=r, column=col, value=by_run[rid])
    wb.save(path)
    return path


def _rule(adopted, out, rulings):
    """Type rulings into the Disagreements sheet: ``{(field, their value): {column: value}}``."""
    wb = openpyxl.load_workbook(adopted)
    wd = wb["Disagreements"]
    for r in range(2, wd.max_row + 1):
        for column, value in rulings.get((wd[f"C{r}"].value, wd[f"D{r}"].value), {}).items():
            wd[f"{column}{r}"] = value
    wb.save(out)
    return out


def _latest(bdir):
    return B._decisions_by_target(M.load(bdir)["scientist_decisions"])


def _technique_question(bdir):
    b = M.load(bdir)
    q = next(q for q in Q.consolidate(b)["questions"] if q.get("target_field_path") == TECH
             and list(q["per_scope_targets"]) == ["experiment"])
    return q["question_id"]


def _without_own_target(bdir, path, runs):
    """The runs of ``runs`` with no candidate or missing value of their own for ``path`` (asked once for the beamtime)."""
    b = M.load(bdir)
    own = {o["run_id"] for o in [*b["candidates"], *b["missing_values"]] if o["target_field_path"] == path}
    return [r for r in runs if r not in own]


def _decide(capsys, *args):
    capsys.readouterr()
    code = cli.main(["bundle", "decide", *map(str, args), "--by", WHO])
    out = capsys.readouterr()
    return code, out.out, out.err


# --- H1: the beamline against the folder's name -----------------------------------------------------


def test_a_beamline_naming_another_one_than_the_folder_hint_is_a_disagreement_never_adopted(plain, tmp_path, capsys):
    out = tmp_path / "adopted.xlsx"
    a = _adopt_json(plain, _sheet(plain, tmp_path / THEIRS), out, capsys)
    bl = _d(a, "Beamline/Instrument", kind="hint")
    assert bl["their_value"] == "62b" and bl["scope"] == "experiment"
    assert any("the imported folder's name 'SYN_BL152_beamtime'" in r and "'BL152'" in r and "a hint, never an answer" in r
               for r in bl["readings"])
    assert any("acquisition system" in r and "'BL15-2'" in r for r in bl["readings"])
    wb = openpyxl.load_workbook(out)
    assert wb["Campaign Info"]["B5"].value is None
    wd = wb["Disagreements"]
    row = next(r for r in range(2, wd.max_row + 1) if wd[f"A{r}"].value == bl["id"])
    assert "a hint" in wd[f"B{row}"].value and "BL152" in wd[f"F{row}"].value and wd[f"H{row}"].value is None
    # their multi-word facility name is not compared (a folder names an acronym), and is adopted
    assert wb["Campaign Info"]["B3"].value == "Synthetic Light Source"
    code, text, _ = _adopt(plain, tmp_path / THEIRS, tmp_path / "again.xlsx", capsys)
    assert code == 0 and "Disagreements with a hint (1;" in text and "yours '62b'" in text


@pytest.mark.parametrize("spelled", ["BL15-2", "BL152", "15-2", "bl 15.2"])
def test_the_same_beamline_in_another_spelling_is_adopted(plain, tmp_path, capsys, spelled):
    out = tmp_path / "adopted.xlsx"
    a = _adopt_json(plain, _sheet(plain, tmp_path / THEIRS, beamline=spelled), out, capsys)
    assert not [d for d in a["disagreements"] if d["header"] == "Beamline/Instrument"]
    assert openpyxl.load_workbook(out)["Campaign Info"]["B5"].value == spelled


def test_a_facility_acronym_naming_another_facility_than_the_folder_is_a_disagreement(plain, tmp_path, capsys):
    a = _adopt_json(plain, _sheet(plain, tmp_path / THEIRS, beamline="BL15-2", facility="ZZLS"), tmp_path / "a.xlsx", capsys)
    fac = _d(a, "Facility Name", kind="hint")
    assert fac["their_value"] == "ZZLS" and any("names facility 'SYN'" in r for r in fac["readings"])
    assert T.same_beamline("BL15-2", "BL152") and T.same_beamline("15-2", "BL152") and T.same_beamline("BL15-2", "15-2")
    assert T.beamline_form("62b") == "62b" and not T.same_beamline("62b", "BL152")


def test_a_typed_beamline_naming_another_one_is_named_once_in_the_preview_and_recorded(plain, tmp_path, capsys):
    xlsx = tmp_path / "t.xlsx"
    assert cli.main(["bundle", "template", "export", str(plain), "--out", str(xlsx)]) == 0
    wb = openpyxl.load_workbook(xlsx)
    wb["Campaign Info"]["B5"] = "62b"
    wb.save(tmp_path / "typed.xlsx")
    capsys.readouterr()
    assert cli.main(["bundle", "template", "import", str(plain), str(tmp_path / "typed.xlsx"), "--by", WHO]) == 0
    text = capsys.readouterr().out
    assert len([ln for ln in text.splitlines() if "'Campaign Info'!B5" in ln and "BL152" in ln]) == 1
    code, rep, err = _import(plain, tmp_path / "typed.xlsx", capsys, "--apply")
    assert code == 0, err
    [named] = [x for x in rep["replaces_or_contradicts_a_file_reading"] if x["target_field_path"] == BEAMLINE]
    assert named["outcome"] == "recorded_as_typed" and named["run_id"] is None
    assert any("BL152" in r for r in named["readings"])
    assert {d["value"] for d in rep["decisions"] if d["target_field_path"] == BEAMLINE} == {"62b"}  # named, not blocked


def test_an_adopted_beamline_a_hint_contradicts_is_held_back_at_import(plain, tmp_path, capsys):
    # an adopted file whose beamline cell was copied before this check (here: adopted as BL15-2, then changed)
    out = tmp_path / "adopted.xlsx"
    _adopt_json(plain, _sheet(plain, tmp_path / THEIRS, beamline="BL15-2"), out, capsys)
    wb = openpyxl.load_workbook(out)
    assert wb["Campaign Info"]["B5"].value == "BL15-2"
    wb["Campaign Info"]["B5"] = "62b"
    wb.save(tmp_path / "older.xlsx")
    code, rep, err = _import(plain, tmp_path / "older.xlsx", capsys, "--apply")
    assert code == 0, err
    assert not [d for d in rep["decisions"] if d["target_field_path"] == BEAMLINE]
    [waiting] = [x for x in rep["awaiting_ruling"] if x["target_field_path"] == BEAMLINE]
    assert waiting["their_value"] == "62b" and "template adopt" in waiting["why"]
    assert not [d for d in M.load(plain)["scientist_decisions"] if d.get("value") == "62b"]


# --- M2: every reading on the row, the agreeing ones too -----------------------------------------------


def test_a_disagreement_lists_the_readings_that_agree_with_their_value_too(grouped, tmp_path, capsys):
    rel = _annotate(grouped, "run-005", POT, 0.8)
    a = _adopt_json(grouped, _their_sheet(grouped, tmp_path / THEIRS), tmp_path / "a.xlsx", capsys)
    row = _d(a, "Potential (V)", run_id="run-005")
    assert any(r.startswith("differs: -0.8") and "m800mV" in r for r in row["readings"])
    assert any(r.startswith("agrees: 0.8") and rel in r for r in row["readings"])
    wd = openpyxl.load_workbook(tmp_path / "a.xlsx")["Disagreements"]
    f = next(wd[f"F{r}"].value for r in range(2, wd.max_row + 1) if wd[f"A{r}"].value == row["id"])
    assert "agrees: 0.8" in f and "differs: -0.8" in f


# --- M3: the electrolyte's concentration has its own ruling cell --------------------------------------


def test_an_electrolyte_rows_concentration_ruling_is_recorded_for_exactly_its_runs(grouped, tmp_path, capsys):
    out = tmp_path / "adopted.xlsx"
    a = _adopt_json(grouped, _their_sheet(grouped, tmp_path / THEIRS), out, capsys)
    name = _d(a, "Electrolyte", kind="name_only")
    assert name["concentration_ruling_cell"].endswith(f"!{T.DISAGREE_CONC_COLUMN}{int(name['ruling_cell'].split('H')[-1])}")
    wd = openpyxl.load_workbook(out)["Disagreements"]
    assert wd[f"{T.DISAGREE_CONC_COLUMN}1"].value == "Your ruling: concentration (M)"
    ruled = _rule(out, tmp_path / "ruled.xlsx", {("Electrolyte", "0.1 M KOH"): {"H": "KOH", "J": 0.1},
                                                ("Environment", "aqueous_electrochemistry"): {"J": 0.5}})
    code, rep, err = _import(grouped, ruled, capsys, "--apply")
    assert code == 0, err
    def runs_of(path, value):
        return {d["run_id"] for d in rep["decisions"] if d["target_field_path"] == path and d["value"] == value}

    assert runs_of(CONC, 0.1) == set(name["runs"]) == set(RUNS)  # run-009 is on none of their rows
    assert runs_of(NAME, "KOH") == set(RUNS)
    notes = [d["note"] for d in M.load(grouped)["scientist_decisions"] if "concentration ruling" in (d.get("note") or "")]
    assert len(notes) == len(RUNS) and all(name["id"] in n for n in notes)
    # a concentration typed on a row that is not an electrolyte's is recorded nowhere, and said
    [nowhere] = [x for x in rep["recorded_nowhere"] if x["header"] == "Your ruling: concentration (M)"]
    assert "only an Electrolyte row" in nowhere["reason"]


def test_a_workbook_adopted_before_the_concentration_column_imports_as_before(grouped, tmp_path, capsys):
    out = tmp_path / "adopted.xlsx"
    _adopt_json(grouped, _their_sheet(grouped, tmp_path / THEIRS), out, capsys)
    ruled = _rule(out, tmp_path / "ruled.xlsx", {("Electrolyte", "0.1 M KOH"): {"H": "KOH", "J": 0.1}})
    wb = openpyxl.load_workbook(ruled)
    wb["Disagreements"][f"{T.DISAGREE_CONC_COLUMN}1"] = None  # nine columns, as adopt wrote them before
    wb.save(tmp_path / "older.xlsx")
    code, rep, err = _import(grouped, tmp_path / "older.xlsx", capsys)
    assert code == 0, err
    assert not [d for d in rep["decisions"] if d["target_field_path"] == CONC]
    assert {d["value"] for d in rep["decisions"] if d["target_field_path"] == NAME} == {"KOH"}


# --- M4: an unread column that holds an allowed value is a hint on the row ------------------------------


def test_an_unread_column_holding_an_allowed_value_is_a_hint_on_the_row_never_adopted(grouped, tmp_path, capsys):
    mode = {r: "insitu" for r in RUNS[:6]} | {"run-007": "nonsense"}
    out = tmp_path / "adopted.xlsx"
    a = _adopt_json(grouped, _sheet(grouped, tmp_path / THEIRS, extra={"Environment mode": mode}), out, capsys)
    env = _d(a, "Environment", kind="not_allowed")
    [hint] = env["unread_hints"]
    assert (hint["header"], hint["value"], hint["allowed_value"], hint["rows"]) == ("Environment mode", "insitu", "in_situ", 6)
    assert "Environment mode" in {c["header"] for c in a["their_columns_not_read"]}
    wb = openpyxl.load_workbook(out)
    wd = wb["Disagreements"]
    b_cell = next(wd[f"B{r}"].value for r in range(2, wd.max_row + 1) if wd[f"A{r}"].value == env["id"])
    assert f"column {hint['column']} ('Environment mode') says 'insitu' (= `in_situ`) on 6 of these 8 run(s)" in b_cell
    assert all(wb["File List"][f"{_col('Environment')}{_row_of(out, r)}"].value is None for r in RUNS)
    code, rep, err = _import(grouped, out, capsys, "--apply")
    assert code == 0, err
    assert not [d for d in M.load(grouped)["scientist_decisions"] if d.get("value") == "in_situ"]
    code, text, _ = _adopt(grouped, tmp_path / THEIRS, tmp_path / "again.xlsx", capsys)
    assert "hint: your sheet's column" in text and "(= `in_situ`)" in text


# --- M1: one group of a question asked once for the beamtime; rulings; the reach ------------------------


def test_only_one_group_of_a_beamtime_wide_question_records_per_run_for_that_group(plain, capsys):
    qid = _technique_question(plain)
    before = {q["question_id"] for q in Q.consolidate(M.load(plain))["questions"]}
    b = M.load(plain)
    run_of = {r["run_key"]: r["run_id"] for r in b["runs"]}
    g0 = Q.live_groups(b)[0]
    gid, members = g0["group_id"], [run_of[k] for k in g0["member_run_keys"]]
    code, out, err = _decide(capsys, plain, "--question", qid, "--only", f"group:{gid}", "--decision", "correct", "--value", "XAS")
    assert code == 0, err
    assert f"It reaches group:{gid} only" in out and "covers the other 8 groups" in out
    b = M.load(plain)
    assert B.validate_bundle(b) == []
    slots = [m for m in b["missing_values"] if m["target_field_path"] == TECH]
    assert sorted(m["run_id"] for m in slots) == sorted(members)
    assert all(m["target_write_scope"] == "run" and qid in m["reason"] for m in slots)
    latest = _latest(plain)
    assert all(latest[("missing_value", m["value_id"])]["value"] == "XAS" for m in slots)
    beamtime = next(c for c in b["candidates"] if c["run_id"] is None and c["target_field_path"] == TECH)
    assert B.candidate_decision(b, beamtime, latest) is None  # the beamtime-wide target is untouched
    # the question keeps its id and still asks the rest; no question id changed
    assert Q.question(Q.consolidate(b), qid) is not None
    assert before <= {q["question_id"] for q in Q.consolidate(b)["questions"]} | {qid}
    # the group's record takes it; another group's does not
    index = R.composition_index(b)
    techs = {g["group_id"]: R.compose_group(b, g, plain, index)[0]["fields"].get(TECH) for g in Q.live_groups(b)}
    assert techs[gid]["value"] == "XAS" and sum(1 for v in techs.values() if v and v.get("value")) == 1
    # a repeat records nothing new and adds no slot
    code, out, err = _decide(capsys, plain, "--question", qid, "--only", f"group:{gid}", "--decision", "correct", "--value", "XAS")
    assert code == 0 and "Nothing new recorded" in out
    assert len([m for m in M.load(plain)["missing_values"] if m["target_field_path"] == TECH]) == len(slots)


def test_only_a_group_refuses_reject_or_defer_and_an_unknown_group(plain, capsys):
    qid = _technique_question(plain)
    n = len(M.load(plain)["missing_values"])
    gid = Q.live_groups(M.load(plain))[0]["group_id"]
    code, _, err = _decide(capsys, plain, "--question", qid, "--only", f"group:{gid}", "--decision", "defer")
    assert code == 2 and "once for the whole beamtime" in err and "Nothing was recorded" in err
    code, _, err = _decide(capsys, plain, "--question", qid, "--only", "group:rg-nope", "--decision", "correct", "--value", "XAS")
    assert code == 6
    code, _, err = _decide(capsys, plain, "--question", qid, "--only", f"group:{gid}", "--decision", "correct",
                           "--value", "not_a_technique")
    assert code != 0  # the same value check as any decide
    b = M.load(plain)
    assert len(b["missing_values"]) == n and not b.get("scientist_decisions")


def test_a_reply_on_a_beamtime_wide_question_says_how_many_groups_and_runs_it_reaches(plain, capsys):
    qid = _technique_question(plain)
    code, out, err = _decide(capsys, plain, "--question", qid, "--decision", "correct", "--value", "HERFD-XAS")
    assert code == 0, err
    assert "It reaches every record group of the beamtime: 9 groups, 9 runs." in out


def test_a_technique_ruling_records_for_exactly_the_runs_its_row_names(plain, tmp_path, capsys):
    out = tmp_path / "adopted.xlsx"
    a = _adopt_json(plain, _sheet(plain, tmp_path / THEIRS, beamline="BL15-2", technique="XANES"), out, capsys)
    tech = _d(a, "Technique", kind="not_allowed")
    assert set(tech["runs"]) == set(RUNS)
    ruled = _rule(out, tmp_path / "ruled.xlsx", {("Technique", "XANES"): {"H": "HERFD-XAS"}})
    n = len(M.load(plain)["missing_values"])
    capsys.readouterr()
    assert cli.main(["bundle", "template", "import", str(plain), str(ruled), "--by", WHO]) == 0
    text = capsys.readouterr().out
    code, rep, err = _import(plain, ruled, capsys)
    assert code == 0, err
    assert not rep["asked_beamtime_wide"]
    [per_run] = rep["ruled_per_run"]
    slotted = _without_own_target(plain, TECH, RUNS)
    assert slotted and len(slotted) < len(RUNS)  # some runs have a Technique target of their own, some do not
    assert per_run["runs"] == per_run["slots_added"] == len(slotted) and per_run["groups_left_to_the_question"] == 1
    assert len(M.load(plain)["missing_values"]) == n  # a preview adds nothing to the bundle
    code, rep, err = _import(plain, ruled, capsys, "--apply")
    assert code == 0, err
    b = M.load(plain)
    assert B.validate_bundle(b) == []
    latest = _latest(plain)
    techs = {m["run_id"]: latest[("missing_value", m["value_id"])]["value"] for m in b["missing_values"] if m["target_field_path"] == TECH}
    assert techs == {r: "HERFD-XAS" for r in slotted}
    assert {d["run_id"] for d in rep["decisions"] if d["target_field_path"] == TECH and d["value"] == "HERFD-XAS"} == set(RUNS)
    index = R.composition_index(b)
    composed = {rid: (R.compose_group(b, g, plain, index)[0]["fields"].get(TECH) or {}).get("value")
                for g in Q.live_groups(b) for rid in {r["run_id"] for r in b["runs"] if r["run_key"] in g["member_run_keys"]}}
    assert composed == {**{r: "HERFD-XAS" for r in RUNS}, "run-009": None}  # run-009 is on none of their rows
    assert f"recorded per run for {len(slotted)} run(s)" in text and "Asked once for the whole beamtime" not in text


def test_a_refused_technique_ruling_leaves_no_answer_slot_behind(plain, tmp_path, capsys):
    out = tmp_path / "adopted.xlsx"
    _adopt_json(plain, _sheet(plain, tmp_path / THEIRS, beamline="BL15-2", technique="XANES"), out, capsys)
    n = len(M.load(plain)["missing_values"])
    ruled = _rule(out, tmp_path / "ruled.xlsx", {("Technique", "XANES"): {"H": "XANES"},
                                                ("Electrolyte", "0.1 M KOH"): {"H": "KOH"}})
    code, rep, err = _import(plain, ruled, capsys, "--apply")
    assert code == 0, err
    assert [x for x in rep["invalid_cells"] if x["header"] == "Technique"]
    assert len(M.load(plain)["missing_values"]) == n


# --- L1-L3 ------------------------------------------------------------------------------------------------


def test_rows_say_runs_and_cells_apart_and_a_cell_reached_by_many_rows_is_named_once(plain, tmp_path, capsys):
    code, text, err = _adopt(plain, _sheet(plain, tmp_path / THEIRS, beamline="BL15-2"), tmp_path / "a.xlsx", capsys)
    assert code == 0, err
    assert "Environment 'aqueous_electrochemistry' for 8 run(s), from 8 of your cell(s)" in text
    wd = openpyxl.load_workbook(tmp_path / "a.xlsx")["Disagreements"]
    runs_cell = next(wd[f"G{r}"].value for r in range(2, wd.max_row + 1) if wd[f"C{r}"].value == "Environment")
    assert runs_cell.endswith("(8 runs)")
    # an ordinary template: one Technique cell reached through every row is named once, with its runs
    xlsx = tmp_path / "t.xlsx"
    assert cli.main(["bundle", "template", "export", str(plain), "--out", str(xlsx)]) == 0
    wb = openpyxl.load_workbook(xlsx)
    wb["Configurations"]["A2"], wb["Configurations"]["B2"] = "C1", "XAS"
    for run_id in [r["run_id"] for r in T.live_rows(M.load(plain))]:
        wb["File List"][f"B{_row_of(xlsx, run_id)}"] = "C1"
    wb.save(tmp_path / "typed.xlsx")
    capsys.readouterr()
    assert cli.main(["bundle", "template", "import", str(plain), str(tmp_path / "typed.xlsx"), "--by", WHO]) == 0
    text = capsys.readouterr().out
    [line] = [ln for ln in text.splitlines() if ln.startswith("  Technique -> system.technique")]
    n = len(_without_own_target(plain, TECH, [r["run_id"] for r in M.load(plain)["runs"]]))
    assert f"{n} run(s) from 1 cell(s) ('Configurations'!B2)" in line and line.count("'Configurations'!B2") == 1


def test_a_reaction_says_it_makes_drive_and_catalysis_required_with_the_groups(plain, tmp_path, capsys):
    reaction = {r: "oxidation" for r in RUNS}
    theirs = _sheet(plain, tmp_path / THEIRS, beamline="BL15-2", extra={"Reaction": reaction})
    code, text, err = _adopt(plain, theirs, tmp_path / "a.xlsx", capsys)
    assert code == 0, err
    assert "makes context.reaction.catalysis and context.reaction.drive required for 8 group(s)" in text
    ruled = _rule(tmp_path / "a.xlsx", tmp_path / "ruled.xlsx", {("Reaction", "oxidation"): {"H": "surface_oxidation"}})
    code, rep, err = _import(plain, ruled, capsys)
    assert code == 0, err
    assert rep["reaction_needs"] == {"fields": sorted(B.REACTION_REQUIRED_PATHS), "runs": 8, "groups": 8}
    capsys.readouterr()
    assert cli.main(["bundle", "template", "import", str(plain), str(ruled), "--by", WHO]) == 0
    assert "Then required: A reaction recorded by this import makes" in capsys.readouterr().out
