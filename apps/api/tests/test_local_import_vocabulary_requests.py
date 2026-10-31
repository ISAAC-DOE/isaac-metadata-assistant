"""A scientist's word the list lacks is kept as a VOCABULARY REQUEST (bundle 1.10.0), and drafted. Synthetic only.

Every test runs over a bundle imported here from the committed SYNTHETIC multi-operator fixture.

* ``decide`` (by target or ``--question``) refuses an answer off its field's list and records no decision and
  no value. Only with ``--as-said`` (the scientist's own word) does it keep the word as a vocabulary request:
  who (``--by``), when, the field, the targets and runs, and where it was said. Without ``--as-said`` it keeps
  NOTHING and says how to keep it (c10-d H1). Any other refusal (a word in a number field) keeps nothing.
* ``template import --apply`` keeps an off-list cell the same way (a preview keeps nothing).
* ``bundle propose-terms`` drafts the requests, with the scientist's own words (``scientist_notes``),
  ``runs``/``run_ids`` and citations naming a file (no folders) and run ids. A word that names an
  existing term (``reversible_hydrogen_electrode``) is never drafted: "use `RHE` (an existing term)".
  A close term (``XANES``) is drafted with ``check_existing``. An AI reading its scientist already
  answered with an allowed value is not counted. ``--merge`` updates a drafts file, keeps a description
  in it (who wrote it is not known), reads no note from it, and never marks a draft ready to send.
* ``validate`` names a missing value that waits on such a request (``pending_vocabulary``).
* A 1.9.0 bundle still loads and validates, and its first save says what 1.10.0 added.
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import openpyxl
import pytest

from isaac_api.local_import import bundle as B
from isaac_api.local_import import cli
from isaac_api.local_import import mutations as M
from isaac_api.local_import import propose_terms as P
from isaac_api.local_import import template as T

from tests_support_records import pristine_bundle  # noqa: F401

WHO = "Synthetic Scientist"
REACTION = "context.electrochemistry.reaction"
SCALE = "context.electrochemistry.potential_scale"
ENV = "context.environment"
CELL = "context.electrochemistry.cell_type"
PH = "context.electrochemistry.pH"
TECH = "system.technique"
KEPT = "kept as a vocabulary request; `isaac bundle propose-terms` can draft it"


@pytest.fixture()
def b(pristine_bundle, tmp_path):
    copy = tmp_path / "bundle"
    shutil.copytree(pristine_bundle, copy)
    return copy


@pytest.fixture(scope="module")
def grouped_base(pristine_bundle, tmp_path_factory):
    bdir = tmp_path_factory.mktemp("vocab-grouped") / "b"
    shutil.copytree(pristine_bundle, bdir)
    M.group(bdir)
    return bdir


@pytest.fixture()
def grouped(grouped_base, tmp_path):
    bdir = tmp_path / "g"
    shutil.copytree(grouped_base, bdir)
    return bdir


def _miss(bdir: Path, field: str, index: int = 0) -> dict:
    return [m for m in M.load(bdir)["missing_values"] if m["target_field_path"] == field and m["run_id"]][index]


def _decide(bdir: Path, capsys, *args: str) -> tuple[int, str]:
    capsys.readouterr()
    code = cli.main(["bundle", "decide", str(bdir), *args, "--by", WHO])
    return code, capsys.readouterr().err


def _said(bdir: Path, capsys, *args: str) -> tuple[int, str]:
    """``decide`` with ``--as-said``: the value is the scientist's own word."""
    return _decide(bdir, capsys, *args, "--as-said")


def _report(bdir: Path, capsys, *extra: str) -> dict:
    capsys.readouterr()
    assert cli.main(["bundle", "propose-terms", str(bdir), *extra]) == cli.EXIT_OK
    return json.loads(capsys.readouterr().out)


def _term(report: dict, field: str, term: str) -> dict | None:
    return next((t for f in report["fields"] if f["field_path"] == field for t in f["terms"] if t["term"] == term), None)


# --- decide ---------------------------------------------------------------------------------------


def test_an_off_list_answer_is_kept_as_a_vocabulary_request_never_a_decision(b, capsys):
    target = _miss(b, REACTION)
    decisions_before = len(M.load(b)["scientist_decisions"])
    code, err = _said(b, capsys, "--target", target["value_id"], "--decision", "correct", "--value", "cathodic corrosion",
                        "--note", "the Pt dissolves under cathodic bias")
    assert code == 2
    assert "not one of the schema's allowed values" in err and "Nothing was recorded" in err and KEPT in err
    held = M.load(b)
    assert len(held["scientist_decisions"]) == decisions_before  # never a decision
    assert next(m for m in held["missing_values"] if m["value_id"] == target["value_id"])["value_state"] == B.STATE_MISSING
    [req] = held["vocabulary_requests"]
    assert req["request_id"] == "vr-00001" and req["field_path"] == REACTION and req["term"] == "cathodic corrosion"
    assert req["requested_by"]["name"] == WHO and req["requested_by"]["identity"] == "local_unverified"
    assert req["targets"] == [{"kind": "missing_value", "id": target["value_id"]}]
    assert req["run_ids"] == [target["run_id"]] and req["scope"] == "run"
    assert req["source"] == {"kind": "decide_answer", "command": "bundle decide", "decision": "correct", "question_id": None}
    assert req["note"] == "the Pt dissolves under cathodic bias"
    assert B.check_bundle(b) == []
    # the same answer again keeps nothing new and writes nothing
    before = (b / "bundle.json").read_bytes()
    code, err = _said(b, capsys, "--target", target["value_id"], "--decision", "correct", "--value", "cathodic corrosion",
                        "--note", "the Pt dissolves under cathodic bias")
    assert code == 2 and "kept earlier" in err and KEPT in err
    assert (b / "bundle.json").read_bytes() == before


def test_without_as_said_an_off_list_answer_keeps_nothing_and_says_how_to_keep_it(b, capsys):
    """c10-d H1: a value tried while working out a mapping is never attributed to the scientist."""
    before = (b / "bundle.json").read_bytes()
    target = _miss(b, REACTION)
    code, err = _decide(b, capsys, "--target", target["value_id"], "--decision", "correct", "--value", "cathodic corrosion")
    assert code == 2 and "not one of the schema's allowed values" in err and "Nothing was recorded" in err
    assert "vocabulary request" in err and KEPT not in err and "kept as a vocabulary request" not in err
    assert "if this is the scientist's own word, repeat with --as-said" in err
    assert (b / "bundle.json").read_bytes() == before
    # by --question too
    from isaac_api.local_import import questions as Q

    q = next(x for x in Q.consolidate(M.load(b))["questions"] if x.get("target_field_path") == ENV)
    code, err = _decide(b, capsys, "--question", q["question_id"], "--decision", "correct", "--value", "aqueous_electrochemistry")
    assert code == 2 and "repeat with --as-said" in err and KEPT not in err
    assert (b / "bundle.json").read_bytes() == before
    # and through the library: the refusal is still an off-list one, keeping nothing
    with pytest.raises(M.OffListRefused) as got:
        M.decide(b, target_ids=[target["value_id"]], decision="correct", value="cathodic corrosion", by=WHO, orcid=None)
    assert M.AS_SAID_HINT in str(got.value)
    assert (b / "bundle.json").read_bytes() == before


def test_as_said_needs_a_value_and_changes_nothing_for_an_allowed_one(b, capsys):
    before = (b / "bundle.json").read_bytes()
    code, err = _decide(b, capsys, "--target", _miss(b, ENV)["value_id"], "--decision", "defer", "--as-said")
    assert code == 2 and "--as-said keeps the scientist's own word given with --value" in err
    assert (b / "bundle.json").read_bytes() == before
    # an allowed value with --as-said is recorded as ever and keeps no request
    assert _said(b, capsys, "--target", _miss(b, ENV)["value_id"], "--decision", "correct", "--value", "ex_situ")[0] == 0
    assert not M.load(b).get("vocabulary_requests")


def test_any_other_refusal_keeps_nothing(b, capsys):
    before = (b / "bundle.json").read_bytes()
    code, err = _said(b, capsys, "--target", _miss(b, PH)["value_id"], "--decision", "correct", "--value", "acidic")
    assert code == 2 and "must be number" in err and "vocabulary request" not in err
    assert (b / "bundle.json").read_bytes() == before
    # an allowed value is recorded as ever, and keeps no request
    assert _decide(b, capsys, "--target", _miss(b, ENV)["value_id"], "--decision", "correct", "--value", "ex_situ")[0] == 0
    assert not M.load(b).get("vocabulary_requests")


def test_a_question_answer_off_the_list_is_kept_with_its_question_and_every_run(b, capsys):
    from isaac_api.local_import import questions as Q

    q = next(x for x in Q.consolidate(M.load(b))["questions"] if x.get("target_field_path") == ENV)
    code, err = _said(b, capsys, "--question", q["question_id"], "--decision", "correct", "--value", "aqueous_electrochemistry")
    assert code == 2 and KEPT in err
    [req] = M.load(b)["vocabulary_requests"]
    assert req["source"]["question_id"] == q["question_id"] and req["term"] == "aqueous_electrochemistry"
    assert len(req["targets"]) == len(q["targets"]) and len(req["run_ids"]) == len(q["targets"])


def test_a_word_for_an_existing_term_says_use_it(b, capsys):
    code, err = _said(b, capsys, "--target", _miss(b, SCALE)["value_id"], "--decision", "correct",
                        "--value", "reversible_hydrogen_electrode")
    assert code == 2 and KEPT in err and "use `RHE` (an existing term)" in err
    entry = _term(_report(b, capsys), SCALE, "reversible_hydrogen_electrode")
    assert entry["draft"]["drafted"] is False and entry["draft"]["use_existing"] == "RHE"
    assert "Use `RHE` (an existing term)" in entry["draft"]["reason"]
    assert "equivalent_existing_term" in {p["code"] for p in entry["draft"]["problems"]}


# --- propose-terms reads the requests --------------------------------------------------------------


def test_propose_terms_drafts_a_request_with_the_scientists_words_runs_and_citations(b, capsys):
    first, second = _miss(b, CELL, 0), _miss(b, CELL, 1)
    for t in (first, second):
        assert _said(b, capsys, "--target", t["value_id"], "--decision", "correct",
                       "--value", "three_electrode_static_cell", "--note", "a static cell, no flow")[0] == 2
    report = _report(b, capsys)
    entry = _term(report, CELL, "three_electrode_static_cell")
    assert entry["count"] == 2 and entry["by_origin"][P.ORIGIN_REQUEST] == 2
    assert entry["runs"] == 2 and entry["run_ids"] == sorted({first["run_id"], second["run_id"]})
    [note1, note2] = entry["scientist_notes"]
    # STRUCTURED: their words (the term as said, their own --note) apart from ISAAC's attribution, never one sentence
    assert note1["term_as_said"] == "three_electrode_static_cell" and note1["note"] == "a static cell, no flow"
    assert note1["by"] == WHO and note1["when"].endswith("Z") and note1["source"]["kind"] == "decide_answer"
    assert "text" not in note1
    assert note1["request_id"] == "vr-00001" and note2["request_id"] == "vr-00002"
    assert all(set(c) == {"file", "locator", "run_ids", "origin", "ref"} for c in entry["citations"])
    assert {tuple(c["run_ids"]) for c in entry["citations"]} == {(first["run_id"],), (second["run_id"],)}
    draft = entry["draft"]
    assert draft["drafted"] is True and draft["portal_payload"]["body"]["description"] == ""
    # close, not the same: drafted, with the existing term named and a question for the description
    assert draft["check_existing"] == ["three_electrode"]
    assert "why none of these fits" in draft["check_existing_question"]
    assert draft["check_existing_question"] in draft["portal_payload"]["notes"]
    assert report["vocabulary_requests_read"] == 2


def test_the_drafts_file_has_the_keys_the_websites_loader_reads(b, capsys, tmp_path):
    assert _said(b, capsys, "--target", _miss(b, CELL)["value_id"], "--decision", "correct",
                   "--value", "three_electrode_static_cell")[0] == 2
    out = tmp_path / "drafts.json"
    assert cli.main(["bundle", "propose-terms", str(b), "--out", str(out)]) == 0
    doc = json.loads(out.read_text(encoding="utf-8"))
    # apps/web/src/lib/termDrafts.ts: format, format_version "1", fields[].terms[] with these keys
    assert doc["format"] == "isaac.term-proposal-drafts" and doc["format_version"] == "1"
    [field] = doc["fields"]
    [t] = field["terms"]
    assert isinstance(t["runs"], int) and all(isinstance(r, str) for r in t["run_ids"])
    assert all(isinstance(n, dict) and {"term_as_said", "note", "by", "when", "source"} <= set(n) for n in t["scientist_notes"])
    assert all(isinstance(c["file"], (str, type(None))) and "/" not in (c["file"] or "") for c in t["citations"])
    body = t["draft"]["portal_payload"]["body"]
    assert {"proposal_type", "section", "category", "term"} <= set(body) and t["draft"]["drafted"] is True


def test_a_close_technique_term_is_drafted_with_the_existing_ones_to_check(b, capsys):
    from isaac_api import schema_proposal_rules as R
    import isaac_api.workspace as ws

    sources = R.load_sources_at(ws.REPO_ROOT / "schema/isaac_record_v1.json")
    lists = {"schema": B.allowed_values(TECH), "portal": B.allowed_values(TECH), "portal_section": "System"}
    draft = P._draft(sources, TECH, "XANES", lists)
    assert draft["drafted"] is True and draft["check_existing"] == ["XAS", "HERFD-XAS"]
    # a genuinely new term is never blocked, and gets no list to check
    fresh = P._draft(sources, TECH, "XMCD", lists)
    assert fresh["drafted"] is True and "check_existing" not in fresh


def test_a_term_with_a_space_says_what_to_ask_for(b, capsys):
    assert _said(b, capsys, "--target", _miss(b, REACTION)["value_id"], "--decision", "correct",
                   "--value", "cathodic corrosion")[0] == 2
    report = _report(b, capsys)
    draft = _term(report, REACTION, "cathodic corrosion")["draft"]
    assert draft["drafted"] is False and "e.g. cathodic_corrosion" in draft["ask"] and "one word" in draft["ask"]
    # an example to ask for, never a term ISAAC wrote
    assert _term(report, REACTION, "cathodic_corrosion") is None


def test_an_ai_reading_the_scientist_answered_with_an_allowed_value_is_not_counted(b, capsys):
    target = _miss(b, CELL)
    assert cli.main(["bundle", "annotate", str(b), "--target", target["value_id"], "--value", "three_electrode_static_cell",
                     "--model", "synthetic-test-model", "--excerpt-locator", "page 1 (synthetic)",
                     "--rationale", "synthetic test reading"]) == 0
    assert _term(_report(b, capsys), CELL, "three_electrode_static_cell")["count"] == 1
    assert _decide(b, capsys, "--target", target["value_id"], "--decision", "correct", "--value", "three_electrode")[0] == 0
    report = _report(b, capsys)
    assert _term(report, CELL, "three_electrode_static_cell") is None
    assert report["skipped"]["ai_readings_already_answered"] == 1


# --- --merge ---------------------------------------------------------------------------------------


def test_merge_adds_new_terms_updates_counts_and_keeps_a_written_description(b, capsys, tmp_path):
    out = tmp_path / "drafts.json"
    assert _said(b, capsys, "--target", _miss(b, CELL, 0)["value_id"], "--decision", "correct",
                   "--value", "three_electrode_static_cell")[0] == 2
    assert cli.main(["bundle", "propose-terms", str(b), "--out", str(out)]) == 0
    doc = json.loads(out.read_text(encoding="utf-8"))
    mine = "A three-electrode cell with a static electrolyte, as our sheet names it."
    doc["fields"][0]["terms"][0]["draft"]["portal_payload"]["body"]["description"] = mine
    doc["fields"][0]["terms"].append({"term": "gone_term", "count": 1, "draft": {"drafted": False, "reason": "x"}})
    out.write_text(json.dumps(doc), encoding="utf-8")
    # without --merge an existing file is refused, and the message names --merge
    capsys.readouterr()
    assert cli.main(["bundle", "propose-terms", str(b), "--out", str(out)]) == cli.EXIT_OUTPUT_REFUSED
    assert "--merge" in capsys.readouterr().err
    assert _said(b, capsys, "--target", _miss(b, CELL, 1)["value_id"], "--decision", "correct",
                   "--value", "three_electrode_static_cell")[0] == 2
    assert _said(b, capsys, "--target", _miss(b, ENV)["value_id"], "--decision", "correct",
                   "--value", "aqueous_electrochemistry")[0] == 2
    capsys.readouterr()
    assert cli.main(["bundle", "propose-terms", str(b), "--out", str(out), "--merge"]) == 0
    said = capsys.readouterr().out
    assert "1 term(s) added" in said and "a description in the file was kept (who wrote it is not known)" in said
    assert "1 carried over" in said and "you wrote" not in said
    merged = json.loads(out.read_text(encoding="utf-8"))
    cell = _term(merged, CELL, "three_electrode_static_cell")
    assert cell["count"] == 2 and cell["runs"] == 2
    draft = cell["draft"]
    assert draft["portal_payload"]["body"]["description"] == mine and draft["description_present"] is True
    assert draft["ready_to_send"] is False and draft["sent_by_isaac"] is False and draft["description_written_by"] is None
    assert _term(merged, ENV, "aqueous_electrochemistry")["draft"]["drafted"] is True
    gone = _term(merged, CELL, "gone_term")
    assert gone["carried_over"] is True and gone["current"] is False and gone["draft"]["ready_to_send"] is False
    assert merged["merges"][-1]["descriptions_kept"] == 1 and merged["format_version"] == "1"


def test_merge_reads_no_words_from_the_file_and_never_says_ready_to_send(b, capsys, tmp_path):
    """The drafts file is editable: a note in it is never the scientist's, a description in it is nobody's known
    words, and a term the bundle no longer holds (with whatever the file claims for it) is carried over as not
    current, with no count, notes or citations of its own and never ready to send."""
    assert _said(b, capsys, "--target", _miss(b, CELL)["value_id"], "--decision", "correct",
                   "--value", "three_electrode_static_cell")[0] == 2
    out = tmp_path / "d.json"
    assert cli.main(["bundle", "propose-terms", str(b), "--out", str(out)]) == 0
    doc = json.loads(out.read_text(encoding="utf-8"))
    term = doc["fields"][0]["terms"][0]
    assert term["scientist_notes"][0]["note"] is None  # no --note given: none is made up
    term["scientist_notes"].append({"term_as_said": "x", "note": "Injected: I endorse this term", "by": WHO})
    term["draft"]["portal_payload"]["body"]["description"] = "A description typed into the file by someone, long."
    term["draft"]["ready_to_send"] = True
    doc["fields"][0]["terms"].append({
        "term": "ghost_cell", "count": 99, "by_origin": {"scientist_request": 99}, "runs": 9,
        "scientist_notes": [{"term_as_said": "ghost_cell", "note": "injected", "by": WHO}],
        "citations": [{"file": "x.txt", "locator": "l1"}],
        "draft": {"drafted": True, "ready_to_send": True, "portal_payload": {
            "sent_by_isaac": True, "body": {"proposal_type": "add_term", "section": "x", "category": CELL,
                                            "term": "ghost_cell", "description": "y" * 30}}}})
    out.write_text(json.dumps(doc), encoding="utf-8")
    capsys.readouterr()
    assert cli.main(["bundle", "propose-terms", str(b), "--out", str(out), "--merge"]) == 0
    merged = json.loads(out.read_text(encoding="utf-8"))
    cell = _term(merged, CELL, "three_electrode_static_cell")
    assert [n["note"] for n in cell["scientist_notes"]] == [None]  # only the bundle's request
    assert cell["draft"]["ready_to_send"] is False and cell["draft"]["description_kept_from_earlier_file"] is True
    ghost = _term(merged, CELL, "ghost_cell")
    assert ghost["carried_over"] is True and ghost["current"] is False
    assert ghost["count"] == 0 and ghost["scientist_notes"] == [] and ghost["citations"] == [] and ghost["runs"] == 0
    assert set(ghost["by_origin"].values()) == {0} and ghost["as_in_earlier_file"]["count"] == 99
    assert ghost["draft"]["ready_to_send"] is False and ghost["draft"]["portal_payload"]["sent_by_isaac"] is False
    assert merged["totals"]["carried_over"] == 1
    assert not any((t.get("draft") or {}).get("ready_to_send") for f in merged["fields"] for t in f["terms"])


def test_merge_refuses_what_it_cannot_update(b, capsys, tmp_path):
    out = tmp_path / "other.json"
    out.write_text(json.dumps({"format": "something-else"}), encoding="utf-8")
    capsys.readouterr()
    assert cli.main(["bundle", "propose-terms", str(b), "--out", str(out), "--merge"]) == cli.EXIT_OUTPUT_REFUSED
    assert "not a drafts file" in capsys.readouterr().err and json.loads(out.read_text())["format"] == "something-else"
    assert cli.main(["bundle", "propose-terms", str(b), "--merge"]) == cli.EXIT_USAGE
    assert cli.main(["bundle", "propose-terms", str(b), "--out", str(out), "--merge", "--force"]) == cli.EXIT_USAGE


# --- template import ---------------------------------------------------------------------------------


def _row_of(xlsx, run_id):
    for values in openpyxl.load_workbook(xlsx)["_isaac"].iter_rows(values_only=True):
        if values[0] == "File List" and values[2] == run_id:
            return int(values[1])
    raise AssertionError(run_id)


def _col(header):
    from openpyxl.utils import get_column_letter

    return get_column_letter([c.header for c in T.FILE_COLUMNS].index(header) + 1)


def test_template_import_keeps_an_off_list_cell_on_apply_only(grouped, capsys, tmp_path):
    xlsx = tmp_path / "t.xlsx"
    assert cli.main(["bundle", "template", "export", str(grouped), "--out", str(xlsx)]) == 0
    wb = openpyxl.load_workbook(xlsx)
    rows = {rid: _row_of(xlsx, rid) for rid in ("run-001", "run-002")}
    for row in rows.values():
        wb["File List"][f"{_col('Environment')}{row}"] = "aqueous_electrochemistry"
    wb.save(xlsx)
    before = (grouped / "bundle.json").read_bytes()
    capsys.readouterr()
    assert cli.main(["bundle", "template", "import", str(grouped), str(xlsx), "--by", WHO]) == 0
    preview = capsys.readouterr().out
    assert "With --apply, kept as vocabulary requests" in preview and "would be kept" in preview
    assert (grouped / "bundle.json").read_bytes() == before  # a preview writes nothing
    capsys.readouterr()
    assert cli.main(["bundle", "template", "import", str(grouped), str(xlsx), "--by", WHO, "--apply", "--json"]) == 0
    rep = json.loads(capsys.readouterr().out)
    [kept] = rep["vocabulary_requests"]
    assert kept["term"] == "aqueous_electrochemistry" and kept["new"] is True and len(kept["cells"]) == 2
    [req] = M.load(grouped)["vocabulary_requests"]
    assert req["source"]["kind"] == "template_cell" and req["source"]["file"] == "t.xlsx"
    assert sorted(req["run_ids"]) == ["run-001", "run-002"] and req["requested_by"]["name"] == WHO
    assert any(c["value"] == "aqueous_electrochemistry" for c in rep["invalid_cells"])  # still refused as a value


# --- validate ---------------------------------------------------------------------------------------


def test_pending_vocabulary_names_the_groups_still_missing_the_field():
    bundle = {"vocabulary_requests": [
        {"request_id": "vr-00001", "field_path": CELL, "term": "three_electrode_static_cell", "run_ids": ["run-1"]},
        {"request_id": "vr-00002", "field_path": SCALE, "term": "reversible_hydrogen_electrode", "run_ids": []}]}
    fresh = [("rg-a", {}, {"member_run_ids": ["run-1"], "filled": []}),
             ("rg-b", {}, {"member_run_ids": ["run-2"], "filled": []}),
             ("rg-c", {}, {"member_run_ids": ["run-1"], "filled": [{"path": CELL}, {"path": SCALE}]})]
    rows = P.pending_vocabulary(bundle, fresh)
    cell = next(r for r in rows if r["field_path"] == CELL)
    assert cell["groups"] == ["rg-a"] and cell["groups_total"] == 1 and cell["request_ids"] == ["vr-00001"]
    assert "propose-terms" in cell["statement"]
    scale = next(r for r in rows if r["field_path"] == SCALE)
    assert scale["groups"] == ["rg-a", "rg-b"] and "use `RHE` (an existing term)" in scale["statement"]


def test_validate_says_pending_vocabulary(grouped, capsys):
    from isaac_api.local_import import records as R

    b = M.load(grouped)
    group = next(g for g in b["record_groups"] if g["status"] in B.LIVE_GROUP_STATUSES)
    assert _decide(grouped, capsys, "--target", group["group_id"], "--decision", "confirm")[0] == 0
    run_key = group["member_run_keys"][0]
    run_id = next(r["run_id"] for r in b["runs"] if r["run_key"] == run_key)
    target = next(m for m in b["missing_values"] if m["target_field_path"] == CELL and m["run_id"] == run_id)
    assert _said(grouped, capsys, "--target", target["value_id"], "--decision", "correct",
                   "--value", "three_electrode_static_cell")[0] == 2
    out = R.validate_records(grouped, run_portal=lambda record: {"status": "ran", "result": {"valid": True, "errors": [], "warnings": []},
                                                              "outcome": "publish", "held": [], "verdict": "publish"})
    [row] = out["pending_vocabulary"]
    assert row["field_path"] == CELL and row["term"] == "three_electrode_static_cell" and row["groups"] == [group["group_id"]]


# --- the bundle format ------------------------------------------------------------------------------


def test_a_1_9_0_bundle_still_validates_and_its_first_save_says_what_1_10_0_added(b, capsys):
    data = json.loads((b / "bundle.json").read_text(encoding="utf-8"))
    data["format_version"] = "1.9.0"
    data.pop("format_upgrades", None)
    assert B.validate_bundle(data) == []
    (b / "bundle.json").write_text(json.dumps(data, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    assert B.validate_bundle(B.load_bundle(b)) == []  # loads and validates as stored
    code, err = _said(b, capsys, "--target", _miss(b, REACTION)["value_id"], "--decision", "correct",
                        "--value", "cathodic corrosion")
    assert code == 2 and KEPT in err
    after = M.load(b)
    assert after["format_version"] == B.FORMAT_VERSION == "1.10.0" and len(after["vocabulary_requests"]) == 1
    [u] = after["format_upgrades"]
    assert u["from_format_version"] == "1.9.0" and B.CHANGES_ADDED_IN["1.10.0"] in u["statement"]
    # the section is part of v1.10.0 only as an optional addition: a bundle without it is valid at 1.10.0
    after.pop("vocabulary_requests")
    assert B.validate_bundle(after) == []


# --- review c9-r: order, length, equivalences ---------------------------------------------------------


def test_a_bulk_answer_with_another_refusal_keeps_nothing_whatever_the_order(b):
    """Every target is checked first: a request is kept only when every refusal is off its list."""
    env, ph = _miss(b, ENV)["value_id"], _miss(b, PH)["value_id"]
    for order in ([env, ph], [ph, env]):
        with pytest.raises(M.MutationRefused) as got:
            M.decide(b, target_ids=order, decision="correct", value="aqueous_electrochemistry", by=WHO, orcid=None, as_said=True)
        assert not isinstance(got.value, M.OffListRefused)
        assert not M.load(b).get("vocabulary_requests")
    # the off-list target alone is kept
    with pytest.raises(M.OffListRefused):
        M.decide(b, target_ids=[env], decision="correct", value="aqueous_electrochemistry", by=WHO, orcid=None, as_said=True)
    assert len(M.load(b)["vocabulary_requests"]) == 1


def test_the_longest_upgrade_statement_fits_the_bundle_schema():
    """Every older version (and an unknown one), every newly writable field and every missing import rule."""
    limit = (json.loads((Path(B.__file__).parent / "bundle-v1.schema.json").read_text(encoding="utf-8"))
             ["properties"]["format_upgrades"]["items"]["properties"]["statement"]["maxLength"])
    assert limit == 1500
    rules = [(st, rid) for _v, listed in B.IMPORT_RULES_ADDED_IN.items() for st, rid in listed]
    for version in (None, "0.0.1", *B._VERSION_ORDER[:-1]):
        statement = B.upgrade_statement(version, sorted(B._writable_paths()), rules)
        assert len(statement) <= limit, (version, len(statement))


def test_only_the_electrode_written_out_names_ag_agcl():
    for path in (SCALE, "context.electrochemistry.reference_electrode.type"):
        assert P.equivalent_term(path, "silver silver chloride", ["Ag/AgCl"]) == "Ag/AgCl"
        assert P.equivalent_term(path, "silver_chloride", ["Ag/AgCl"]) is None
        assert P.equivalent_term(path, "silver chloride", ["Ag/AgCl"]) is None
