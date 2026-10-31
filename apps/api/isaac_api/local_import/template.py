"""The ISAAC metadata template round trip: ``isaac bundle template export|import``.

A scientist fills ONE workbook per beamtime with what no file states (facility, beamline,
environment, cell type, reaction, sample form ...). The workbook has the sheets and header
rows of ISAAC's own metadata template (Campaign Info, Configurations, File List,
Vocabulary), plus two sheets ISAAC adds:

* ``ISAAC Suggestions`` (visible): for each answer cell, what the import already holds —
  an extracted value or literal with its file and locator, a hint from the consolidated
  questions, or an AI reading. Evidence beside the blank cell; never in it.
* ``_isaac`` (hidden): the format id and version, the ``bundle_id``, the export time, which
  row is which run, and which column is which official field.

EXPORT never writes a value into an answer cell. Printing the suggestions shows parser
values, so the export records an exposure first (as ``bundle show`` does).

IMPORT refuses any workbook whose ``_isaac`` sheet is missing or names another bundle — so
the example values in ISAAC's raw template can never be imported as answers. It reads only
non-empty answer cells, validates each one (an enum column must hold an allowed value
exactly; a numeric column must hold a number), and turns each valid cell into scientist
decisions through :func:`mutations.record_decision` — the same code as ``isaac bundle
decide``. Each decision's note cites its cell. Without ``apply`` nothing is written. A cell
that would change an earlier scientist decision is a clash and is skipped unless
``supersede``. A column with no official field, or no target in this bundle, is reported
with its cells: nothing is dropped silently, and no mapping is guessed. A cell that REPLACES an
extracted value, or contradicts the run's own file-name token (``m800mV``) or an AI reading of a file,
is named on its own line; one ``adopt`` copied from their sheet before that reading existed is not
recorded but awaits a ruling.

ADOPT (``isaac bundle template adopt``) reads the scientist's OWN filled copy of ISAAC's template
(no ``_isaac`` sheet) and writes this bundle's template from it: their value goes into an answer
cell only where no file reading contradicts it; every other one is a row of the ``Disagreements``
sheet with an empty "Your ruling" cell, which IMPORT reads for exactly the runs that row names. A row lists every
reading of its field (agreeing ones marked); a Campaign Info beamline or facility acronym that names another one
than a facility hint (the folder's name, the acquisition system's) is a row too; an Electrolyte row that states a
concentration takes it in its own ruling cell ("Your ruling: concentration (M)"). A ruling for a field the bundle
asks once for the whole beamtime is recorded per run (``mutations.run_answer_targets``), never beamtime-wide.
"""

from __future__ import annotations

import copy
import math
import os
import re
import shlex
import tempfile
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Mapping, Sequence

from . import bundle as B

FORMAT_ID = "isaac.metadata_template"
#: 1.2.0: File List gains the ISAAC-added column "Control Mode" (``context.electrochemistry.control_mode``,
#: locally answerable since bundle 1.5.0), after pH. A 1.1.0 workbook is still read: its columns
#: are the ones 1.1.0 wrote (:data:`COLUMNS_ADDED_IN`), and nothing past pH is read from it.
#: 1.3.0: File List gains the ISAAC-added column "Current Set-point (mA/cm2)"
#: (``context.electrochemistry.current_setpoint_mA_cm2``, locally answerable since bundle 1.7.0,
#: DEC-92), after Control Mode. A 1.2.0 workbook is read up to Control Mode, a 1.1.0 one up to pH.
#: 1.4.0: the template can FINISH a record. Every field the record composer requires a scientist to
#: supply has a cell: Campaign Info gains "Your group", "Your organization" (``attribution.produced_by.*``)
#: and "Domain" (``system.domain``), appended after the template's own five rows; File List gains
#: "Reaction Drive" and "Reaction Catalysis" (``context.reaction.*``, required once a run states a
#: reaction), appended after the current set-point. Campaign Info's "Organization" was always the
#: FACILITY's (``system.facility.organization``) and now says so; a 1.3.0 workbook, which still reads
#: "Organization" and has five rows, imports as it did (:func:`campaign_fields_of`).
FORMAT_VERSION = "1.4.0"

SHEET_CAMPAIGN = "Campaign Info"
SHEET_CONFIG = "Configurations"
SHEET_FILES = "File List"
SHEET_VOCAB = "Vocabulary"
SHEET_SUGGEST = "ISAAC Suggestions"
SHEET_MARKER = "_isaac"
#: ``template adopt`` only: every value of the scientist's own sheet that a file reading contradicts, or
#: that the field does not take, with an empty "Your ruling" column (:data:`DISAGREE_RULING_COLUMN`).
SHEET_DISAGREE = "Disagreements"
DISAGREE_HEADERS = ("ID", "What", "Field", "Their value", "Their cell(s)", "File reading(s)", "Run(s)", "Your ruling",
                    "Where your ruling goes", "Your ruling: concentration (M)")
DISAGREE_RULING_COLUMN = "H"
#: c7-s: an Electrolyte row whose value also states a concentration ("0.01 M HClO4") takes TWO rulings, the name in
#: "Your ruling" and the concentration (mol/L) here, both recorded for exactly that row's runs. Appended after
#: "Where your ruling goes", so a workbook adopt wrote before it (nine columns) still imports as it did.
DISAGREE_CONC_COLUMN = "J"

EXIT_USAGE = 2
EXIT_INPUT_REFUSED = 3
EXIT_OUTPUT_REFUSED = 4
EXIT_NOT_FOUND = 6

#: Rows of the Configurations sheet that carry the Technique and Config ID dropdowns.
CONFIG_ROWS = 50
_SUGGESTION_CHARS = 3_000
_SUGGESTIONS_PER_CELL = 4
_NUMBER_RE = re.compile(r"\A[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?\Z")


class TemplateRefused(ValueError):
    """Nothing was read or written. ``code`` is the exit code."""

    def __init__(self, message: str, code: int = EXIT_USAGE):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class Column:
    header: str
    path: str | None
    #: why there is no official field (``path is None`` only)
    why_none: str | None = None
    #: a column ISAAC adds AFTER the template's own columns (labelled as such in the workbook)
    isaac_added: bool = False
    #: a note shown on the header cell
    note: str | None = None
    #: other labels the SAME cell had in another template version (same cell, same path): an
    #: already-exported workbook still imports
    aliases: tuple[str, ...] = ()


# ISAAC's own template: sheet names and header text, in its order. No example value.
CAMPAIGN_TITLE = "ISAAC Campaign Metadata"
CAMPAIGN_FIRST_ROW = 3
_FACILITY_ORG_NOTE = ("The organization that RUNS the facility (e.g. SLAC / SSRL for SSRL), not yours: your own "
                      "institution goes in 'Your organization (who produced the data)'.")
#: Campaign Info as template 1.3.0 and earlier wrote it: the template's own five rows.
CAMPAIGN_FIELDS_1_3: tuple[Column, ...] = (
    Column("Facility Name", "system.facility.facility_name"),
    Column("Organization", "system.facility.organization", aliases=("Facility organization (e.g. SLAC / SSRL)",)),
    Column("Beamline/Instrument", "system.facility.beamline"),
    Column("Start Date", None, "the schema has no campaign start date (timestamps.acquired_start_utc is each run's own UTC instant)"),
    Column("Proposal ID", None, "the schema has no proposal id field"),
)
#: 1.4.0: the facility's organization says whose it is, and the record-scope fields the composer
#: requires (the producer, the domain) follow the template's own rows. Rows are read by position,
#: so a row is only ever APPENDED.
CAMPAIGN_FIELDS: tuple[Column, ...] = (
    CAMPAIGN_FIELDS_1_3[0],
    Column("Facility organization (e.g. SLAC / SSRL)", "system.facility.organization", note=_FACILITY_ORG_NOTE,
           aliases=("Organization",)),
    *CAMPAIGN_FIELDS_1_3[2:],
    Column("Your group (who produced the data)", "attribution.produced_by.group", isaac_added=True,
           note=("Added by ISAAC (not in ISAAC's base template). The group or PI that produced these data, as you "
                 "would write it on the record. ISAAC never takes it from a folder name.")),
    Column("Your organization (who produced the data)", "attribution.produced_by.organization", isaac_added=True,
           note=("Added by ISAAC (not in ISAAC's base template). YOUR institution (e.g. your university), "
                 "not the facility's: that is 'Facility organization'.")),
    Column("Domain", "system.domain", isaac_added=True,
           note="Added by ISAAC (not in ISAAC's base template). One of the listed values; beamline data are 'experimental'."),
)


def campaign_fields_of(version: Any) -> tuple[Column, ...]:
    """The Campaign Info rows a workbook of template ``version`` has, in order: 1.3.0 and earlier
    (or an unreadable version) have the template's five, with the facility's organization labelled
    "Organization"."""
    have = _version_key(version)
    return CAMPAIGN_FIELDS if have is not None and have >= (1, 4, 0) else CAMPAIGN_FIELDS_1_3


CONFIG_COLUMNS: tuple[Column, ...] = (
    Column("Config ID", None, "a key that joins File List rows to this sheet; not itself a value"),
    Column("Technique", "system.technique"),
    Column("Mode", None, "the schema has no acquisition-mode field"),
    Column("Primary Signal Channel", None, "a channel choice, not a schema field"),
    Column("Normalization Channel", None, "a channel choice, not a schema field"),
    Column("Detectors", None, "the schema has no detector field"),
    Column("Energy Calibration", None, "the schema has no calibration-note field"),
    Column("Beamline Setup", None, "the schema has no setup-note field"),
    Column("Notes", None, "free notes; no schema field"),
)
FILE_COLUMNS: tuple[Column, ...] = (
    Column("File Name", None, "identifies the row's run; not a value"),
    Column("Config ID", None, "a key that joins this row to a Configurations row; not itself a value"),
    Column("Sample Name", "sample.material.name"),
    Column("Sample Formula", "sample.material.formula"),
    Column("Sample Form", "sample.sample_form"),
    Column("Environment", "context.environment"),
    Column("Reaction", "context.electrochemistry.reaction"),
    Column("Cell Type", "context.electrochemistry.cell_type"),
    Column("Flow Mode", "context.transport.flow_mode"),
    Column("Electrolyte", "context.electrochemistry.electrolyte.name",
           note="The electrolyte NAME only (e.g. KOH). Its concentration goes in 'Electrolyte Concentration (M)', the pH in 'pH'."),
    Column("Potential (V)", "context.electrochemistry.potential_setpoint_V"),
    Column("Potential Scale", "context.electrochemistry.potential_scale"),
    Column("Gas Feed", None, "no single schema field (feed phase and feed composition are separate fields)"),
    # ISAAC's additions, after the template's own columns: "0.1 M KOH" is two schema fields.
    Column("Electrolyte Concentration (M)", "context.electrochemistry.electrolyte.concentration_M", isaac_added=True,
           note="Added by ISAAC (not in ISAAC's base template). A number in mol/L, e.g. 0.1 for 0.1 M KOH."),
    Column("pH", "context.electrochemistry.pH", isaac_added=True,
           note="Added by ISAAC (not in ISAAC's base template). A number."),
    Column("Control Mode", B.CONTROL_MODE_PATH, isaac_added=True,
           note=("Added by ISAAC (not in ISAAC's base template). How the cell was driven for this run, written "
                 "exactly as one of the listed values. Leave it blank if no cell was driven or nobody knows.")),
    Column("Current Set-point (mA/cm2)", B.CURRENT_SETPOINT_PATH, isaac_added=True,
           note=("Added by ISAAC (not in ISAAC's base template). For a galvanostatic run: the commanded current "
                 "DENSITY in mA/cm2, a number, with its sign as stated. Not a total current in mA: ISAAC never "
                 "divides by an electrode area. Leave it blank if you do not know it.")),
    Column("Reaction Drive", "context.reaction.drive", isaac_added=True,
           note=("Added by ISAAC (not in ISAAC's base template). For a run that states a reaction: what drove it, "
                 "exactly as one of the listed values (e.g. electrochemical). Leave it blank otherwise.")),
    Column("Reaction Catalysis", "context.reaction.catalysis", isaac_added=True,
           note=("Added by ISAAC (not in ISAAC's base template). For a run that states a reaction: its catalysis, "
                 "exactly as one of the listed values (e.g. heterogeneous). Leave it blank otherwise.")),
)
#: File List columns a later template version added: ``header -> version``. A workbook of an older
#: version does not have them; its import reads only the columns that version wrote. Such a
#: column is only ever APPENDED: cells are read by position, so one inserted mid-sheet would shift them.
COLUMNS_ADDED_IN: Mapping[str, str] = {"Control Mode": "1.2.0", "Current Set-point (mA/cm2)": "1.3.0",
                                       "Reaction Drive": "1.4.0", "Reaction Catalysis": "1.4.0"}


def _version_key(version: Any) -> tuple[int, ...] | None:
    try:
        return tuple(int(x) for x in str(version).split("."))
    except ValueError:
        return None


def file_columns_of(version: Any) -> tuple[Column, ...]:
    """The File List columns a workbook of template ``version`` has, in order. An unreadable
    version gets none of the later columns: a cell is read only where its column is known."""
    have = _version_key(version)
    return tuple(c for c in FILE_COLUMNS
                 if c.header not in COLUMNS_ADDED_IN or (have is not None and have >= _version_key(COLUMNS_ADDED_IN[c.header])))
#: The electrolyte is one reading in the files ("0.1 M KOH") but two fields here: each column's
#: suggestions also show the AI readings of the other, whole, so the scientist can split it.
_SIBLING_PATHS = {
    "context.electrochemistry.electrolyte.name": ("context.electrochemistry.electrolyte.concentration_M",),
    "context.electrochemistry.electrolyte.concentration_M": ("context.electrochemistry.electrolyte.name",),
}
#: Vocabulary sheet: ISAAC's six lists in the template's order, plus Technique (for the
#: Configurations dropdown). The VALUES are bundle.allowed_values (what decide accepts and export validates against).
VOCAB_COLUMNS: tuple[tuple[str, str], ...] = (
    ("Sample Form", "sample.sample_form"),
    ("Environment", "context.environment"),
    ("Cell Type", "context.electrochemistry.cell_type"),
    ("Flow Mode", "context.transport.flow_mode"),
    ("Potential Scale", "context.electrochemistry.potential_scale"),
    ("Reaction", "context.electrochemistry.reaction"),
    ("Technique", "system.technique"),
    ("Control Mode", B.CONTROL_MODE_PATH),  # ISAAC-added (1.2.0), after ISAAC's own lists
    ("Domain", "system.domain"),  # ISAAC-added (1.4.0): Campaign Info
    ("Reaction Drive", "context.reaction.drive"),  # ISAAC-added (1.4.0): File List
    ("Reaction Catalysis", "context.reaction.catalysis"),
)
_ID_COLUMNS = ("File Name", "Config ID")


def _openpyxl():
    try:
        import openpyxl  # core dependency (pyproject `dependencies`)
    except ImportError as exc:  # pragma: no cover
        raise TemplateRefused(f"openpyxl is not installed: {exc}", code=EXIT_USAGE) from exc
    return openpyxl


def _letter(i: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(i)


def _now_utc() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


# --- what each field accepts ------------------------------------------------------------------


def field_rules() -> dict[str, dict]:
    """``path -> {allowed, numeric, integer, in_schema}`` for every template path (:func:`rule_for_path`),
    read once from the schema the local export targets (``questions._target_schema``)."""
    from . import questions as Q

    schema, _basis = Q._target_schema()
    paths = {c.path for c in (*CAMPAIGN_FIELDS, *CAMPAIGN_FIELDS_1_3, *CONFIG_COLUMNS, *FILE_COLUMNS) if c.path}
    paths |= {p for _, p in VOCAB_COLUMNS}
    return {path: rule_for_path(path, schema=schema) for path in sorted(paths)}


def rule_for_path(path: str, *, schema: Mapping | None = None) -> dict:
    """What one field accepts. ``allowed`` is :func:`bundle.allowed_values` — THE list ``decide`` and
    eligibility accept and the export validates against, the same one every question offers — and
    the type is read from the target schema, then the vendored one. Used by the template import and
    by the local questions screen, so a value typed in either is parsed the same way
    (:func:`parse_answer`) and an offered value is never one ``decide`` refuses."""
    from . import questions as Q

    if schema is None:
        schema, _basis = Q._target_schema()
    node = Q._node(schema, path) or B._schema_node(path)
    enum = B.allowed_values(path)
    declared = node.get("type") if isinstance(node, Mapping) else None
    types = declared if isinstance(declared, list) else [declared] if declared else []
    return {
        "allowed": list(enum) if enum else None,
        "numeric": any(t in ("number", "integer") for t in types),
        "integer": "integer" in types and "number" not in types,
        "in_schema": node is not None,
    }


def parse_answer(raw: Any, rule: Mapping | None) -> tuple[Any, str | None]:
    """``(value, None)`` or ``(None, why it is refused)`` for one typed answer: the template import's
    own cell parsing (:func:`_parse`), so a number is stored as a number exactly as ``--json`` would
    store it, and an enum field takes only one of its allowed values, written exactly."""
    return _parse(raw, rule)


def _parse(raw: Any, rule: Mapping | None) -> tuple[Any, str | None]:
    """``(value, None)`` or ``(None, why it is refused)``. Nothing is coerced into another value:
    surrounding spaces are trimmed, and that is all."""
    if isinstance(raw, bool):
        return None, "a TRUE/FALSE cell; this field does not take one"
    if isinstance(raw, datetime):
        return None, "Excel stored a date here; this field takes text or a number. Format the cell as Text and retype it"
    numeric = bool(rule and rule["numeric"])
    if numeric:
        if isinstance(raw, (int, float)):
            value = raw
        elif isinstance(raw, str) and _NUMBER_RE.match(raw.strip()):
            text = raw.strip()
            value = int(text) if re.fullmatch(r"[+-]?\d+", text) else float(text)
        else:
            return None, f"not a number: {raw!r}"
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None, "not a finite number"
        if rule["integer"] and not (isinstance(value, int) or float(value).is_integer()):
            return None, f"this field takes a whole number; got {raw!r}"
        if rule["integer"]:
            value = int(value)
        return value, None
    if not isinstance(raw, str):
        return None, (f"Excel stored a number ({raw!r}) where text is expected, which can lose leading zeros or "
                      "digits. Format the cell as Text and retype it")
    value = raw.strip()
    if not value:
        return None, None
    allowed = rule["allowed"] if rule else None
    if allowed and value not in allowed:
        near = [a for a in allowed if isinstance(a, str) and a.casefold() == value.casefold()]
        hint = f" (did you mean {near[0]!r}? it must be written exactly)" if near else ""
        return None, f"{value!r} is not an allowed value{hint}; allowed: {', '.join(map(str, allowed))}"
    return value, None


ELECTROLYTE_NAME_PATH = "context.electrochemistry.electrolyte.name"
ELECTROLYTE_CONC_PATH = "context.electrochemistry.electrolyte.concentration_M"
POTENTIAL_PATH = "context.electrochemistry.potential_setpoint_V"
PH_PATH = "context.electrochemistry.pH"
#: A concentration written into the electrolyte NAME ("0.01 M HClO4", "1M KOH", "KOH 100 mM").
_CONCENTRATION_IN_NAME = re.compile(r"(?<![\w.])(\d+(?:\.\d+)?|\.\d+)\s*(mM|M|mol\s*/\s*L)(?![A-Za-z])")


def electrolyte_name_problem(value: Any) -> str | None:
    """Why an Electrolyte (name) cell holding a concentration is refused, or ``None``. The split is the
    scientist's: the number is named back to them, never moved into the concentration column."""
    if not isinstance(value, str):
        return None
    m = _CONCENTRATION_IN_NAME.search(value)
    if not m:
        return None
    if m.group(2) == "mM":
        return (f"name only: {value!r} also states a concentration ({m.group(1)} mM); put it in "
                "'Electrolyte Concentration (M)' in mol/L, and only the name here")
    return f"name only: put {m.group(1)} in 'Electrolyte Concentration (M)', and only the name here"


def electrolyte_statement_parts(value: Any) -> tuple[str, float] | None:
    """``(name, concentration in mol/L)`` of an electrolyte STATEMENT that carries exactly one concentration
    (``"0.5 M H2SO4"`` -> ``("H2SO4", 0.5)``; ``"KOH 100 mM"`` -> ``("KOH", 0.1)``), else ``None``. The template
    asks the two apart (name only; the concentration in its own column), so a cell that is one part of a
    reading agrees with it; a statement with two concentrations is not split."""
    if not isinstance(value, str):
        return None
    found = list(_CONCENTRATION_IN_NAME.finditer(value))
    if len(found) != 1:
        return None
    m = found[0]
    name = " ".join((value[:m.start()] + " " + value[m.end():]).split())
    if not name:
        return None
    number = float(m.group(1))
    return name, (number / 1000 if m.group(2) == "mM" else number)


# --- file-name tokens: a READING to compare against, never a value ---------------------------------
#
# ``m800mV`` is -0.8 V in these beamtimes' naming, but the parser reads no value from it, so a typed
# 0.8 used to clash with nothing. These readings exist ONLY so ``template adopt`` and the import
# preview can say that a typed value contradicts the run's own file name. They are never offered,
# stored or recorded as a value.
_POTENTIAL_TOKEN = re.compile(r"(m?)(\d+)(?:p(\d+))?(mV|V)")
_PH_TOKEN = re.compile(r"pH(\d+)(?:p(\d+))?")
NAME_TOKEN_PATHS = (POTENTIAL_PATH, PH_PATH)


def _token_number(whole: str, frac: str | None) -> Decimal:
    return Decimal(f"{whole}.{frac}") if frac else Decimal(whole)


def name_token_readings(names: Sequence[str], path: str) -> list[dict]:
    """``[{value, token, name}]``: what the whole ``_``-separated tokens of ``names`` (a run's file
    and folder names) read as for ``path`` (potential: ``m?\d+(p\d+)?mV|V``, ``m`` = minus; pH:
    ``pH\d+``). One entry per distinct value; nothing for any other path."""
    if path not in NAME_TOKEN_PATHS:
        return []
    out: dict[str, dict] = {}
    for name in names:
        for token in re.split(r"[_\s]+", name or ""):
            if path == POTENTIAL_PATH and (m := _POTENTIAL_TOKEN.fullmatch(token)):
                volts = _token_number(m.group(2), m.group(3)) / (1000 if m.group(4) == "mV" else 1)
                value = float(-volts if m.group(1) else volts)
            elif path == PH_PATH and (m := _PH_TOKEN.fullmatch(token)):
                value = float(_token_number(m.group(1), m.group(2)))
            else:
                continue
            value = int(value) if path == PH_PATH and value.is_integer() else value
            out.setdefault(repr(float(value)), {"value": value, "token": token, "name": name})
    return list(out.values())


def run_names(bundle: Mapping, run: Mapping) -> list[str]:
    """The names a run's file is known by: its stem, and each part of its file's path (extension dropped)."""
    sources = {s["source_id"]: s for s in bundle["sources"]}
    rel = (sources.get(run.get("acquisition_source_id")) or {}).get("relative_path") or ""
    parts = [p for p in rel.split("/") if p]
    if parts:
        parts[-1] = parts[-1].rsplit(".", 1)[0] if "." in parts[-1] else parts[-1]
    return list(dict.fromkeys([str(run.get("stem") or ""), *parts]))


# --- facility hints: a HINT to compare against, never a value ---------------------------------------
#
# The folder's name (``SSRL_BL152_2019-05_…``) and, when its scan headers match, the acquisition system's name are
# what ``ISAAC Suggestions`` shows beside the facility cells: hints, never answers. A Campaign Info beamline (or a
# facility acronym) that names ANOTHER beamline than such a hint is not adopted (``template adopt``) and is named in
# the import preview: ``62b`` beside ``BL152`` is a different beamline in any spelling, ``BL15-2`` / ``BL152`` /
# ``15-2`` / ``SSRL BL15–2`` are one, ``BL11-2`` and ``BL1-12`` two (:func:`same_beamline`). The organization that runs a facility is in no hint, so it is not compared.
BEAMLINE_PATH = "system.facility.beamline"
FACILITY_NAME_PATH = "system.facility.facility_name"
HINTED_PATHS = (BEAMLINE_PATH, FACILITY_NAME_PATH)
_BEAMLINE_TOKEN = re.compile(r"(?i)BL[-.]?\d+(?:[-.]\d+)?[a-z]?")
_FACILITY_TOKEN = re.compile(r"[A-Z][A-Za-z]{1,9}(?:-[A-Z0-9]{1,4})?")


_HINT_WHERE = {"folder_name": "the imported folder's name",
               "acquisition_system": "the acquisition system the scan headers' column names match"}


#: Every Unicode dash a scientist's sheet may hold (en and em dash, minus sign, hyphen variants): read as ``-``.
_DASHES = re.compile("[\u2010\u2011\u2012\u2013\u2014\u2015\u2212\ufe58\ufe63\uff0d]")
#: The FIRST beamline-shaped token of a value (``BL15-2``, ``bl 15.2``, ``15-2``, ``62b``), a facility word around it
#: (``SSRL BL15-2``, ``BL15-2 (SSRL)``) ignored: its number, the number after a separator, and a trailing letter.
_BEAMLINE_SHAPE = re.compile(r"(?<![a-z0-9])(?:beamline|bl)?[\s_]*(\d+)(?:\s*[-.]\s*(\d+))?([a-z]?)(?![a-z0-9])")


def _beamline_parts(value: Any) -> tuple[str, int | None] | None:
    """``(digits and letter, separator position or None)`` of the first beamline-shaped token of ``value``
    (``SSRL BL15-2`` -> ``("152", 2)``, ``BL152`` -> ``("152", None)``, ``62b`` -> ``("62b", None)``), or ``None``
    when it holds none."""
    if not isinstance(value, str):
        return None
    m = _BEAMLINE_SHAPE.search(_DASHES.sub("-", value).casefold())
    if m is None:
        return None
    first, second, letter = m.group(1), m.group(2) or "", m.group(3) or ""
    return first + second + letter, (len(first) if second else None)


def beamline_form(value: Any) -> str | None:
    """A beamline as it is compared: the first beamline-shaped token, case-folded, a leading ``BL`` / ``beamline``
    and a facility word around it dropped, its separator kept as ``-`` (``BL15-2``, ``SSRL BL15–2``, ``bl 15.2`` and
    ``15-2`` are ``15-2``; ``BL152`` is ``152``; ``62b`` is ``62b``). A value with no such token: the whole value,
    separators dropped (it is then compared as written)."""
    parts = _beamline_parts(value)
    if parts is None:
        if not isinstance(value, str):
            return None
        return re.sub(r"[\s\-_.]", "", _DASHES.sub("-", value.strip().casefold())) or None
    digits, sep = parts
    return digits if sep is None else f"{digits[:sep]}-{digits[sep:]}"


def same_beamline(a: Any, b: Any) -> bool:
    """One beamline: the same digits and letter, split at the same place — or one of them written with no
    separator, which says nothing about where it splits (``BL152`` is ``BL15-2``; ``BL11-2`` is never ``BL1-12``,
    nor ``15-2`` ``1-52``)."""
    pa, pb = _beamline_parts(a), _beamline_parts(b)
    if pa is None or pb is None:
        fa, fb = beamline_form(a), beamline_form(b)
        return fa is not None and fa == fb
    return pa[0] == pb[0] and (pa[1] is None or pb[1] is None or pa[1] == pb[1])


def facility_hint_readings(bundle: Mapping, path: str) -> list[dict]:
    """``[{value, form, source}]``: the beamline (``BL\\d…``) or facility acronym (the word right before such a
    beamline: ``SSRL`` in ``SSRL_BL152``) each facility hint names, with the hint and where it comes from. Nothing
    for any other path, or for a hint that names neither."""
    if path not in HINTED_PATHS:
        return []
    from . import suggestions as S

    out: dict[tuple, dict] = {}
    for h in S.facility_hints(bundle):
        literal = str(h.get("literal") or "")
        tokens = [t for t in re.split(r"[_\s/,;()]+", _DASHES.sub("-", literal)) if t]
        for i, tok in enumerate(tokens):
            if not _BEAMLINE_TOKEN.fullmatch(tok):
                continue
            value = tok if path == BEAMLINE_PATH else (tokens[i - 1] if i and _FACILITY_TOKEN.fullmatch(tokens[i - 1]) else None)
            if value is None:
                continue
            form = beamline_form(value) if path == BEAMLINE_PATH else value.casefold()
            where = _HINT_WHERE.get(str(h.get("kind")), str(h.get("kind")))
            out.setdefault((form, literal), {"value": value, "form": form, "kind": "hint",
                                             "source": f"{where} {literal!r}"})
    return list(out.values())


def facility_hint_lines(bundle: Mapping, path: str, value: Any) -> list[str]:
    """One line per facility hint that names a DIFFERENT beamline (or facility acronym) than ``value``; empty when
    every hint agrees, none names one, or ``value`` is not comparable (a facility name of several words is never
    compared: a folder names an acronym, not a facility's full name)."""
    readings = facility_hint_readings(bundle, path)
    if not readings or not isinstance(value, str):
        return []
    if path == BEAMLINE_PATH:
        mine = beamline_form(value)
    else:
        mine = value.strip().casefold() if _FACILITY_TOKEN.fullmatch(value.strip()) else None
    if not mine:
        return []
    what = "beamline" if path == BEAMLINE_PATH else "facility"
    return [f"differs from a hint: {r['source']} names {what} {r['value']!r}"
            + (f" (compared as {r['form']!r}, yours {mine!r})" if path == BEAMLINE_PATH else "")
            + "; a hint, never an answer" for r in readings
            if (not same_beamline(r["value"], value) if path == BEAMLINE_PATH else r["form"] != mine)]


def _agrees(a: Any, b: Any) -> bool:
    """The same value: numbers by number (within rounding), text after trimming. A number never
    equals text."""
    def num(v):
        return isinstance(v, (int, float)) and not isinstance(v, bool)

    if num(a) and num(b):
        return math.isclose(float(a), float(b), rel_tol=1e-9, abs_tol=1e-12)
    if isinstance(a, str) and isinstance(b, str):
        return a.strip() == b.strip()
    return False


# --- what the bundle holds ----------------------------------------------------------------------


def _live_group_of(bundle: Mapping) -> dict[str, str]:
    """``run_key -> group_id`` of the first live (proposed or confirmed) group holding it."""
    group_of: dict[str, str] = {}
    for g in bundle.get("record_groups") or ():
        if g.get("status") in B.LIVE_GROUP_STATUSES:
            for key in g.get("member_run_keys") or ():
                group_of.setdefault(key, g["group_id"])
    return group_of


def live_run_ids(bundle: Mapping) -> set[str]:
    group_of = _live_group_of(bundle)
    return {r["run_id"] for r in bundle["runs"] if r["run_key"] in group_of}


def live_rows(bundle: Mapping) -> list[dict]:
    """One row per run that belongs to a live record group (proposed or confirmed), in run order.

    ROW IDENTITY IS UNIQUE. A row is named by its acquisition file's path in the import, as the
    bundle records it (``SampleA/scan_001`` and ``SampleB/scan_001`` stay apart). Where two runs
    still share that path (one file holding two acquisitions), the run's label is added. If the
    names still repeat, the export is refused: a scientist could not tell the rows apart."""
    group_of = _live_group_of(bundle)
    sources = {s["source_id"]: s for s in bundle["sources"]}
    rows = []
    for run in bundle["runs"]:
        gid = group_of.get(run["run_key"])
        if gid is None:
            continue
        rel = (sources.get(run.get("acquisition_source_id")) or {}).get("relative_path")
        rows.append({
            "run_id": run["run_id"],
            "run_key": run["run_key"],
            "group_id": gid,
            "file_name": rel or str(run.get("stem") or run["run_id"]),
            "label": run.get("label"),
            "relative_path": rel,
        })
    counts: dict[str, int] = {}
    for r in rows:
        counts[r["file_name"]] = counts.get(r["file_name"], 0) + 1
    for r in rows:
        if counts[r["file_name"]] > 1 and r["label"]:
            r["file_name"] = f"{r['file_name']} · {r['label']}"
    seen: dict[str, list[str]] = {}
    for r in rows:
        seen.setdefault(r["file_name"], []).append(r["run_id"])
    repeated = {name: ids for name, ids in seen.items() if len(ids) > 1}
    if repeated:
        name, ids = next(iter(sorted(repeated.items())))
        raise TemplateRefused(
            f"{len(repeated)} File List name(s) would repeat (e.g. {name!r} for runs {', '.join(ids)}): the file path and "
            "the run label do not tell these runs apart, so a scientist could not tell which row is which. No template "
            "was written.",
            code=EXIT_INPUT_REFUSED,
        )
    return rows


class _Index:
    def __init__(self, bundle: Mapping):
        self.bundle = bundle
        self.missing_by_path: dict[str, list[dict]] = {}
        self.missing_by_run: dict[tuple, list[dict]] = {}
        self.cands_by_path: dict[str, list[dict]] = {}
        self.cands_by_run: dict[tuple, list[dict]] = {}
        for m in bundle["missing_values"]:
            self.missing_by_path.setdefault(m["target_field_path"], []).append(m)
            self.missing_by_run.setdefault((m["run_id"], m["target_field_path"]), []).append(m)
        for c in bundle["candidates"]:
            if c["target_field_path"]:
                self.cands_by_path.setdefault(c["target_field_path"], []).append(c)
                self.cands_by_run.setdefault((c["run_id"], c["target_field_path"]), []).append(c)
        self.paths = set(self.missing_by_path) | set(self.cands_by_path)
        self.sources = {s["source_id"]: s for s in bundle["sources"]}

    def status(self, col: Column) -> tuple[str, str | None]:
        """``(status, reason)``: mapped | no_official_field | no_target_in_bundle."""
        if col.path is None:
            return "no_official_field", col.why_none
        if col.path not in self.paths:
            return "no_target_in_bundle", f"this bundle holds no missing value or extracted value for {col.path}"
        return "mapped", None


# --- export -------------------------------------------------------------------------------------


def _ev_where(index: _Index, ev: Mapping | None) -> str:
    if not ev:
        return "no locator recorded"
    path = ev.get("relative_path") or (index.sources.get(ev.get("source_id") or "") or {}).get("relative_path")
    return f"{path or '(source)'} @ {ev.get('locator') or '(no locator)'}"


def _shown(value: Any) -> str:
    text = repr(value) if not isinstance(value, str) else f"'{value}'"
    return text if len(text) <= 120 else text[:119] + "…"


class _Suggester:
    """What ISAAC already has for one cell. Every candidate or conflict whose value is shown is
    collected in :attr:`exposed`, recorded before the workbook is written."""

    def __init__(self, bundle: Mapping, index: _Index, *, live: set[str]):
        from . import questions as Q

        self.index = index
        self.live = live
        self.exposed: list[tuple[str, str]] = []
        self.hints_of: dict[str, list[dict]] = {}
        self.field_hints_of: dict[str, list[dict]] = {}
        payload = Q.consolidate(bundle)
        for q in payload["questions"]:
            for t in q["targets"]:
                if t["kind"] == "missing_value" and q.get("hints"):
                    self.hints_of.setdefault(t["id"], q["hints"])
            # an OPTIONAL field (the facility's name, organization, beamline) is asked in an optional-values
            # question, whose hints are per field (``field_hints``), not on the question: without these the
            # facility hints (the folder's name, the acquisition system's column names) never reached the sheet.
            # An AI reading is not a hint from the files: it is shown from the target's own interpretations.
            for path, hints in (q.get("field_hints") or {}).items():
                files = [h for h in hints if h.get("kind") != "ai_suggestion"]
                if files:
                    self.field_hints_of.setdefault(path, files)
        self.run_of_candidate = {c["candidate_id"]: c["run_id"] for c in bundle["candidates"]}
        self.ai_of: dict[tuple[str, str], list[dict]] = {}
        for i in bundle.get("ai_interpretations") or ():
            t = i.get("target") or {}
            self.ai_of.setdefault((str(t.get("kind")), str(t.get("id"))), []).append(i)

    def _expose(self, kind: str, tid: str | None) -> None:
        if tid and (kind, tid) not in self.exposed:
            self.exposed.append((kind, tid))

    def _ai(self, kind: str, tid: str, label: str = "") -> list[str]:
        out = []
        for i in self.ai_of.get((kind, tid), [])[:2]:
            ref = (i.get("source_refs") or [{}])[0] if i.get("source_refs") else {}
            where = i.get("source_path") or (self.index.sources.get(ref.get("source_id") or "") or {}).get("relative_path")
            loc = i.get("excerpt_locator") or ref.get("locator")
            # THE WHOLE READING, never clipped: "0.1 M KOH" must reach both electrolyte columns intact.
            value = i.get("proposed_value")
            shown = f"'{value}'" if isinstance(value, str) else repr(value)
            out.append(f"AI suggestion ({(i.get('model') or {}).get('model_id')}, needs review){label}: {shown}"
                       + (f" — {where or '(source)'} @ {loc or '(no locator)'}" if (where or loc) else ""))
        return out

    def _candidate(self, c: Mapping) -> list[str]:
        ev = next((e for e in c["evidence"] if e["method"] != "ai_interpretation"), None)
        if c["value"] is not None:
            what = f"extracted {_shown(c['value'])}"
        elif c["observed_literals"]:
            what = f"literal {_shown(c['observed_literals'][0])} (no value derived)"
        else:
            return self._ai("candidate", c["candidate_id"])
        self._expose("candidate", c["candidate_id"])
        return [f"{what} [{c['value_state']}] — {_ev_where(self.index, ev)}", *self._ai("candidate", c["candidate_id"])]

    def _row_hints(self, hints: Sequence[Mapping], run_id: str) -> list[Mapping]:
        """The hints of a consolidated question that are THIS run's own evidence. A beamtime-wide
        question (e.g. the environment of every run) carries readings of OTHER runs ("'dry' in
        run 83") and beamtime-wide hints (the folder's name) as context; beside one row's cell
        those read as that row's evidence, so they are not shown there. The row's own AI
        readings are shown from its own targets (:meth:`_ai`)."""
        return [h for h in hints if h.get("candidate_id") and self.run_of_candidate.get(h["candidate_id"]) == run_id]

    def _missing(self, m: Mapping, *, with_hints: bool = True, row_of: str | None = None) -> list[str]:
        """``row_of``: the File List row's run — its cell shows only that run's own hints. The
        Campaign Info block (``row_of=None``) shows the question's beamtime-wide hints."""
        out = []
        if with_hints:
            hints = self.hints_of.get(m["value_id"]) or self.field_hints_of.get(m["target_field_path"], [])
            if row_of is not None:
                hints = self._row_hints(hints, row_of)
            for h in hints[:2]:
                if h.get("candidate_id"):
                    self._expose("candidate", h["candidate_id"])
                src = (self.index.sources.get(h.get("source_id") or "") or {}).get("relative_path")
                where = f"{src} @ {h.get('locator') or '(no locator)'}" if src else (h.get("locator") or "(source)")
                out.append(f"hint ({h.get('kind')}, never an answer): {_shown(h.get('literal'))} — {where}")
        return out + self._ai("missing_value", m["value_id"])

    def for_run(self, run_id: str, path: str) -> str:
        lines: list[str] = []
        for c in self.index.cands_by_run.get((run_id, path), []):
            lines += self._candidate(c)
        for m in self.index.missing_by_run.get((run_id, path), []):
            lines += self._missing(m, row_of=run_id)
        for other in _SIBLING_PATHS.get(path, ()):
            label = f" of {other.rsplit('.', 2)[-2]}.{other.rsplit('.', 1)[-1]} (the same reading; split it)"
            sib = [x for c in self.index.cands_by_run.get((run_id, other), []) for x in self._ai("candidate", c["candidate_id"], label)]
            sib += [x for m in self.index.missing_by_run.get((run_id, other), []) for x in self._ai("missing_value", m["value_id"], label)]
            # the same reading recorded on both fields is shown once
            lines += [x for x in dict.fromkeys(sib) if x.replace(label, "") not in lines]
        return self._join(lines)

    def for_campaign(self, path: str) -> str:
        lines: list[str] = []
        for c in self.index.cands_by_run.get((None, path), []):
            lines += self._candidate(c)
        values: dict[str, int] = {}
        for c in self.index.cands_by_path.get(path, []):
            if c["run_id"] is not None and c["value"] is not None:
                values[_shown(c["value"])] = values.get(_shown(c["value"]), 0) + 1
                self._expose("candidate", c["candidate_id"])
        if values:
            lines.append("extracted per run: " + "; ".join(f"{v} in {n} run(s)" for v, n in sorted(values.items())[:5]))
        missing = [m for m in self.index.missing_by_path.get(path, []) if m["run_id"] is None or m["run_id"] in self.live]
        if missing:
            lines += self._missing(missing[0])
        if not lines:
            return ""  # nothing to show: no evidence is not a suggestion
        if missing:
            lines.insert(0, f"no source states it for {len(missing)} target(s) of the File List runs; your answer fills all of them")
        return self._join(lines)

    @staticmethod
    def _join(lines: list[str]) -> str:
        if not lines:
            return ""
        text = "\n".join(lines[:_SUGGESTIONS_PER_CELL])
        if len(lines) > _SUGGESTIONS_PER_CELL:
            text += f"\n… {len(lines) - _SUGGESTIONS_PER_CELL} more"
        return text if len(text) <= _SUGGESTION_CHARS else text[: _SUGGESTION_CHARS - 1] + "…"


def _write_atomic(path: Path, workbook) -> None:
    fd, tmp = tempfile.mkstemp(dir=str(path.parent), prefix=f".{path.name}.", suffix=".tmp")
    os.close(fd)
    try:
        workbook.save(tmp)
        os.chmod(tmp, 0o644)
        os.replace(tmp, path)
    except BaseException:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


def export_template(bundle_dir: Path, out: Path, *, suggestions: bool = True, force: bool = False,
                    _fill=None, _command: str = "bundle template export", _keep_off_list: Mapping | None = None) -> dict:
    """Write the template for ``bundle_dir`` to ``out``. Answer cells are left EMPTY.

    ``_fill`` (``template adopt`` only): called with the built workbook before anything is written;
    it may fill answer cells and add sheets, and returns a dict whose ``_exposed`` targets are recorded
    with the suggestions' before the file exists."""
    from . import mutations as M

    openpyxl = _openpyxl()
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    out = Path(out).expanduser()
    if out.suffix.lower() != ".xlsx":
        raise TemplateRefused("--out must name a .xlsx file.", code=EXIT_USAGE)
    if out.exists() and (out.is_dir() or not force):
        raise TemplateRefused(
            f"{out} already exists. A filled template would be overwritten; pass --force to replace it. Nothing was written.",
            code=EXIT_OUTPUT_REFUSED,
        )
    if not out.parent.is_dir():
        raise TemplateRefused(f"No directory {out.parent}. Nothing was written.", code=EXIT_OUTPUT_REFUSED)
    try:
        bundle = M.load(Path(bundle_dir))
    except M.MutationRefused as exc:
        raise TemplateRefused(str(exc), code=exc.code) from exc
    rows = live_rows(bundle)
    if not rows:
        raise TemplateRefused(
            "This bundle has no live record groups, so there is no run to list. Run `isaac bundle group "
            "<bundle_dir>` first. Nothing was written.",
            code=EXIT_USAGE,
        )
    index = _Index(bundle)
    rules = field_rules()

    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    # Campaign Info: the labels, EMPTY answer cells.
    ws = wb.active
    ws.title = SHEET_CAMPAIGN
    ws["A1"] = CAMPAIGN_TITLE
    ws["A1"].font = bold
    for i, col in enumerate(CAMPAIGN_FIELDS):
        label = ws.cell(row=CAMPAIGN_FIRST_ROW + i, column=1, value=col.header)
        label.font = Font(italic=col.isaac_added)
        if col.note:
            label.comment = Comment(col.note, "ISAAC")
        ws.cell(row=CAMPAIGN_FIRST_ROW + i, column=2).number_format = "@"
    instructions = (
        "Instructions:",
        "1. Fill in Campaign Info above. One answer applies to every run listed in File List.",
        "   'Facility organization' is the facility's (e.g. SLAC / SSRL); 'Your organization' and 'Your group' are who produced the data.",
        "2. Configurations: add one row per setup (Config ID, Technique, ...).",
        "3. File List: one row per run ISAAC found, named by its file's path. Do not edit File Name or reorder rows.",
        "4. Choose each row's Config ID, then fill what you know. Leave a cell blank if you do not know.",
        "5. Electrolyte is the NAME only (e.g. KOH). Put 0.1 M in 'Electrolyte Concentration (M)' as 0.1, and the pH in 'pH'.",
        "   Those two columns, 'Control Mode', 'Current Set-point (mA/cm2)', 'Reaction Drive' and 'Reaction Catalysis', after Gas Feed,",
        "   are added by ISAAC; they are not in ISAAC's base template. Drive and catalysis are needed for a run that states a reaction.",
        "6. 'ISAAC Suggestions' shows what the files already say, with where. It is evidence, not an answer.",
        "7. Then: isaac bundle template import <bundle_dir> <this file> --by \"<your name>\" (preview), add --apply to record.",
    )
    first_instruction = CAMPAIGN_FIRST_ROW + len(CAMPAIGN_FIELDS) + 2
    for i, line in enumerate(instructions):
        ws.cell(row=first_instruction + i, column=1, value=line)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 48

    # Configurations: header only.
    wc = wb.create_sheet(SHEET_CONFIG)
    for j, col in enumerate(CONFIG_COLUMNS, start=1):
        wc.cell(row=1, column=j, value=col.header).font = bold
        wc.column_dimensions[_letter(j)].width = 20
        for r in range(2, 2 + CONFIG_ROWS):
            wc.cell(row=r, column=j).number_format = "@"

    # File List: one row per live-group run; File Name only.
    wf = wb.create_sheet(SHEET_FILES)
    for j, col in enumerate(FILE_COLUMNS, start=1):
        head = wf.cell(row=1, column=j, value=col.header)
        head.font = Font(bold=True, italic=col.isaac_added)
        if col.note:
            head.comment = Comment(col.note, "ISAAC")
        wf.column_dimensions[_letter(j)].width = 18
    wf.column_dimensions["A"].width = 48
    for i, row in enumerate(rows, start=2):
        wf.cell(row=i, column=1, value=row["file_name"])
        for j, col in enumerate(FILE_COLUMNS[1:], start=2):
            rule = rules.get(col.path or "")
            if not (rule and rule["numeric"]):
                wf.cell(row=i, column=j).number_format = "@"
    wf.freeze_panes = "B2"
    last = 1 + len(rows)

    # Vocabulary: ISAAC's own lists (target schema / pinned vocabulary), hidden as in ISAAC's template.
    wv = wb.create_sheet(SHEET_VOCAB)
    vocab_ref: dict[str, str] = {}
    for j, (header, path) in enumerate(VOCAB_COLUMNS, start=1):
        wv.cell(row=1, column=j, value=header).font = bold
        allowed = (rules.get(path) or {}).get("allowed") or []
        for i, v in enumerate(allowed, start=2):
            wv.cell(row=i, column=j, value=v)
        if allowed:
            vocab_ref[path] = f"{SHEET_VOCAB}!${_letter(j)}$2:${_letter(j)}${1 + len(allowed)}"
    wv.sheet_state = "hidden"

    def dropdown(sheet, path: str | None, ref: str) -> None:
        if path in vocab_ref:
            dv = DataValidation(type="list", formula1=vocab_ref[path], allow_blank=True, showErrorMessage=True,
                                errorTitle="Not an allowed value", error="Choose a value from the list.")
        elif path and (rules.get(path) or {}).get("numeric"):
            dv = DataValidation(type="decimal", operator="between", formula1="-1E+300", formula2="1E+300",
                                allow_blank=True, showErrorMessage=True, errorTitle="Not a number", error="Enter a number.")
        else:
            return
        sheet.add_data_validation(dv)
        dv.add(ref)

    for j, col in enumerate(FILE_COLUMNS, start=1):
        dropdown(wf, col.path, f"{_letter(j)}2:{_letter(j)}{last}")
    for j, col in enumerate(CONFIG_COLUMNS, start=1):
        dropdown(wc, col.path, f"{_letter(j)}2:{_letter(j)}{1 + CONFIG_ROWS}")
    for i, col in enumerate(CAMPAIGN_FIELDS):
        dropdown(ws, col.path, f"B{CAMPAIGN_FIRST_ROW + i}")
    cfg = DataValidation(type="list", formula1=f"'{SHEET_CONFIG}'!$A$2:$A${1 + CONFIG_ROWS}", allow_blank=True)
    wf.add_data_validation(cfg)
    cfg.add(f"B2:B{last}")

    # ISAAC Suggestions: the File List layout, evidence in each cell; then the campaign block.
    suggester = _Suggester(bundle, index, live={r["run_id"] for r in rows}) if suggestions else None
    coverage: dict[str, int] = {}
    wsg = wb.create_sheet(SHEET_SUGGEST)
    for j, col in enumerate(FILE_COLUMNS, start=1):
        wsg.cell(row=1, column=j, value=col.header).font = bold
        wsg.column_dimensions[_letter(j)].width = 36
    for i, row in enumerate(rows, start=2):
        wsg.cell(row=i, column=1, value=f"{row['run_id']} · group {row['group_id']} · {row['relative_path'] or row['file_name']}")
        if suggester is None:
            continue
        for j, col in enumerate(FILE_COLUMNS, start=1):
            if col.header == "Config ID":
                text = suggester.for_run(row["run_id"], "system.technique")
                text = ("Technique: " + text) if text else ""
                key = "Configurations/Technique"
            elif col.path:
                text = suggester.for_run(row["run_id"], col.path)
                key = f"{SHEET_FILES}/{col.header}"
            else:
                continue
            if text:
                cell = wsg.cell(row=i, column=j, value=text)
                cell.alignment = wrap
                coverage[key] = coverage.get(key, 0) + 1
    block = last + 3
    wsg.cell(row=block, column=1, value=SHEET_CAMPAIGN).font = bold
    for i, col in enumerate(CAMPAIGN_FIELDS, start=1):
        wsg.cell(row=block + i, column=1, value=col.header)
        if suggester is not None and col.path:
            text = suggester.for_campaign(col.path)
            if text:
                wsg.cell(row=block + i, column=2, value=text).alignment = wrap
                coverage[f"{SHEET_CAMPAIGN}/{col.header}"] = 1
    if suggester is None:
        wsg.cell(row=block + len(CAMPAIGN_FIELDS) + 2, column=1,
                 value="Exported with --no-suggestions: no parser value is shown, so no exposure was recorded.")
    wsg.freeze_panes = "B2"

    # _isaac: the marker and the mappings.
    exported = _now_utc()
    wm = wb.create_sheet(SHEET_MARKER)
    meta = [
        ("format", FORMAT_ID), ("format_version", FORMAT_VERSION), ("bundle_id", bundle["bundle_id"]),
        ("bundle_format_version", bundle.get("format_version")), ("exported_utc", exported),
        ("notice", "Written by `isaac bundle template export`. Do not edit this sheet: import reads it to know "
                   "which row is which run and which column is which official field."),
    ]
    r = 1
    wm.cell(row=r, column=1, value="[meta]")
    for k, v in meta:
        r += 1
        wm.cell(row=r, column=1, value=k)
        wm.cell(row=r, column=2, value=v)
    r += 2
    wm.cell(row=r, column=1, value="[rows]")
    r += 1
    for j, h in enumerate(("sheet", "row", "run_id", "run_key", "group_id", "file_name"), start=1):
        wm.cell(row=r, column=j, value=h)
    for i, row in enumerate(rows, start=2):
        r += 1
        for j, v in enumerate((SHEET_FILES, i, row["run_id"], row["run_key"], row["group_id"], row["file_name"]), start=1):
            wm.cell(row=r, column=j, value=v)
    r += 2
    wm.cell(row=r, column=1, value="[columns]")
    r += 1
    for j, h in enumerate(("sheet", "cell_or_column", "header", "target_field_path", "status", "accepts"), start=1):
        wm.cell(row=r, column=j, value=h)
    columns_report = []

    def column_row(sheet, where, col):
        nonlocal r
        status, reason = index.status(col)
        rule = rules.get(col.path or "") or {}
        accepts = ("one of: " + ", ".join(map(str, rule["allowed"]))) if rule.get("allowed") else (
            "a number" if rule.get("numeric") else "text")
        r += 1
        for j, v in enumerate((sheet, where, col.header, col.path, status, accepts), start=1):
            wm.cell(row=r, column=j, value=v)
        columns_report.append({"sheet": sheet, "where": where, "header": col.header, "target_field_path": col.path,
                               "status": status, "reason": reason})

    for i, col in enumerate(CAMPAIGN_FIELDS):
        column_row(SHEET_CAMPAIGN, f"B{CAMPAIGN_FIRST_ROW + i}", col)
    for j, col in enumerate(CONFIG_COLUMNS, start=1):
        column_row(SHEET_CONFIG, _letter(j), col)
    for j, col in enumerate(FILE_COLUMNS, start=1):
        column_row(SHEET_FILES, _letter(j), col)
    wm.sheet_state = "hidden"

    filled = _fill(wb=wb, bundle=bundle, rows=rows, index=index, rules=rules, marker=wm) if _fill else None
    exposed = list(suggester.exposed) if suggester is not None else []
    for target in (filled or {}).pop("_exposed", ()):
        if target not in exposed:
            exposed.append(target)
    # THE SUGGESTIONS PRINT PARSER VALUES: that is an exposure, recorded before the file exists.
    if exposed:
        try:
            M.expose_and_load(Path(bundle_dir), exposed, _command)
        except M.MutationRefused as exc:
            raise TemplateRefused(f"Nothing written: {exc}", code=exc.code) from exc
    off_list = list((filled or {}).pop("_off_list", ()))
    if filled is not None and _keep_off_list is not None:
        filled["vocabulary_requests"] = []
        if off_list:
            # THEIR OFF-LIST WORDS ARE KEPT (template adopt), recorded before the file exists, as exposures are
            keep = dict(_keep_off_list)

            def change(b):
                kept = keep_off_list_cells(b, off_list, by=keep.get("by"), orcid=keep.get("orcid"), file=keep["file"],
                                           kind="adopted_sheet_cell", command=_command)
                return kept

            try:
                with M.locked(Path(bundle_dir)):
                    held = M.load(Path(bundle_dir))
                    kept = change(held)
                    if any(k["new"] for k in kept):
                        M.save_bundle(Path(bundle_dir), held)
            except M.MutationRefused as exc:
                raise TemplateRefused(f"Nothing written: {exc}", code=exc.code) from exc
            filled["vocabulary_requests"] = kept
    _write_atomic(out, wb)
    return {
        **({"adopted": filled} if filled is not None else {}),
        "written": str(out),
        "format": FORMAT_ID,
        "format_version": FORMAT_VERSION,
        "bundle_id": bundle["bundle_id"],
        "exported_utc": exported,
        "file_list_rows": len(rows),
        "sheets": list(wb.sheetnames),
        "hidden_sheets": [n for n in wb.sheetnames if wb[n].sheet_state == "hidden"],
        "bundle_dir": str(bundle_dir),
        "columns": columns_report,
        "suggestions": {"shown": suggester is not None, "cells_by_column": coverage,
                        "exposures_recorded": len(suggester.exposed) if suggester else 0},
        "statement": "Every answer cell is empty. ISAAC Suggestions is evidence, not an answer.",
    }


def render_export_summary(result: Mapping) -> str:
    """The short text ``template export`` prints: where, how many rows, which sheets, how to import.
    The column mapping is in ``--json`` (and in the workbook's hidden ``_isaac`` sheet)."""
    hidden = set(result.get("hidden_sheets") or ())
    visible = [n for n in result.get("sheets") or () if n not in hidden]
    mapped = [c for c in result["columns"] if c["status"] == "mapped"]
    not_here = [c for c in result["columns"] if c["status"] == "no_target_in_bundle"]
    sug = result["suggestions"]
    lines = [
        f"Wrote {result['written']} (template {result['format_version']}, bundle {result['bundle_id']}).",
        f"{result['file_list_rows']} File List row(s), one per run of a live record group. Sheets: {', '.join(visible)}"
        + (f" (hidden: {', '.join(n for n in result['sheets'] if n in hidden)})." if hidden else "."),
        f"{len(mapped)} answer cell(s)/column(s) map to a field of this bundle; {len(not_here)} map to a field this bundle "
        "holds no target for, and the rest to no official field (--json lists every column).",
        (f"ISAAC Suggestions: {sum(sug['cells_by_column'].values())} cell(s) of evidence; "
         f"{sug['exposures_recorded']} exposure(s) recorded." if sug["shown"] else
         "ISAAC Suggestions: not shown (--no-suggestions); no exposure recorded."),
        result["statement"],
        f"When it is filled: isaac bundle template import {result.get('bundle_dir') or '<bundle_dir>'} {result['written']} "
        '--by "<your name>"   (a preview; add --apply to record)',
    ]
    return "\n".join(lines)


# --- import -------------------------------------------------------------------------------------


def _marker(wb) -> dict:
    if SHEET_MARKER not in wb.sheetnames:
        raise TemplateRefused(
            "This workbook has no `_isaac` sheet, so ISAAC cannot tell which bundle, run or field any cell "
            "belongs to (ISAAC's blank template also holds example values that are not answers). Run "
            "`isaac bundle template export <bundle_dir> --out <file.xlsx>` and fill that file. Nothing was recorded.",
            code=EXIT_INPUT_REFUSED,
        )
    ws = wb[SHEET_MARKER]
    meta: dict[str, Any] = {}
    rows: dict[int, dict] = {}
    columns: list[dict] = []
    #: ``template adopt`` only: the Disagreements rows (``[disagreements]``) and where each adopted cell came from (``[adopted]``)
    extra: dict[str, list[dict]] = {"[disagreements]": [], "[adopted]": []}
    section = None
    header: list[str] = []
    for values in ws.iter_rows(values_only=True):
        first = values[0] if values else None
        if first in ("[meta]", "[rows]", "[columns]", *extra):
            section, header = first, []
            continue
        if first is None:
            continue
        if section == "[meta]":
            meta[str(first)] = values[1] if len(values) > 1 else None
        elif section in ("[rows]", "[columns]"):
            if not header:
                header = [str(v) for v in values if v is not None]
                continue
            entry = dict(zip(header, values))
            if section == "[rows]":
                try:
                    rows[int(entry["row"])] = entry
                except (TypeError, ValueError, KeyError) as exc:
                    raise TemplateRefused(f"The `_isaac` sheet's row mapping is damaged ({entry!r}). Re-export the template.",
                                          code=EXIT_INPUT_REFUSED) from exc
            else:
                columns.append(entry)
        elif section in extra:
            if not header:
                header = [str(v) for v in values if v is not None]
                continue
            extra[section].append(dict(zip(header, values)))
    return {"meta": meta, "rows": rows, "columns": columns, "disagreements": extra["[disagreements]"],
            "adopted": {str(e.get("cell")): e for e in extra["[adopted]"] if e.get("cell")}}


def _cell_ref(sheet: str, ref: str) -> str:
    return f"'{sheet}'!{ref}"


def _read_workbook(xlsx: Path):
    openpyxl = _openpyxl()
    import warnings
    from zipfile import BadZipFile

    xlsx = Path(xlsx).expanduser()
    if not xlsx.is_file():
        raise TemplateRefused(f"No file at {xlsx}.", code=EXIT_NOT_FOUND)
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            return openpyxl.load_workbook(xlsx, data_only=False)
    except (BadZipFile, OSError, KeyError, ValueError) as exc:
        raise TemplateRefused(f"{xlsx.name} is not a readable .xlsx workbook ({type(exc).__name__}). Nothing was recorded.",
                              code=EXIT_INPUT_REFUSED) from exc


def _raw(cell) -> tuple[Any, str | None]:
    """The cell's typed value, or a refusal for a formula. ``(None, None)`` = empty."""
    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
        return None, "a formula, not a value: type the value itself"
    value = cell.value
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    return value, None


def read_cells(wb, marker: Mapping, bundle: Mapping, file_label: str) -> dict:
    """Every non-empty answer cell, parsed and validated, as ``planned`` / ``invalid`` /
    ``unmapped`` / ``nowhere`` entries. Pure: decides nothing."""
    index = _Index(bundle)
    rules = field_rules()
    by_header = {(c.get("sheet"), c.get("header")): c for c in marker["columns"]}
    known = {c.header: c for c in (*CAMPAIGN_FIELDS, *CONFIG_COLUMNS, *FILE_COLUMNS)}
    out: dict[str, Any] = {"cells": [], "invalid": [], "unmapped": {}, "nowhere": [], "rows_refused": [], "awaiting_ruling": [],
                           "off_list": [], "adopted_from": (marker.get("meta") or {}).get("adopted_from")}
    adopted = marker.get("adopted") or {}
    for name in (SHEET_CAMPAIGN, SHEET_CONFIG, SHEET_FILES):
        if name not in wb.sheetnames:
            raise TemplateRefused(f"The workbook has no {name!r} sheet. Re-export the template.", code=EXIT_INPUT_REFUSED)

    def cite(sheet, ref):
        return f"template {file_label} {_cell_ref(sheet, ref)}"

    def unmapped(sheet, col: Column, ref: str, value: Any, reason: str | None = None):
        status, why = index.status(col)
        key = f"{sheet}/{col.header}"
        entry = out["unmapped"].setdefault(key, {
            "sheet": sheet, "header": col.header, "target_field_path": col.path,
            "reason": ("recorded nowhere: no official field" + (f" ({col.why_none})" if col.why_none else ""))
            if status == "no_official_field" else f"recorded nowhere: {why}",
            "cells": [],
        })
        entry["cells"].append({"cell": _cell_ref(sheet, ref), "value": value, **({"note": reason} if reason else {})})

    def check_header(sheet, ws, where: str, col: Column, label_cell: str | None = None):
        labels = (col.header, *col.aliases)  # an alias is the same cell and path under another version's label
        mapped = next((by_header[(sheet, h)] for h in labels if (sheet, h) in by_header), None)
        if mapped is None or str(mapped.get("cell_or_column")) != where or mapped.get("target_field_path") != col.path:
            raise TemplateRefused(
                f"The `_isaac` sheet does not map {sheet!r} {col.header!r} to {where} as this build writes it. "
                "Re-export the template with this build.", code=EXIT_INPUT_REFUSED)
        shown = ws[label_cell].value if label_cell else ws[f"{where}1"].value
        if (shown or "").strip() not in labels:
            raise TemplateRefused(
                f"{_cell_ref(sheet, label_cell or where + '1')} should read {col.header!r} but reads {shown!r}: "
                "the columns were moved or renamed, so no cell can be mapped safely. Re-export the template.",
                code=EXIT_INPUT_REFUSED)

    def take(sheet, ref, col: Column, raw: Any, *, run: Mapping | None, scope: str, via: str | None = None,
             because: str | None = None):
        rule = rules.get(col.path or "")
        value, why = _parse(raw, rule)
        if not why and col.path == ELECTROLYTE_NAME_PATH:
            why = electrolyte_name_problem(value)
        if why:
            if not any(i["cell"] == _cell_ref(sheet, ref) for i in out["invalid"]):  # a Technique cell many rows reach
                out["invalid"].append({"cell": _cell_ref(sheet, ref), "header": col.header, "value": raw, "reason": why})
            word = raw.strip() if isinstance(raw, str) else None
            if col.path and B.off_list_only(col.path, word):
                # ONLY off the field's list: a word the vocabulary lacks, kept as a vocabulary request (1.10.0)
                out["off_list"].append({"cell": _cell_ref(sheet, ref), "header": col.header, "path": col.path, "value": word,
                                        "run_id": run["run_id"] if run else None, "scope": scope})
            return
        if value is None:
            return
        status, reason = index.status(col)
        if status != "mapped":
            unmapped(sheet, col, ref, value)
            return
        note = cite(sheet, ref) + (f" (via {via})" if via else "") + (f" ({because})" if because else "")
        origin = adopted.get(_cell_ref(sheet, ref)) if not because else None
        if origin:
            note += f" (adopted from {origin.get('their_file')} {origin.get('their_cell')})"
        out["cells"].append({"cell": _cell_ref(sheet, ref) + (f" via {via}" if via else ""), "sheet": sheet,
                             "header": col.header, "path": col.path, "value": value, "scope": scope,
                             "run_id": run["run_id"] if run else None, "note": note,
                             # ``template adopt`` copied it from their sheet (no file reading contradicted it then)
                             **({"adopted": {"their_file": origin.get("their_file"), "their_cell": origin.get("their_cell")}}
                                if origin else {})})

    # Campaign Info: the rows THIS workbook's version wrote (1.3.0 and earlier: five, "Organization")
    wsc = wb[SHEET_CAMPAIGN]
    workbook_version = (marker.get("meta") or {}).get("format_version")
    for i, col in enumerate(campaign_fields_of(workbook_version)):
        row = CAMPAIGN_FIRST_ROW + i
        check_header(SHEET_CAMPAIGN, wsc, f"B{row}", col, label_cell=f"A{row}")
        raw, why = _raw(wsc[f"B{row}"])
        if why:
            out["invalid"].append({"cell": _cell_ref(SHEET_CAMPAIGN, f"B{row}"), "header": col.header, "value": wsc[f"B{row}"].value, "reason": why})
            continue
        if raw is None:
            continue
        out["read"] = out.get("read", 0) + 1
        if col.path is None:
            unmapped(SHEET_CAMPAIGN, col, f"B{row}", raw)
        else:
            take(SHEET_CAMPAIGN, f"B{row}", col, raw, run=None, scope="experiment")

    # Configurations: read every row; Technique reaches a run through File List's Config ID.
    wcf = wb[SHEET_CONFIG]
    for j, col in enumerate(CONFIG_COLUMNS, start=1):
        check_header(SHEET_CONFIG, wcf, _letter(j), col)
    configs: dict[str, dict] = {}
    duplicate: set[str] = set()
    for r in range(2, wcf.max_row + 1):
        cid_raw, _ = _raw(wcf.cell(row=r, column=1))
        cid = str(cid_raw).strip() if cid_raw is not None else None
        cells = {}
        for j, col in enumerate(CONFIG_COLUMNS, start=1):
            raw, why = _raw(wcf.cell(row=r, column=j))
            ref = f"{_letter(j)}{r}"
            if why:
                out["invalid"].append({"cell": _cell_ref(SHEET_CONFIG, ref), "header": col.header, "value": wcf[ref].value, "reason": why})
                continue
            if raw is None:
                continue
            out["read"] = out.get("read", 0) + 1
            cells[col.header] = (ref, raw)
        if not cells:
            continue
        if not cid:
            for header, (ref, raw) in cells.items():
                out["invalid"].append({"cell": _cell_ref(SHEET_CONFIG, ref), "header": header, "value": raw,
                                       "reason": "this Configurations row has no Config ID, so no run can refer to it"})
            continue
        if cid in configs:
            duplicate.add(cid)
        configs[cid] = {"row": r, "cells": cells}
        for header, (ref, raw) in cells.items():
            col = known[header]
            if col.path is None and header != "Config ID":
                unmapped(SHEET_CONFIG, col, ref, raw)
    for cid in duplicate:
        out["invalid"].append({"cell": _cell_ref(SHEET_CONFIG, f"A{configs[cid]['row']}"), "header": "Config ID", "value": cid,
                               "reason": "this Config ID is defined on more than one row, so which row a File List row means is unknown"})

    # File List
    wfl = wb[SHEET_FILES]
    file_columns = file_columns_of(workbook_version)
    for j, col in enumerate(file_columns, start=1):
        check_header(SHEET_FILES, wfl, _letter(j), col)
    runs = {r["run_id"]: r for r in bundle["runs"]}
    rows = marker["rows"]
    # ROW IDENTITY MUST BE UNIQUE in the mapping itself, or a moved row could match a wrong run.
    for key in ("file_name", "run_id"):
        seen: dict[str, int] = {}
        for rnum, entry in sorted(rows.items()):
            v = str(entry.get(key))
            if v in seen:
                raise TemplateRefused(
                    f"The `_isaac` sheet names {key} {v!r} on rows {seen[v]} and {rnum}, so a row's run cannot be "
                    "told apart. Re-export the template. Nothing was recorded.", code=EXIT_INPUT_REFUSED)
            seen[v] = rnum
    row_of_name = {str(e.get("file_name")): rnum for rnum, e in rows.items()}
    # the runs a Campaign Info answer reaches: the File List's runs that are still in a live group
    out["file_list_run_ids"] = {str(e.get("run_id")) for e in rows.values()} & live_run_ids(bundle)
    max_row = max([wfl.max_row, *rows]) if rows else wfl.max_row
    for r in range(2, max_row + 1):
        entry = rows.get(r)
        values = []
        for j, col in enumerate(file_columns, start=1):
            raw, why = _raw(wfl.cell(row=r, column=j))
            values.append((col, f"{_letter(j)}{r}", raw, why))
        # A cell typed where a LATER template version has a column this workbook's version does not
        # (e.g. past pH in a 1.1.0 workbook) is not read, and is SAID, never dropped silently.
        for j, col in enumerate(FILE_COLUMNS[len(file_columns):], start=len(file_columns) + 1):
            raw = wfl.cell(row=r, column=j).value
            if raw is not None and not (isinstance(raw, str) and not raw.strip()):
                out["nowhere"].append({
                    "cell": _cell_ref(SHEET_FILES, f"{_letter(j)}{r}"), "header": col.header, "value": raw,
                    "reason": (f"this workbook is template {workbook_version!r}, which has no column {_letter(j)} "
                               f"({col.header} was added in {COLUMNS_ADDED_IN.get(col.header)}), so the cell is not "
                               "read; re-export the template to answer it"),
                })
        answered = [v for v in values[1:] if v[2] is not None or v[3]]
        if entry is None:
            if answered:
                out["rows_refused"].append({"row": r, "reason": "ISAAC did not export this row, so it names no run; its cells are recorded nowhere",
                                            "cells": [_cell_ref(SHEET_FILES, ref) for _, ref, _, _ in answered]})
            continue
        run = runs.get(str(entry.get("run_id")))
        name_cell = wfl.cell(row=r, column=1).value
        # EXACT match, not trimmed or case-folded: the name is the row's identity.
        if run is None or not isinstance(name_cell, str) or name_cell != str(entry.get("file_name")):
            elsewhere = row_of_name.get(name_cell) if isinstance(name_cell, str) else None
            if answered:
                out["rows_refused"].append({
                    "row": r,
                    "reason": (f"File Name reads {name_cell!r} but ISAAC exported {entry.get('file_name')!r} on this row"
                               + (f" ({name_cell!r} was exported on row {elsewhere}: rows were reordered)" if elsewhere else
                                  " (the name was edited)")
                               + ", so the row's run is unknown; nothing on it is recorded")
                    if run is not None else f"the run {entry.get('run_id')!r} is not in this bundle",
                    "cells": [_cell_ref(SHEET_FILES, ref) for _, ref, _, _ in answered],
                })
            continue
        for col, ref, raw, why in values[1:]:
            if why:
                out["invalid"].append({"cell": _cell_ref(SHEET_FILES, ref), "header": col.header, "value": wfl[ref].value, "reason": why})
                continue
            if raw is None:
                continue
            out["read"] = out.get("read", 0) + 1
            if col.header == "Config ID":
                cid = str(raw).strip()
                if cid in duplicate:
                    continue  # reported on the Configurations row
                config = configs.get(cid)
                if config is None:
                    out["invalid"].append({"cell": _cell_ref(SHEET_FILES, ref), "header": col.header, "value": raw,
                                           "reason": f"no Configurations row has Config ID {cid!r}"})
                    continue
                tech = config["cells"].get("Technique")
                if tech is not None:
                    take(SHEET_CONFIG, tech[0], known["Technique"], tech[1], run=run, scope="run",
                         via=_cell_ref(SHEET_FILES, ref))
                continue
            take(SHEET_FILES, ref, col, raw, run=run, scope="run")
    # Technique cells of configurations no row refers to: read, but they reach no run.
    used = {str(wfl.cell(row=r, column=2).value).strip() for r in rows if wfl.cell(row=r, column=2).value is not None}
    for cid, config in configs.items():
        tech = config["cells"].get("Technique")
        if tech is not None and cid not in used and cid not in duplicate:
            out["nowhere"].append({"cell": _cell_ref(SHEET_CONFIG, tech[0]), "header": "Technique", "value": tech[1],
                                   "reason": f"no File List row chose Config ID {cid!r}, so it reaches no run"})
    _read_rulings(wb, marker, runs, known, take, out)
    return out


def _read_rulings(wb, marker: Mapping, runs: Mapping, known: Mapping[str, Column], take, out: dict) -> None:
    """A workbook ``template adopt`` wrote: each Disagreements row's "Your ruling" cell, typed by the
    scientist, is read like an answer cell for exactly the targets ``_isaac`` names for that row. An
    empty ruling records nothing and is listed (``awaiting_ruling``); nothing ISAAC chose is ever
    read from that sheet."""
    entries = marker.get("disagreements") or ()
    if not entries:
        return
    if SHEET_DISAGREE not in wb.sheetnames:
        raise TemplateRefused(f"The `_isaac` sheet lists disagreements but the workbook has no {SHEET_DISAGREE!r} sheet. "
                              "Run `template adopt` again. Nothing was recorded.", code=EXIT_INPUT_REFUSED)
    ws = wb[SHEET_DISAGREE]
    # a workbook adopt wrote before the concentration ruling column (nine columns) has no column J to read
    conc_column = str(ws[f"{DISAGREE_CONC_COLUMN}1"].value or "").strip() == DISAGREE_HEADERS[9]
    conc_col = next(c for c in FILE_COLUMNS if c.path == ELECTROLYTE_CONC_PATH)
    for e in entries:
        try:
            row = int(e.get("row"))
        except (TypeError, ValueError):
            raise TemplateRefused(f"The `_isaac` sheet's disagreement mapping is damaged ({e!r}). Run `template adopt` "
                                  "again. Nothing was recorded.", code=EXIT_INPUT_REFUSED) from None
        did = str(e.get("id"))
        ref = f"{DISAGREE_RULING_COLUMN}{row}"
        conc_ref = f"{DISAGREE_CONC_COLUMN}{row}"
        shown = ws[f"A{row}"].value
        if shown != did:
            out["rows_refused"].append({"row": row, "cells": [_cell_ref(SHEET_DISAGREE, ref)],
                                        "reason": f"{_cell_ref(SHEET_DISAGREE, f'A{row}')} should read {did!r} but reads "
                                                  f"{shown!r}: the Disagreements rows were moved, so no ruling on it is read"})
            continue
        col = known.get(str(e.get("header")))
        if col is None or col.path != e.get("target_field_path"):
            raise TemplateRefused(f"The `_isaac` sheet maps disagreement {did} to a column this build does not know "
                                  f"({e.get('header')!r}). Run `template adopt` again. Nothing was recorded.",
                                  code=EXIT_INPUT_REFUSED)
        because = f"your ruling on disagreement {did}" + (f": {e.get('their_file')} {e.get('their_cells')}" if e.get("their_cells") else "")
        runs_named = [] if e.get("scope") == "experiment" else [t for t in str(e.get("targets") or "").split(";") if t]

        def record(cell_ref: str, column: Column, raw: Any, why_said: str):
            out["read"] = out.get("read", 0) + 1
            if e.get("scope") == "experiment":
                take(SHEET_DISAGREE, cell_ref, column, raw, run=None, scope="experiment", because=why_said)
                return
            for run_id in runs_named:
                run = runs.get(run_id)
                if run is None:
                    out["nowhere"].append({"cell": _cell_ref(SHEET_DISAGREE, cell_ref), "header": column.header, "value": raw,
                                           "reason": f"the run {run_id!r} is not in this bundle"})
                    continue
                take(SHEET_DISAGREE, cell_ref, column, raw, run=run, scope="run", because=why_said)

        # M3: the CONCENTRATION of an electrolyte row ("0.01 M HClO4"), ruled in its own column, for the same runs
        if conc_column:
            craw, cwhy = _raw(ws[conc_ref])
            if cwhy:
                out["invalid"].append({"cell": _cell_ref(SHEET_DISAGREE, conc_ref), "header": conc_col.header,
                                       "value": ws[conc_ref].value, "reason": cwhy})
            elif craw is not None and e.get("kind") != "name_only":
                out["nowhere"].append({"cell": _cell_ref(SHEET_DISAGREE, conc_ref), "header": DISAGREE_HEADERS[9], "value": craw,
                                       "reason": f"only an Electrolyte row that also states a concentration takes a "
                                                 f"concentration ruling; {did} is about {col.header}"})
            elif craw is not None:
                record(conc_ref, conc_col, craw, f"your concentration ruling on disagreement {did}"
                       + (f": {e.get('their_file')} {e.get('their_cells')}" if e.get("their_cells") else ""))
        raw, why = _raw(ws[ref])
        if why:
            out["invalid"].append({"cell": _cell_ref(SHEET_DISAGREE, ref), "header": col.header, "value": ws[ref].value, "reason": why})
            continue
        if raw is None:
            reading = ws[f"F{row}"].value
            out["awaiting_ruling"].append({"id": did, "cell": _cell_ref(SHEET_DISAGREE, ref), "header": col.header,
                                           "target_field_path": col.path, "their_value": ws[f"D{row}"].value,
                                           "why": ("readings: " + "; ".join(str(reading).splitlines())) if reading else ws[f"B{row}"].value})
            continue
        record(ref, col, raw, because)


def _same_value(a: Any, b: Any) -> bool:
    """Equal as values: numbers by number (``0.1 == 0.1``, ``1 == 1.0``), everything else by type
    and value. ``True`` is not ``1``; ``"0.1"`` is not ``0.1``."""
    def num(v):
        return isinstance(v, (int, float)) and not isinstance(v, bool)

    if num(a) and num(b):
        return a == b
    return type(a) is type(b) and a == b


def _prior(bundle: Mapping, kind: str, obj: Mapping, latest: Mapping) -> Mapping | None:
    if kind == "candidate":
        return B.candidate_decision(bundle, obj, latest)
    return latest.get((kind, obj["value_id"]))


def plan_and_record(bundle: dict, cells: Mapping, *, by: str, orcid: str | None, supersede: bool) -> dict:
    """Turn read cells into decisions ON ``bundle`` (in memory) through
    :func:`mutations.record_decision`. The caller saves it (``--apply``) or discards it (preview)."""
    from . import mutations as M

    index = _Index(bundle)
    latest = B._decisions_by_target(bundle.get("scientist_decisions") or ())
    report: dict[str, list] = {"decisions": [], "clashes": [], "already": [], "refused": [], "nowhere": list(cells["nowhere"]),
                               "left_alone": [], "campaign_not_filled": [], "contradictions": []}
    runs_by_id = {r["run_id"]: r for r in bundle["runs"]}
    contradicted: set[tuple[str, str]] = set()
    awaiting: list[dict] = []
    # a file reading is what adopt compares (:class:`_Adopter`): a parser value, a file-name token, an AI
    # reading of a FILE; an AI reading of the sheet the cells were adopted from is that sheet, not a file
    adopter = _Adopter(bundle, index, field_rules(), str(cells.get("adopted_from") or ""))

    def file_readings(cell: Mapping, decision: str, kind: str, obj: Mapping) -> list[str]:
        """What a file says against this cell for this target, one line each. A literal from which no value
        was derived (a macro's file name, a sample token) is not a reading of a value: answering it
        replaces nothing."""
        readings = []
        if decision == "correct" and kind == "candidate" and obj.get("value") is not None:
            ev = next((e for e in obj["evidence"] if e["method"] != "ai_interpretation"), None)
            readings.append(f"replaces extracted {_shown(obj['value'])} — {_ev_where(index, ev)}")
        run = runs_by_id.get(obj.get("run_id") or "")
        if run is not None:
            for t in name_token_readings(run_names(bundle, run), cell["path"]):
                if not _agrees(t["value"], cell["value"]):
                    readings.append(f"contradicts the file name token {t['token']} ({t['value']!r}) in {t['name']}")
        for r in adopter.against(cell["value"], obj.get("run_id"), cell["path"]):
            if r["kind"] == "ai_reading":
                readings.append(f"contradicts the {r['source']}: {_shown(r['value'])}")
        return readings

    def contradiction(cell: Mapping, obj: Mapping, readings: list[str], outcome: str) -> None:
        """A cell that REPLACES an extracted value, or contradicts the run's file-name token or an AI reading
        of a file, is named on its own line: a count of ``correct`` decisions never shows it. ``outcome``
        says what happens to it: ``recorded_as_typed`` (the scientist typed it, or ruled it) or
        ``awaiting_ruling`` (adopted from their sheet before the reading existed: not recorded)."""
        key = (cell["cell"], obj.get("run_id") or "")
        if readings and key not in contradicted:
            contradicted.add(key)
            report["contradictions"].append({"cell": cell["cell"], "run_id": obj.get("run_id"), "target_field_path": cell["path"],
                                             "value": cell["value"], "readings": readings, "outcome": outcome,
                                             **({"adopted": cell["adopted"]} if cell.get("adopted") else {})})

    report["asked_beamtime_wide"] = []
    asked_of: dict[str, dict | None] = {}
    payload_box: list[Mapping] = []

    def beamtime_question(path: str) -> dict | None:
        """The question that asks ``path`` ONCE for the whole beamtime, when this bundle holds the field only
        at experiment scope (a beamtime candidate, its shared conflict, a record-scope missing value): a run's
        cell has no target of its own then, and that question is what answers it (an answered question too,
        with the same id). ``None`` when no such question exists. Built once, only when a cell needs it."""
        if path in asked_of:
            return asked_of[path]
        from . import questions as Q

        ids = {m["value_id"] for m in index.missing_by_run.get((None, path), [])}
        ids |= {c["candidate_id"] for c in index.cands_by_run.get((None, path), [])}
        ids |= {cf["conflict_id"] for cf in bundle.get("conflicts") or () if cf.get("candidate_id") in ids}
        found = None
        if ids:
            if not payload_box:
                payload_box.append(Q.consolidate(bundle, include_answered=True, explain_counter=False))
            for q in payload_box[0]["questions"]:
                if q["kind"] not in Q.VALUE_KINDS or q.get("informational"):
                    continue
                hit = next((tids for field, tids in Q._value_paths(q) if field == path and ids.intersection(tids)), None)
                if hit is not None:
                    found = {"question_id": q["question_id"], "command": Q.scoped_answer_command(q, path),
                             "answered": any(tid in hit and d.get("decision") in (B.CONFIRMING_DECISIONS | {"reject"})
                                             for (_kind, tid), d in latest.items())}
                    break
        asked_of[path] = found
        return found

    reached = cells.get("file_list_run_ids") or set()
    planned: list[tuple[dict, str, str, Mapping, str]] = []  # (cell, decision, kind, obj, target_id)
    # H1: a Campaign Info beamline (or facility acronym) naming another one than a facility hint — named once per
    # cell (never once per run it fills); held back when adopt copied it, recorded when the scientist typed it
    hinted = {c["cell"]: lines for c in cells["cells"]
              if c["scope"] == "experiment" and (lines := facility_hint_lines(bundle, c["path"], c["value"]))}
    report["ruled_per_run"] = []
    slots_added: list[str] = []
    for cell in cells["cells"]:
        path, value = cell["path"], cell["value"]
        targets: list[tuple[str, str, Mapping, str]] = []
        if cell["scope"] == "experiment":
            outside = []
            for m in index.missing_by_path.get(path, []):
                if m["run_id"] is None or m["run_id"] in reached:  # record scope, or a File List run
                    targets.append(("correct", "missing_value", m, m["value_id"]))
                else:
                    outside.append(m)
            if outside:
                report["campaign_not_filled"].append({
                    "cell": cell["cell"], "target_field_path": path, "targets": len(outside),
                    "run_ids": sorted({m["run_id"] for m in outside})[:20],
                    "reason": "these runs are not in the template's File List or no longer in a live record group, so a "
                              "Campaign Info answer does not reach them; nothing was recorded for them",
                })
            extracted = [c for c in index.cands_by_path.get(path, [])]
            if extracted:
                report["left_alone"].append({
                    "cell": cell["cell"], "target_field_path": path, "candidates": len(extracted),
                    "reason": "a Campaign Info answer fills only the targets no source states; these extracted values "
                              "are left for you to confirm or correct per run",
                })
        else:
            for m in index.missing_by_run.get((cell["run_id"], path), []):
                targets.append(("correct", "missing_value", m, m["value_id"]))
            for c in index.cands_by_run.get((cell["run_id"], path), []):
                targets.append(("confirm" if _same_value(c["value"], value) else "correct", "candidate", c, c["candidate_id"]))
        if not targets:
            asked = beamtime_question(path) if cell["scope"] == "run" else None
            if asked is not None and cell["sheet"] == SHEET_DISAGREE and cell.get("run_id"):
                # M1(b): A RULING NAMES ITS RUNS. Its value is recorded for exactly them — on a per-run answer slot
                # when the import asks the field once for the beamtime (as `decide --question --only group:<id>`
                # does), never on the beamtime-wide target, which would reach every run of the beamtime
                ids, added = M.run_answer_targets(bundle, path, [cell["run_id"]], question_id=asked["question_id"],
                                                  how=f"their ruling {cell['cell'].split(' via ', 1)[0]}")
                slots_added += added
                slot_of = {m["value_id"]: m for m in bundle["missing_values"]}
                for vid in ids:
                    targets.append(("correct", "missing_value", slot_of[vid], vid))
                report["ruled_per_run"].append({"cell": cell["cell"], "header": cell["header"], "target_field_path": path,
                                                "run_id": cell["run_id"], "question_id": asked["question_id"],
                                                "answered": asked["answered"], "slots_added": len(added)})
            elif asked is not None:
                generic = asked["command"].replace('--by "<your name>"', f"--by {shlex.quote(by)}")
                command = generic
                if isinstance(value, str) and "--json" not in command:
                    command = re.sub(r"--value <[^>]*>", lambda _m: f"--value {shlex.quote(value)}", command)
                report["asked_beamtime_wide"].append({
                    "cell": cell["cell"], "header": cell["header"], "value": value, "run_id": cell["run_id"],
                    "target_field_path": path, "question_id": asked["question_id"], "answered": asked["answered"],
                    "command": command, "_generic_command": generic,
                    "reason": (f"this bundle asks {path} once for the whole beamtime (question {asked['question_id']}"
                               + (", answered" if asked["answered"] else "") + "), not per run, so this cell is not "
                               "recorded; " + ("to change that answer" if asked["answered"] else "answer that question")
                               + f": {command}"),
                })
                continue
            else:
                report["nowhere"].append({"cell": cell["cell"], "header": cell["header"], "value": value,
                                          "reason": f"this run has no missing or extracted value for {path} in this bundle"})
                continue
        for decision, kind, obj, tid in targets:
            planned.append((cell, decision, kind, obj, tid))

    # ONE QUESTION, CELLS THAT DIFFER: one command per value would be wrong (running them in turn leaves the
    # last standing, for every run). The cells name the values; the scientist says which applies.
    by_question: dict[str, set[str]] = {}
    for x in report["asked_beamtime_wide"]:
        by_question.setdefault(x["question_id"], set()).add(repr(x["value"]))
    for x in report["asked_beamtime_wide"]:
        generic = x.pop("_generic_command")
        values = sorted(by_question[x["question_id"]])
        if len(values) > 1:
            x["command"], x["values_differ"] = generic, values
            x["reason"] = (f"this bundle asks {x['target_field_path']} once for the whole beamtime (question "
                           f"{x['question_id']}" + (", answered" if x["answered"] else "") + "), not per run, and this "
                           f"workbook's cells for it hold {len(values)} different values ({', '.join(values)}): nothing is "
                           "recorded, and the scientist says which applies to the whole beamtime, then: " + generic)

    # Two cells of THIS workbook giving one target different values: neither is recorded.
    values_of: dict[str, set[str]] = {}
    for cell, _d, _k, _o, tid in planned:
        values_of.setdefault(tid, set()).add(repr(cell["value"]))
    split = {tid for tid, vs in values_of.items() if len(vs) > 1}
    calls: dict[tuple, list[str]] = {}
    seen_target: set[str] = set()
    for cell, decision, kind, obj, tid in planned:
        if tid in split:
            report["clashes"].append({"cell": cell["cell"], "target": {"kind": kind, "id": tid}, "value": cell["value"],
                                      "reason": "another cell of this workbook gives this target a different value; neither is recorded"})
            continue
        if tid in seen_target:
            continue
        seen_target.add(tid)
        prior = _prior(bundle, kind, obj, latest)
        if prior is not None and prior.get("decision") in (B.CONFIRMING_DECISIONS | {"reject"}):
            base = obj.get("value") if kind == "candidate" else None
            before = B.final_value(base, prior)
            if _same_value(before, cell["value"]) and prior.get("decision") != "reject":
                report["already"].append({"cell": cell["cell"], "target": {"kind": kind, "id": tid}, "value": cell["value"],
                                          "decided_by": (prior.get("decided_by") or {}).get("name"), "decision_id": prior.get("decision_id")})
                continue
            if not supersede:
                report["clashes"].append({
                    "cell": cell["cell"], "target": {"kind": kind, "id": tid}, "value": cell["value"],
                    "existing": {"decision": prior.get("decision"), "value": before, "decision_id": prior.get("decision_id"),
                                 "decided_by": (prior.get("decided_by") or {}).get("name")},
                    "reason": "an earlier scientist decision gives a different value; skipped (pass --supersede to record this one)",
                })
                continue
        if cell["cell"] in hinted:
            # once per cell: a Campaign Info answer fills every run's target, and one hint line covers them all
            key = (cell["cell"], "")
            if key not in contradicted:
                contradicted.add(key)
                outcome = "awaiting_ruling" if cell.get("adopted") else "recorded_as_typed"
                report["contradictions"].append({"cell": cell["cell"], "run_id": None, "target_field_path": cell["path"],
                                                 "value": cell["value"], "readings": hinted[cell["cell"]], "outcome": outcome,
                                                 "hint": True, **({"adopted": cell["adopted"]} if cell.get("adopted") else {})})
                if cell.get("adopted"):
                    origin = cell["adopted"]
                    awaiting.append({"id": None, "cell": cell["cell"], "header": cell["header"], "target_field_path": cell["path"],
                                     "run_id": None, "target": None, "their_value": cell["value"],
                                     "why": (f"adopted from {origin.get('their_file')} {origin.get('their_cell')}, but it "
                                             f"{'; '.join(hinted[cell['cell']])}. Run `isaac bundle template adopt` again: it "
                                             "lists this as a disagreement to rule on")})
            if cell.get("adopted"):
                continue
        readings = file_readings(cell, decision, kind, obj)
        if readings and cell.get("adopted"):
            # adopted because no file reading contradicted it THEN; one does now. The scientist never typed
            # this cell here, so it is not recorded until they rule (a fresh adopt lists it as a disagreement).
            origin = cell["adopted"]
            contradiction(cell, obj, readings, "awaiting_ruling")
            awaiting.append({"id": None, "cell": cell["cell"], "header": cell["header"], "target_field_path": cell["path"],
                             "run_id": obj.get("run_id"), "target": {"kind": kind, "id": tid}, "their_value": cell["value"],
                             "why": (f"adopted from {origin.get('their_file')} {origin.get('their_cell')}, but a file reading "
                                     f"now contradicts it: {'; '.join(readings)}. Run `isaac bundle template adopt` again: "
                                     "it lists this as a disagreement to rule on")})
            continue
        call = calls.setdefault((cell["cell"], decision), {"cell": cell, "decision": decision, "tids": []})
        call["tids"].append(tid)
        report["decisions"].append({"cell": cell["cell"], "sheet": cell["sheet"], "header": cell["header"],
                                    "target": {"kind": kind, "id": tid}, "target_field_path": cell["path"],
                                    "run_id": obj.get("run_id"), "decision": decision, "value": cell["value"]})
        contradiction(cell, obj, readings, "recorded_as_typed")

    recorded: list[str] = []
    unchanged: list[str] = []
    refused_targets: set[str] = set()
    for call in calls.values():
        cell, decision, tids = call["cell"], call["decision"], call["tids"]
        try:
            res = M.record_decision(bundle, target_ids=tids, decision=decision, value=cell["value"], by=by, orcid=orcid,
                                    note=cell["note"])
        except M.MutationRefused as exc:
            report["refused"].append({"cell": cell["cell"], "targets": tids, "decision": decision, "reason": str(exc)})
            refused_targets.update(tids)
            continue
        recorded += res["recorded"]
        unchanged += res["unchanged"]
    report["decisions"] = [d for d in report["decisions"] if d["target"]["id"] not in refused_targets]
    # a per-run slot whose answer was not recorded (refused, a clash) is never left behind as an empty question
    gone = set(M.drop_unrecorded_slots(bundle, slots_added))
    report["decisions"] = [d for d in report["decisions"] if d["target"]["id"] not in gone]
    kept = {(d["cell"], d.get("run_id") or "") for d in report["decisions"]}
    kept_cells = {d["cell"] for d in report["decisions"]}
    report["contradictions"] = [c for c in report["contradictions"]
                                if c["outcome"] == "awaiting_ruling" or (c["cell"], c.get("run_id") or "") in kept
                                or (c.get("hint") and c["cell"] in kept_cells)]
    report["ruled_per_run"] = _ruled_per_run_summary(bundle, report["ruled_per_run"], report["decisions"])
    # c7-r HIGH-1: A BEAMTIME-WIDE ANSWER DOES NOT REACH A GROUP THAT KEEPS ITS OWN. A Campaign Info value recorded on
    # a beamtime-wide target fills every group but one whose run holds a differing answer of its own (an answer slot
    # from `--only group:<id>` or a ruling, or its own target): that group's record would hold two values and fill
    # nothing. It is NAMED, with the command that changes it — never counted as reached, nor only listed as a clash.
    report["not_reached"] = []
    recorded_once = {(d["cell"], d["target_field_path"]): d["value"] for d in report["decisions"] if d.get("run_id") is None}
    if recorded_once:
        view = B.apply_derivations(copy.deepcopy(bundle))  # the states the save derives
        for (cell_name, path), value in recorded_once.items():
            asked = beamtime_question(path)
            for k in M.groups_keeping_their_own_answer(view, path, value,
                                                       question_id=asked["question_id"] if asked else None):
                report["not_reached"].append({"cell": cell_name, "target_field_path": path, "value": value, **k,
                                              "reason": M.not_reached_sentence(path, [k], lead="")})
    report["reaction_needs"] = reaction_followups(bundle, report["decisions"])
    report["awaiting_ruling"] = awaiting
    report["recorded_target_ids"] = recorded
    report["unchanged_target_ids"] = unchanged
    return report


REACTION_PATHS = ("context.electrochemistry.reaction", "context.reaction.name")


def _run_group_of(bundle: Mapping) -> dict[str, str]:
    """``run_id -> group_id`` of the first live group holding it."""
    by_key = _live_group_of(bundle)
    return {r["run_id"]: by_key[r["run_key"]] for r in bundle["runs"] if r["run_key"] in by_key}


def reaction_lacking(bundle: Mapping, run_ids: Sequence[str], answered_now: set[tuple[str, str]]) -> dict | None:
    """L3: a run that states a reaction needs its drive and catalysis (``bundle.REACTION_REQUIRED_PATHS``). For the
    runs ``run_ids`` that get a reaction now: ``{fields, runs, groups}`` still unanswered — no scientist-confirmed
    value for the run (or the beamtime) and none in ``answered_now`` (``(run_id, path)``) — or ``None``."""
    group_of = _run_group_of(bundle)
    objs = [*(bundle.get("candidates") or ()), *(bundle.get("missing_values") or ())]
    confirmed = {(o.get("run_id"), o.get("target_field_path")) for o in objs
                 if o.get("value_state") == B.STATE_SCIENTIST_CONFIRMED}
    lacking: dict[str, set[str]] = {}
    for rid in dict.fromkeys(run_ids):
        for path in B.REACTION_REQUIRED_PATHS:
            if (rid, path) in answered_now or (rid, path) in confirmed or (None, path) in confirmed:
                continue
            lacking.setdefault(path, set()).add(rid)
    if not lacking:
        return None
    runs = set().union(*lacking.values())
    return {"fields": sorted(lacking), "runs": len(runs), "groups": len({group_of[r] for r in runs if r in group_of})}


def reaction_followups(bundle: Mapping, decisions: Sequence[Mapping]) -> dict | None:
    """:func:`reaction_lacking` for the runs an import gives a reaction."""
    runs = [d["run_id"] for d in decisions if d.get("target_field_path") in REACTION_PATHS and d.get("run_id")]
    if not runs:
        return None
    return reaction_lacking(bundle, runs, {(d.get("run_id"), d.get("target_field_path")) for d in decisions})


def _reaction_needs_line(x: Mapping, verb: str) -> str:
    return (f"A reaction {verb} makes {' and '.join(x['fields'])} required for {x['groups']} group(s) ({x['runs']} run(s)) "
            "that have no answer for it yet: fill 'Reaction Drive' and 'Reaction Catalysis' (File List) for those runs, or "
            "answer their questions (`isaac bundle questions <bundle_dir>`).")


def _ruled_per_run_summary(bundle: Mapping, entries: Sequence[Mapping], decisions: Sequence[Mapping]) -> list[dict]:
    """One entry per ruling cell recorded per run for a field asked once for the beamtime: how many runs it reached,
    how many answer slots it added, and how many groups the beamtime-wide question still covers."""
    if not entries:
        return []
    recorded = {(d["cell"], d.get("run_id")) for d in decisions}
    group_of = _run_group_of(bundle)
    total = len({g for g in group_of.values()})
    out: dict[tuple, dict] = {}
    for x in entries:
        if (x["cell"], x["run_id"]) not in recorded:
            continue
        e = out.setdefault((x["cell"], x["question_id"]), {k: x[k] for k in ("cell", "header", "target_field_path",
                                                                           "question_id", "answered")} | {"runs": 0, "slots_added": 0})
        e["runs"] += 1
        e["slots_added"] += x["slots_added"]
    for e in out.values():
        touched = {group_of.get(d.get("run_id")) for d in decisions if d.get("target_field_path") == e["target_field_path"]}
        e["groups_left_to_the_question"] = max(total - len(touched - {None}), 0)
    return list(out.values())


def off_list_requests(cells: Sequence[Mapping]) -> list[dict]:
    """Off-list cells (``read_cells`` ``off_list``; adopt's not-allowed values) as ONE request per field and word:
    ``{field_path, term, header, cells, run_ids, scope}``. A cell of Campaign Info is about the whole beamtime."""
    out: dict[tuple, dict] = {}
    for c in cells:
        row = out.setdefault((c["path"], c["value"]), {"field_path": c["path"], "term": c["value"], "header": c.get("header"),
                                                       "cells": [], "run_ids": [], "scope": "run"})
        for cell in c.get("cells") or [c["cell"]]:
            if cell not in row["cells"]:
                row["cells"].append(cell)
        for rid in c.get("run_ids") or ([c["run_id"]] if c.get("run_id") else []):
            if rid not in row["run_ids"]:
                row["run_ids"].append(rid)
    for row in out.values():
        row["scope"] = "run" if row["run_ids"] else "experiment"
    return list(out.values())


def keep_off_list_cells(bundle: dict, cells: Sequence[Mapping], *, by: str | None, orcid: str | None, file: str,
                        kind: str, command: str) -> list[dict]:
    """Keep each off-list word of a sheet as a vocabulary request on ``bundle`` (in memory; the caller saves):
    the rows of :func:`off_list_requests`, each with its ``request_id`` and whether it is ``new``."""
    from . import mutations as M

    kept = []
    for row in off_list_requests(cells):
        request, new = M.keep_vocabulary_request(
            bundle, field_path=row["field_path"], term=row["term"], by=by, orcid=orcid, run_ids=row["run_ids"],
            scope=row["scope"], source={"kind": kind, "command": command, "file": file, "cells": row["cells"]})
        kept.append({**row, "request_id": request["request_id"], "new": new})
    return kept


def _request_lines(items: Sequence[Mapping], applied: bool) -> list[str]:
    """One line per kept (or, in a preview, to-be-kept) vocabulary request."""
    from . import propose_terms as P

    out = []
    for r in items:
        reach = f"{len(r['run_ids'])} run(s)" if r["run_ids"] else "the whole beamtime"
        more = f" and {len(r['cells']) - _CELLS_SHOWN} more" if len(r["cells"]) > _CELLS_SHOWN else ""
        tag = (f"{r['request_id']}{'' if r.get('new') else ', kept earlier'}" if applied and r.get("request_id") else
               "would be kept")
        same = P.equivalent_term(r["field_path"], r["term"])
        out.append(f"  {r['header']} {r['term']!r} for {reach} ({', '.join(r['cells'][:_CELLS_SHOWN])}{more}) [{tag}]"
                   + (f"; use `{same}` (an existing term)" if same else ""))
    return out


def import_template(bundle_dir: Path, xlsx: Path, *, by: str, orcid: str | None = None,
                    apply: bool = False, supersede: bool = False) -> dict:
    """Read a filled template. Preview (default) writes nothing; ``apply`` records."""
    from . import mutations as M

    problem = B.decider_name_problem(by)
    if problem:
        raise TemplateRefused(f"--by is required: the name of the person whose answers these are (this one {problem}).")
    if orcid is not None and not M.ORCID_RE.fullmatch(orcid):
        raise TemplateRefused("--orcid must look like 0000-0000-0000-0000 (last character may be X).")
    xlsx = Path(xlsx).expanduser()
    wb = _read_workbook(xlsx)
    marker = _marker(wb)
    meta = marker["meta"]
    if meta.get("format") != FORMAT_ID:
        raise TemplateRefused(f"The `_isaac` sheet names format {meta.get('format')!r}, not {FORMAT_ID}. Run "
                              "`isaac bundle template export` and fill that file. Nothing was recorded.", code=EXIT_INPUT_REFUSED)
    if str(meta.get("format_version") or "").split(".")[0] != FORMAT_VERSION.split(".")[0]:
        raise TemplateRefused(f"Template version {meta.get('format_version')!r} is not readable by this build "
                              f"({FORMAT_VERSION}). Re-export the template. Nothing was recorded.", code=EXIT_INPUT_REFUSED)
    try:
        current = M.load(Path(bundle_dir))
    except M.MutationRefused as exc:
        raise TemplateRefused(str(exc), code=exc.code) from exc
    if meta.get("bundle_id") != current["bundle_id"]:
        raise TemplateRefused(
            f"This template was exported from bundle {meta.get('bundle_id')!r}, not this one "
            f"({current['bundle_id']}). Run `isaac bundle template export {bundle_dir} --out <file.xlsx>` and fill "
            "that file. Nothing was recorded.", code=EXIT_INPUT_REFUSED)

    def run(bundle):
        cells = read_cells(wb, marker, bundle, xlsx.name)
        report = plan_and_record(bundle, cells, by=by, orcid=orcid, supersede=supersede)
        return cells, report

    requests: list[dict] = []
    if apply:
        with M.locked(Path(bundle_dir)):
            bundle = M.load(Path(bundle_dir))
            cells, report = run(bundle)
            requests = keep_off_list_cells(bundle, cells["off_list"], by=by, orcid=orcid, file=xlsx.name,
                                           kind="template_cell", command="bundle template import")
            if report["recorded_target_ids"] or any(r["new"] for r in requests):  # nothing new: not rewritten
                M.save_bundle(Path(bundle_dir), bundle)
    else:
        cells, report = run(current)  # in memory only; nothing is saved
        requests = [{**r, "request_id": None, "new": None} for r in off_list_requests(cells["off_list"])]
    return _as_text_dates({  # an invalid cell is shown as read, and a date cell reads as a datetime
        "applied": apply,
        "template": xlsx.name,
        "bundle_id": current["bundle_id"],
        "exported_utc": meta.get("exported_utc"),
        "cells_read": cells.get("read", 0),
        "decisions": report["decisions"],
        "replaces_or_contradicts_a_file_reading": report["contradictions"],
        "awaiting_ruling": cells["awaiting_ruling"] + report["awaiting_ruling"],
        "recorded": len(report["recorded_target_ids"]) if apply else 0,
        "would_record": len(report["recorded_target_ids"]),
        "unchanged": len(report["unchanged_target_ids"]),
        "already_decided_same_value": report["already"],
        "invalid_cells": cells["invalid"],
        # an invalid cell refused ONLY because its word is not on the field's list: kept as a vocabulary request
        "vocabulary_requests": requests,
        "clashes": report["clashes"],
        "refused": report["refused"],
        "unmapped_columns": list(cells["unmapped"].values()),
        "recorded_nowhere": report["nowhere"],
        # a run's cell for a field this bundle asks once for the beamtime: the question that answers it
        "asked_beamtime_wide": report["asked_beamtime_wide"],
        # a ruling for a field asked once for the beamtime: recorded per run, for exactly the runs its row names
        "ruled_per_run": report["ruled_per_run"],
        # a reaction recorded here makes drive and catalysis required where nothing answers them yet
        "reaction_needs": report["reaction_needs"],
        "rows_refused": cells["rows_refused"],
        "extracted_values_left_alone": report["left_alone"],
        "campaign_not_filled": report["campaign_not_filled"],
        # a beamtime-wide answer recorded here that a group's own, differing answer keeps from its record
        "not_reached": report["not_reached"],
        "statement": ("Recorded as scientist decisions (identity: self_reported_local_user, unverified)."
                      if apply else "PREVIEW: nothing was written. Add --apply to record these decisions."),
    })


#: How many cells a compact report names per column before it says how many more.
_CELLS_SHOWN = 3


def _decision_lines(rep: Mapping) -> list[str]:
    """One line per Campaign Info cell, Configurations cell, or File List column: how many
    decisions it gives, of which kind, to which field."""
    groups: dict[str, dict] = {}
    for d in rep["decisions"]:
        sheet = d.get("sheet") or ""
        if sheet == SHEET_FILES:
            key = f"{SHEET_FILES} / {d.get('header')}"
        else:
            key = f"{d['cell'].split(' via ', 1)[0]} {d.get('header') or ''}".rstrip()
        g = groups.setdefault(key, {"path": d["target_field_path"], "n": 0, "kinds": {}, "runs": set(), "values": set(),
                                    "per_row": sheet == SHEET_FILES or " via " in d["cell"]})
        g["n"] += 1
        g["kinds"][d["decision"]] = g["kinds"].get(d["decision"], 0) + 1
        g["values"].add(repr(d["value"]))
        if d.get("run_id"):
            g["runs"].add(d["run_id"])
    out = []
    for key, g in groups.items():
        kinds = ", ".join(f"{k} {n}" for k, n in sorted(g["kinds"].items()))
        rows = f" over {len(g['runs'])} row(s)" if g["per_row"] else ""
        value = f" = {next(iter(g['values']))}" if len(g["values"]) == 1 else f"; {len(g['values'])} distinct values"
        out.append(f"  {key} -> {g['path']}: {g['n']} decision(s){rows} ({kinds}){value}")
    return out


def _grouped_cells(items: Sequence[Mapping], *, key: str = "header") -> list[str]:
    """One line per column (and reason) with a few of its cells named, for a long list of cells."""
    groups: dict[tuple, list[str]] = {}
    for x in items:
        groups.setdefault((x.get(key), x.get("reason")), []).append(x.get("cell") or f"row {x.get('row')}")
    out = []
    for (header, reason), cells in groups.items():
        more = f" and {len(cells) - _CELLS_SHOWN} more" if len(cells) > _CELLS_SHOWN else ""
        out.append(f"  {header}: {len(cells)} cell(s) ({', '.join(cells[:_CELLS_SHOWN])}{more}): {reason}")
    return out


def _beamtime_lines(items: Sequence[Mapping]) -> list[str]:
    """One line per column, question and command: the cells, the question asking the field once for the
    beamtime, and the command that answers it (with the cell's value, when the cells agree)."""
    groups: dict[tuple, dict] = {}
    for x in items:
        key = (x["header"], x["target_field_path"], x["question_id"], x["answered"], x["command"], tuple(x.get("values_differ") or ()))
        g = groups.setdefault(key, {"cells": [], "runs": set()})
        cell = x["cell"].split(" via ", 1)[0]
        if cell not in g["cells"]:  # one cell (a Technique reached through many rows) is named once
            g["cells"].append(cell)
        if x.get("run_id"):
            g["runs"].add(x["run_id"])
    out = []
    for (header, path, qid, answered, command, differ), g in groups.items():
        cells = g["cells"]
        more = f" and {len(cells) - _CELLS_SHOWN} more" if len(cells) > _CELLS_SHOWN else ""
        runs = f"{len(g['runs'])} run(s) from " if g["runs"] else ""
        line = f"  {header} -> {path}: {runs}{len(cells)} cell(s) ({', '.join(cells[:_CELLS_SHOWN])}{more}); question {qid}"
        if differ:
            # never one command per value: the scientist says which value applies to the whole beamtime
            line += (f": these cells hold {len(differ)} different values ({', '.join(differ)}), so the scientist must say "
                     "which applies" + (", then change the answer" if answered else ", then answer it") + f": {command}")
        else:
            line += (" (answered; to change it)" if answered else "") + f": {command}"
        out.append(line)
    return out


def _invalid_lines(items: Sequence[Mapping]) -> list[str]:
    """Compact: per column, its allowed values ONCE, then each distinct value with a few of its cells.
    The allowed list is not repeated per cell (``--full`` and ``--json`` keep each cell's whole reason)."""
    by_header: dict[str, dict] = {}
    for x in items:
        reason = str(x.get("reason") or "")
        short, sep, allowed = reason.partition("; allowed: ")
        h = by_header.setdefault(str(x.get("header")), {"allowed": None, "values": {}})
        if sep and h["allowed"] is None:
            h["allowed"] = allowed
        h["values"].setdefault((repr(x.get("value")), short if sep else reason), []).append(str(x.get("cell")))
    out = []
    for header, h in by_header.items():
        out.append(f"  {header}" + (f" (allowed: {h['allowed']}):" if h["allowed"] else ":"))
        for (value, reason), cells in h["values"].items():
            more = f" and {len(cells) - _CELLS_SHOWN} more" if len(cells) > _CELLS_SHOWN else ""
            reason = reason.replace(f"{value} is not an allowed value", "not an allowed value")
            out.append(f"    {value} in {len(cells)} cell(s) ({', '.join(cells[:_CELLS_SHOWN])}{more}): {reason}")
    return out


def render_report(rep: Mapping, *, full: bool = False) -> str:
    """Plain text for the terminal. COMPACT by default: one line per cell or column with its number
    of decisions, then every invalid cell, clash, refusal and refused row in full, and the totals.
    ``full``: one line per decision and per cell, as ``--json`` holds them."""
    lines = []
    add = lines.append
    add(("APPLIED" if rep["applied"] else "PREVIEW (nothing written)") + f": template {rep['template']} -> bundle {rep['bundle_id']}")
    add("")
    verb = "recorded" if rep["applied"] else "would record"
    if full:
        add(f"Decisions ({verb}):")
        for d in rep["decisions"]:
            where = f" {d['run_id']}" if d.get("run_id") else ""
            add(f"  {d['cell']}: {d['decision']} {d['target']['kind']} {d['target']['id']}{where} {d['target_field_path']} = {d['value']!r}")
    else:
        add(f"Decisions ({verb}), one line per cell or column:")
        lines += _decision_lines(rep)
    if not rep["decisions"]:
        add("  (none)")
    if rep.get("replaces_or_contradicts_a_file_reading"):
        add("")
        add("Replaces or contradicts a file reading (one line each; read each back before --apply):")
        for x in rep["replaces_or_contradicts_a_file_reading"]:
            run = f" {x['run_id']}" if x.get("run_id") else ""
            if x.get("outcome") == "awaiting_ruling":
                said = "not recorded: adopted from your sheet before this reading, awaiting your ruling"
            else:
                said = ("recorded" if rep["applied"] else "would be recorded") + ": you typed it"
            add(f"  {x['cell']}{run} {x['target_field_path']} = {x['value']!r} [{said}]: {'; '.join(x['readings'])}")
    if rep.get("awaiting_ruling"):
        add("")
        add("Disagreements awaiting your ruling (nothing recorded for them):")
        for x in rep["awaiting_ruling"]:
            run = f" {x['run_id']}" if x.get("run_id") else ""
            add(f"  {x['id'] or '(not yet on the Disagreements sheet)'} {x['cell']}{run} {x['header']}: theirs {x['their_value']!r}; {x['why']}")
    in_full = [("Clashes (skipped)", "clashes"), ("Refused", "refused"), ("Rows refused", "rows_refused")]
    if full:
        in_full.insert(0, ("Invalid cells (recorded as nothing)", "invalid_cells"))
        in_full.insert(3, ("Recorded nowhere (no target for that run/field)", "recorded_nowhere"))
    elif rep["invalid_cells"]:
        add("")
        add("Invalid cells (recorded as nothing):")
        lines += _invalid_lines(rep["invalid_cells"])
    if rep.get("vocabulary_requests"):
        add("")
        add(("Kept" if rep["applied"] else "With --apply, kept") + " as vocabulary requests (not decisions, not values: "
            "each is a word its field's list lacks; `isaac bundle propose-terms` can draft it):")
        lines += _request_lines(rep["vocabulary_requests"], rep["applied"])
    for title, key in in_full:
        if rep[key]:
            add("")
            add(f"{title}:")
            for x in rep[key]:
                add(f"  {x.get('cell', 'row ' + str(x.get('row')))}: {x.get('reason')}"
                    + (f" (value {x['value']!r})" if "value" in x else ""))
    if rep["recorded_nowhere"] and not full:
        add("")
        add("Recorded nowhere (no target for that run/field):")
        lines += _grouped_cells(rep["recorded_nowhere"])
    if rep.get("asked_beamtime_wide"):
        add("")
        add("Asked once for the whole beamtime (these cells are not recorded; the question's command answers it):")
        if full:
            for x in rep["asked_beamtime_wide"]:
                add(f"  {x['cell']} {x.get('run_id') or ''} {x['header']} = {x['value']!r}: {x['reason']}")
        else:
            lines += _beamtime_lines(rep["asked_beamtime_wide"])
    if rep.get("ruled_per_run"):
        add("")
        add("Rulings for a field asked once for the whole beamtime (recorded for exactly the runs each row names):")
        for x in rep["ruled_per_run"]:
            slots = f" ({x['slots_added']} answer slot(s) added)" if x["slots_added"] else ""
            add(f"  {x['cell']} {x['header']} -> {x['target_field_path']}: recorded per run for {x['runs']} run(s) that have no "
                f"target of their own for it{slots}. Question {x['question_id']} "
                + ("(answered) " if x["answered"] else "") + f"still covers {x['groups_left_to_the_question']} group(s) this "
                "import records nothing for; if its beamtime-wide answer differs from this ruling, these runs' groups get two "
                "values and fill nothing for it.")
    if rep.get("not_reached"):
        add("")
        add("Not reached by a beamtime-wide answer (the group keeps its own answer, so its record fills nothing for the "
            "field until the two agree):")
        for x in rep["not_reached"]:
            add(f"  {x['cell']} {x['target_field_path']} = {x['value']!r}: {x['reason']}")
    if rep.get("reaction_needs"):
        add("")
        add("Then required: " + _reaction_needs_line(rep["reaction_needs"], "recorded by this import"))
    if rep["unmapped_columns"]:
        add("")
        add("Unmapped columns:")
        for u in rep["unmapped_columns"]:
            if full:
                add(f"  {u['sheet']} / {u['header']}: {u['reason']}")
                for c in u["cells"]:
                    add(f"    {c['cell']} = {c['value']!r}")
            else:
                cells = [c["cell"] for c in u["cells"]]
                more = f" and {len(cells) - _CELLS_SHOWN} more" if len(cells) > _CELLS_SHOWN else ""
                add(f"  {u['sheet']} / {u['header']}: {len(cells)} cell(s) ({', '.join(cells[:_CELLS_SHOWN])}{more}); {u['reason']}")
    if rep.get("campaign_not_filled"):
        add("")
        add("Campaign Info answers not filled for runs outside the File List:")
        for x in rep["campaign_not_filled"]:
            add(f"  {x['cell']} {x['target_field_path']}: {x['targets']} target(s) (runs {', '.join(map(str, x['run_ids']))}); {x['reason']}")
    if rep["extracted_values_left_alone"]:
        add("")
        add("Extracted values left alone:")
        for x in rep["extracted_values_left_alone"]:
            add(f"  {x['cell']} {x['target_field_path']}: {x['candidates']} extracted value(s); {x['reason']}")
    add("")
    add(
        f"Summary: {rep['cells_read']} cells read; {rep['recorded'] if rep['applied'] else rep['would_record']} decision(s) "
        f"{verb}; {rep['unchanged']} unchanged; {len(rep['already_decided_same_value'])} already decided with the same value; "
        f"{len(rep['invalid_cells'])} invalid; {len(rep['clashes'])} clash(es); {len(rep['refused'])} refused; "
        f"{sum(len(u['cells']) for u in rep['unmapped_columns'])} cell(s) in {len(rep['unmapped_columns'])} unmapped column(s); "
        f"{len(rep['recorded_nowhere'])} recorded nowhere; {len(rep['rows_refused'])} row(s) refused"
        + (f"; {len({x['run_id'] for x in rep['asked_beamtime_wide']})} run(s) from "
           f"{len({x['cell'].split(' via ', 1)[0] for x in rep['asked_beamtime_wide']})} cell(s) asked once for the whole beamtime (its question answers them)"
           if rep.get("asked_beamtime_wide") else "")
        + (f"; {len(rep['replaces_or_contradicts_a_file_reading'])} replace or contradict a file reading"
           if rep.get("replaces_or_contradicts_a_file_reading") else "")
        + (f"; {len(rep['awaiting_ruling'])} disagreement(s) awaiting a ruling" if rep.get("awaiting_ruling") else "")
        + (f"; {len({x['group_id'] for x in rep['not_reached']})} group(s) not reached by a beamtime-wide answer"
           if rep.get("not_reached") else "")
        + "."
    )
    add(rep["statement"])
    if not full:
        add("(--full lists every decision and cell; --json prints the whole report.)")
    return "\n".join(lines)


# --- adopt: the scientist's own filled copy of ISAAC's template ------------------------------------
#
# ``isaac bundle template adopt <bundle> <their.xlsx> --out <file.xlsx>``. A scientist who already filled
# ISAAC's metadata template (no ``_isaac`` sheet, File List keyed by their own folder names) should not
# retype it. Adopt writes the bundle's own template (:func:`export_template`) and copies THEIR value into
# an answer cell only when no file reading contradicts it; every value a reading contradicts, or that the
# field does not take, goes to the Disagreements sheet with an empty "Your ruling" column instead, so
# nothing that contradicts a file is imported until the scientist rules. ``template import`` then reads
# the answer cells and the rulings, and nothing else.

#: Their header (whitespace-collapsed, case-folded) -> this template's File List column, beyond an exact
#: header. ISAAC's template keys File List by "File Name"; filled copies key it by their folders.
THEIR_FILE_HEADER_ALIASES: Mapping[str, str] = {"folder name": "File Name", "sample": "Sample Name"}
_TAKEN_NAME_SUFFIX = "_dir"

#: ISAAC's BLANK metadata template (``ISAAC_Metadata_Template.xlsx`` as it is handed out) holds these example
#: values in Campaign Info (column B, by the label in column A) and these two example Configurations rows
#: (by header). Import never mistakes it for answers because it has no ``_isaac`` sheet (:func:`_marker`);
#: adopt reads exactly such workbooks, so it recognises the blank by these values: a copy whose Campaign
#: Info and Configurations still hold exactly them, and nothing else, is the blank, whatever its File List
#: holds (a later blank may ship example rows there).
BLANK_TEMPLATE_CAMPAIGN_EXAMPLES: Mapping[str, Any] = {
    "Facility Name": "SSRL", "Organization": "SLAC National Accelerator Laboratory", "Beamline/Instrument": "15-2",
    "Start Date": "2026-01-01", "Proposal ID": None}
BLANK_TEMPLATE_CONFIG_EXAMPLES: tuple[Mapping[str, Any], ...] = (
    {"Config ID": "Config_A", "Technique": "XAS", "Mode": "Fluorescence", "Primary Signal Channel": "VortDT",
     "Normalization Channel": "I0", "Detectors": "Canberra 100", "Energy Calibration": "Cu-Foil deriv=0 -> 8979eV",
     "Beamline Setup": "Standard Setup"},
    {"Config ID": "Config_B", "Technique": "XAS", "Mode": "Transmission", "Primary Signal Channel": "I1",
     "Normalization Channel": "I0", "Detectors": "Ion Chambers", "Energy Calibration": "Cu-Foil deriv=0 -> 8979eV",
     "Beamline Setup": "Changed detector"},
)


def _norm_header(value: Any) -> str:
    return " ".join(str(value or "").split()).casefold()


def _blank_value(value: Any) -> Any:
    """A cell as the blank's example values are compared: text trimmed, a date as its ISO day, empty = None."""
    from datetime import date

    if isinstance(value, datetime) and value.time() == datetime.min.time():
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip() or None
    return value


def _still_the_blank(wb) -> bool:
    """Campaign Info and Configurations hold exactly the blank's example values and nothing else."""
    from collections import Counter

    def pairs(items):
        return {(_norm_header(k), _blank_value(v)) for k, v in items if _blank_value(v) is not None}

    wc = wb[SHEET_CAMPAIGN]
    campaign = pairs((wc.cell(row=r, column=1).value, wc.cell(row=r, column=2).value) for r in range(2, wc.max_row + 1))
    if campaign != pairs(BLANK_TEMPLATE_CAMPAIGN_EXAMPLES.items()):
        return False
    wcf = wb[SHEET_CONFIG]
    header = [c.value for c in wcf[1]]
    rows = Counter(frozenset(p) for row in wcf.iter_rows(min_row=2) if (p := pairs(zip(header, (c.value for c in row)))))
    return rows == Counter(frozenset(pairs(r.items())) for r in BLANK_TEMPLATE_CONFIG_EXAMPLES)


def _as_text_dates(obj: Any) -> Any:
    """``obj`` with every date or datetime as its ISO text (a midnight datetime as its day), so the report
    prints as JSON. Nothing else changes."""
    from datetime import date

    if isinstance(obj, Mapping):
        return {k: _as_text_dates(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_as_text_dates(v) for v in obj]
    if isinstance(obj, date):
        return _blank_value(obj)
    return obj


def _refuse_unofficial(name: str, why: str) -> TemplateRefused:
    return TemplateRefused(
        f"{name} is not a filled copy of ISAAC's metadata template: {why}. `template adopt` reads only that template "
        f"(sheets {SHEET_CAMPAIGN!r}, {SHEET_CONFIG!r} and {SHEET_FILES!r}; {CAMPAIGN_TITLE!r} in {SHEET_CAMPAIGN} A1; "
        "a File List keyed by 'File Name' or 'Folder Name'). Nothing was written.", code=EXIT_INPUT_REFUSED)


def _their_sheet(wb, name: str) -> dict:
    """The parts of a filled official template adopt reads, or a plain refusal."""
    if SHEET_MARKER in wb.sheetnames:
        raise TemplateRefused(
            f"{name} has an `_isaac` sheet: it is a template ISAAC exported for a bundle, not the scientist's own "
            f"copy. Read it with `isaac bundle template import <bundle> {name} --by \"<name>\"`. Nothing was written.",
            code=EXIT_INPUT_REFUSED)
    missing = [s for s in (SHEET_CAMPAIGN, SHEET_CONFIG, SHEET_FILES) if s not in wb.sheetnames]
    if missing:
        raise _refuse_unofficial(name, f"it has no {', '.join(repr(m) for m in missing)} sheet")
    if str(wb[SHEET_CAMPAIGN]["A1"].value or "").strip() != CAMPAIGN_TITLE:
        raise _refuse_unofficial(name, f"{SHEET_CAMPAIGN} A1 reads {wb[SHEET_CAMPAIGN]['A1'].value!r}")
    wf = wb[SHEET_FILES]
    header = [c.value for c in wf[1]]
    if not header or _norm_header(header[0]) not in ("file name", "folder name"):
        raise _refuse_unofficial(name, f"File List A1 reads {header[0] if header else None!r}")
    names = [(r, wf.cell(row=r, column=1).value) for r in range(2, wf.max_row + 1)]
    names = [(r, str(v).strip()) for r, v in names if v is not None and str(v).strip()]
    if not names:
        raise TemplateRefused(
            f"{name} is ISAAC's blank metadata template (its File List names no file or folder): its other cells are "
            "the template's example values, not anyone's answers. Nothing was written.", code=EXIT_INPUT_REFUSED)
    if _still_the_blank(wb):
        raise TemplateRefused(
            f"{name} is ISAAC's blank metadata template: its Campaign Info and Configurations hold exactly the "
            "template's example values (Facility Name 'SSRL', Beamline/Instrument '15-2', Start Date 2026-01-01, "
            "Config_A and Config_B), so its File List rows are taken as examples too, not anyone's answers. If you "
            "did fill it, put your own values in Campaign Info and Configurations (or clear the examples) and adopt "
            "it again. Nothing was written.", code=EXIT_INPUT_REFUSED)
    return {"files": wf, "header": header, "names": names}


def _match_runs(bundle: Mapping, their_name: str) -> list[str]:
    """The run ids a File List row names: a run whose file sits DIRECTLY in a folder of that name, or whose
    file or stem IS that name (``<name>_dir`` also names the run ``<name>``: SPEC keeps the scan file beside
    its ``_dir`` folder). An ancestor folder (``raw``, a day's folder above the runs' own) names no run, so a
    row never takes every run below it. Exact names only: nothing is case-folded or guessed."""
    want = {their_name}
    if their_name.endswith(_TAKEN_NAME_SUFFIX) and len(their_name) > len(_TAKEN_NAME_SUFFIX):
        want.add(their_name[: -len(_TAKEN_NAME_SUFFIX)])
    sources = {s["source_id"]: s for s in bundle["sources"]}
    out = []
    for run in bundle["runs"]:
        rel = (sources.get(run.get("acquisition_source_id")) or {}).get("relative_path") or ""
        parts = [p for p in rel.split("/") if p]
        last = parts[-1] if parts else ""
        if ((len(parts) > 1 and parts[-2] == their_name) or rel in want or last in want or str(run.get("stem") or "") in want
                or (last.rsplit(".", 1)[0] in want if "." in last else False)):
            out.append(run["run_id"])
    return out


class _Adopter:
    """Reads their sheet against one bundle. Every reading compared is a file's (a parser value, a
    file-name token, an AI reading of a file); an AI reading OF THEIR SHEET is their sheet, so it is
    never compared."""

    def __init__(self, bundle: Mapping, index: _Index, rules: Mapping, their_file: str):
        self.bundle, self.index, self.rules, self.their_file = bundle, index, rules, their_file
        self.runs = {r["run_id"]: r for r in bundle["runs"]}
        self.ai_of: dict[tuple[str, str], list[dict]] = {}
        for i in bundle.get("ai_interpretations") or ():
            t = i.get("target") or {}
            self.ai_of.setdefault((str(t.get("kind")), str(t.get("id"))), []).append(i)
        self.exposed: list[tuple[str, str]] = []
        self.sheet_readings_skipped = 0

    def _is_their_sheet(self, i: Mapping) -> bool:
        paths = [i.get("source_path")] + [(self.index.sources.get(r.get("source_id") or "") or {}).get("relative_path")
                                          for r in i.get("source_refs") or ()]
        return any(p and Path(str(p)).name == self.their_file for p in paths)

    def _comparable(self, value: Any, path: str) -> Any:
        rule = self.rules.get(path) or {}
        if rule.get("numeric"):
            if isinstance(value, str) and _NUMBER_RE.match(value.strip()):
                return float(value.strip())
            return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None
        return value.strip() if isinstance(value, str) and value.strip() else None

    def readings(self, run_id: str | None, path: str) -> list[dict]:
        """``[{value, kind, source, target}]``: what the files say for one run (or, ``run_id=None``, for
        the whole beamtime), one entry per distinct value and source."""
        if run_id is None:
            cands = self.index.cands_by_path.get(path, [])
            missing = self.index.missing_by_path.get(path, [])
        else:
            cands = self.index.cands_by_run.get((run_id, path), []) + (
                self.index.cands_by_run.get((None, path), []) if path == "system.technique" else [])
            missing = self.index.missing_by_run.get((run_id, path), [])
        out: dict[tuple, dict] = {}
        for c in cands:
            if c["value"] is None:
                continue
            ev = next((e for e in c["evidence"] if e["method"] != "ai_interpretation"), None)
            out.setdefault((repr(c["value"]), "extracted"), {
                "value": c["value"], "kind": "extracted", "source": f"extracted from {_ev_where(self.index, ev)}",
                "target": ("candidate", c["candidate_id"])})
        if run_id is not None:
            for t in name_token_readings(run_names(self.bundle, self.runs[run_id]), path):
                out.setdefault((repr(float(t["value"])), "token"), {
                    "value": t["value"], "kind": "file_name_token", "source": f"file name token {t['token']} in {t['name']}"})
        targets = [("candidate", c["candidate_id"]) for c in cands] + [("missing_value", m["value_id"]) for m in missing]
        for kind, tid in targets:
            for i in self.ai_of.get((kind, tid), []):
                if self._is_their_sheet(i):
                    self.sheet_readings_skipped += 1
                    continue
                value = self._comparable(i.get("proposed_value"), path)
                if value is None and path == ELECTROLYTE_CONC_PATH and electrolyte_statement_parts(i.get("proposed_value")):
                    value = i["proposed_value"].strip()  # the whole statement, compared by its concentration part
                if value is None:
                    continue
                where = i.get("source_path") or (self.index.sources.get(((i.get("source_refs") or [{}])[0]).get("source_id") or "") or {}).get("relative_path")
                source = f"AI reading ({(i.get('model') or {}).get('model_id')}) of {where or '(source)'} @ {i.get('excerpt_locator') or '(no locator)'}"
                out.setdefault((repr(value), source), {"value": value, "kind": "ai_reading", "source": source})
        return list(out.values())

    def _contradicts(self, r: Mapping, path: str, value: Any) -> bool:
        """Does reading ``r`` contradict ``value``? The electrolyte is one statement in the files ("0.5 M H2SO4")
        and two cells here, so a name equal to its name part, or a concentration equal to its concentration part,
        agrees; anything else about a name or concentration is a contradiction."""
        parts = electrolyte_statement_parts(r["value"]) if path in (ELECTROLYTE_NAME_PATH, ELECTROLYTE_CONC_PATH) else None
        if parts is not None:
            if path == ELECTROLYTE_NAME_PATH:
                return not (isinstance(value, str) and " ".join(value.split()) == parts[0])
            return not _agrees(parts[1], value)
        v = self._comparable(r["value"], path)
        return v is not None and not _agrees(v, value)

    def against(self, value: Any, run_id: str | None, path: str) -> list[dict]:
        """The readings that CONTRADICT ``value``. A candidate shown here is an exposure."""
        return self.compare(value, run_id, path)[0]

    def compare(self, value: Any, run_id: str | None, path: str) -> tuple[list[dict], list[dict]]:
        """``(contradicting, agreeing)``: EVERY reading of ``path`` for the run, split by whether it contradicts
        ``value`` (a Disagreements row lists both, so the scientist rules seeing all of them). A reading with no
        comparable value is in neither. A contradicting candidate is an exposure (it is always shown); an agreeing
        one is exposed by :meth:`expose` where it is shown."""
        pool = self.readings(run_id, path)
        if path in (ELECTROLYTE_NAME_PATH, ELECTROLYTE_CONC_PATH):
            # the statement is one reading of both cells, whichever of the two fields it was recorded on
            other = ELECTROLYTE_CONC_PATH if path == ELECTROLYTE_NAME_PATH else ELECTROLYTE_NAME_PATH
            seen = {(repr(r["value"]), r["source"]) for r in pool}
            pool += [r for r in self.readings(run_id, other)
                     if electrolyte_statement_parts(r["value"]) and (repr(r["value"]), r["source"]) not in seen]
        bad, good = [], []
        for r in pool:
            if self._contradicts(r, path, value):
                bad.append(r)
            elif self._comparable(r["value"], path) is not None or (
                    path in (ELECTROLYTE_NAME_PATH, ELECTROLYTE_CONC_PATH) and electrolyte_statement_parts(r["value"])):
                good.append(r)
        self.expose(bad)
        return bad, good

    def expose(self, readings: Sequence[Mapping]) -> None:
        for r in readings:
            if r.get("target") and r["target"] not in self.exposed:
                self.exposed.append(r["target"])


def _compared_values(ad: "_Adopter", readings: Sequence[Mapping], path: str) -> tuple[str, ...]:
    """The distinct values ``readings`` hold, as compared (``_comparable``: ``'0'`` and ``0.0`` are one), each a
    ``repr``. A reading with no comparable form keeps its own value; ``0`` / ``0.0`` is a value (L3, c6-r)."""
    return tuple(sorted({repr(r["value"] if (cv := ad._comparable(r["value"], path)) is None else cv) for r in readings}))


def _shown_reading(r: Mapping) -> str:
    return f"{_shown(r['value'])} ({r['source']})"


def _refs_shown(refs: Sequence[str], n: int = _CELLS_SHOWN, *, noun: str = "cells") -> str:
    more = f" and {len(refs) - n} more ({len(refs)} {noun})" if len(refs) > n else ""
    return ", ".join(refs[:n]) + more


def _loose(text: Any) -> str:
    """A value as an unread column is matched to an allowed one: case-folded, every non-alphanumeric dropped
    (``insitu``, ``in-situ`` and ``In Situ`` are all ``in_situ``'s ``insitu``)."""
    return re.sub(r"[^0-9a-z]", "", str(text).casefold())


def _unread_hint_text(h: Mapping, runs: int) -> str:
    rows = f" on {h['rows']} of these {runs} run(s)" if runs and h["rows"] < runs else ""
    return (f"your sheet's column {h['column']} ({h['header']!r}) says {h['value']!r} (= `{h['allowed_value']}`){rows} "
            "— a hint from a column adopt does not read, never adopted")


def _adopt_name_only_reason(value: Any) -> str | None:
    """Why an Electrolyte value that also states a concentration is not adopted, saying where each part of a
    ruling goes on THIS path: the scientist types only rulings here, so the name goes in "Your ruling" and the
    concentration in "Your ruling: concentration (M)". ISAAC never splits it for them."""
    if electrolyte_name_problem(value) is None:
        return None
    m = _CONCENTRATION_IN_NAME.search(value)
    stated = f"{m.group(1)} {m.group(2)}" if m else "a concentration"
    return (f"name only: {value!r} also states a concentration ({stated}); rule with the NAME in 'Your ruling' and the "
            f"concentration in mol/L in '{DISAGREE_HEADERS[9]}' (column {DISAGREE_CONC_COLUMN}), both for this row's runs")


def adopt_template(bundle_dir: Path, their_xlsx: Path, out: Path, *, force: bool = False, suggestions: bool = True,
                   by: str | None = None, orcid: str | None = None) -> dict:
    """Write ``out``: the bundle's template with the scientist's own values where no file contradicts
    them, and a Disagreements sheet for every one a file contradicts or the field does not take. A value
    the field refuses ONLY because it is not on its list is also kept as a vocabulary request (in ``by``'s
    name when given), recorded before the file is written."""
    from . import mutations as M

    if by is not None and (problem := B.decider_name_problem(by)):
        raise TemplateRefused(f"--by must be the name of the person whose sheet it is (this one {problem}).", code=EXIT_USAGE)
    if orcid is not None and not M.ORCID_RE.fullmatch(orcid):
        raise TemplateRefused("--orcid must look like 0000-0000-0000-0000 (last character may be X).", code=EXIT_USAGE)
    their_xlsx = Path(their_xlsx).expanduser()
    wb_theirs = _read_workbook(their_xlsx)
    theirs = _their_sheet(wb_theirs, their_xlsx.name)
    their_file = their_xlsx.name

    def fill(*, wb, bundle, rows, index, rules, marker) -> dict:
        from openpyxl.comments import Comment
        from openpyxl.styles import PatternFill

        flag = PatternFill("solid", fgColor="FFE699")
        ad = _Adopter(bundle, index, rules, their_file)
        our_row = {r["run_id"]: i for i, r in enumerate(rows, start=2)}
        file_col = {c.header: j for j, c in enumerate(FILE_COLUMNS, start=1)}
        result: dict[str, Any] = {"their_file": their_file, "cells_adopted": 0, "disagreements": [], "not_adopted": [],
                                  "unmatched_rows": [], "unmatched_runs": [], "rows_matched": 0, "runs_matched": 0,
                                  "runs_on_several_rows": [], "rows_reaching_several_runs": [], "runs_outside_live_groups": [],
                                  "their_columns_not_read": [],
                                  "their_columns_mapped": []}
        adopted_cells: list[tuple[str, str]] = []  # (our cell, their cell)
        adopted_paths: set[tuple[str | None, str]] = set()  # (run_id, path) of each adopted answer cell
        groups: dict[tuple, dict] = {}  # disagreement key -> entry

        def disagree(key, **entry):
            g = groups.get(key)
            if g is None:
                g = groups[key] = {**entry, "their_cells": [], "runs": [], "answer_cells": [], "readings": [],
                                   "unread_hints": {}}
            for k in ("their_cells", "runs", "answer_cells", "readings"):
                for v in entry.get("_" + k, ()):
                    if v not in g[k]:
                        g[k].append(v)
            for h in entry.get("_unread_hints", ()):
                g["unread_hints"][h] = g["unread_hints"].get(h, 0) + 1
            return g

        wf_theirs = theirs["files"]

        def unread_hints(their_row: int | None, allowed: Sequence | None) -> list[tuple]:
            """M4: ``[(column, header, their text, the allowed value it is)]`` — a column of THEIR row adopt does not
            read that holds, loosely written (``insitu``), one of the field's allowed values (``in_situ``). A hint
            beside the row, never adopted."""
            if their_row is None or not allowed:
                return []
            out = []
            for c in result["their_columns_not_read"]:
                text = wf_theirs[f"{c['column']}{their_row}"].value
                if not isinstance(text, str) or not text.strip():
                    continue
                match = next((a for a in allowed if isinstance(a, str) and _loose(a) == _loose(text)), None)
                if match is not None:
                    out.append((c["column"], c["header"], text.strip(), match))
            return out

        def judge(sheet, our_ref, col: Column, raw: Any, their_ref: str, run_id: str | None, their_row: int | None = None):
            """Adopt one of their cells into ``our_ref``, or record why not. ``run_id=None``: Campaign Info."""
            path = col.path
            rule = rules.get(path or "")
            if isinstance(raw, str) and raw.startswith("="):
                value, why = None, "a formula, not a value"
            else:
                value, why = _parse(raw, rule)
            kind = None
            if not why and path == ELECTROLYTE_NAME_PATH:
                why = _adopt_name_only_reason(value)
                kind = "name_only" if why else None
            runs = [run_id] if run_id else []
            target_cell = _cell_ref(sheet, our_ref)
            if why:
                allowed = rule.get("allowed") if rule else None
                not_allowed = bool(allowed) and isinstance(raw, str) and raw.strip() not in allowed
                kind = kind or ("not_allowed" if not_allowed else "invalid")
                shown_reason = "not an allowed value" if not_allowed else why.partition("; allowed: ")[0]
                disagree((kind, sheet, col.header, repr(raw)), kind=kind, sheet=sheet, header=col.header, path=path,
                         their_value=raw, reason=shown_reason, allowed=allowed if not_allowed else None,
                         scope="experiment" if run_id is None else "run",
                         _their_cells=[their_ref], _runs=runs, _answer_cells=[target_cell],
                         _unread_hints=unread_hints(their_row, allowed) if not_allowed else ())
                return
            bad, good = ad.compare(value, run_id, path)
            # H1: a Campaign Info beamline (or facility acronym) naming another one than the folder's name or the
            # acquisition system's — a hint, not a reading, but never adopted over it: the scientist rules
            hinted = facility_hint_lines(bundle, path, value) if run_id is None else []
            if bad or hinted:
                # one row per cell of theirs and the VALUES the files read (a scan file and its calibration
                # file under one folder name are one disagreement); every source is kept, the agreeing ones too
                kind = "file_reading" if bad else "hint"
                ad.expose(good)
                key = (kind, sheet, col.header, their_ref, _compared_values(ad, bad, path))
                disagree(key, kind=kind, sheet=sheet, header=col.header, path=path, their_value=value,
                         reason="a file reading says otherwise" if bad else "a hint names another one",
                         scope="experiment" if run_id is None else "run",
                         _readings=[f"differs: {_shown_reading(r)}" for r in bad] + hinted
                         + [f"agrees: {_shown_reading(r)}" for r in good],
                         _their_cells=[their_ref], _runs=runs, _answer_cells=[target_cell])
                return
            ws = wb[sheet]
            ws[our_ref] = value
            ws[our_ref].comment = Comment(f"From your sheet: {their_file} {their_ref}", "ISAAC")
            adopted_cells.append((target_cell, their_ref))
            adopted_paths.add((run_id, path))
            result["cells_adopted"] += 1

        # Campaign Info: by label, whatever row it is on
        wc = wb_theirs[SHEET_CAMPAIGN]
        campaign_by_label = {}
        for col in CAMPAIGN_FIELDS:
            for label in (col.header, *col.aliases):
                campaign_by_label[_norm_header(label)] = col
        for r in range(1, wc.max_row + 1):
            label, raw = wc.cell(row=r, column=1).value, wc.cell(row=r, column=2).value
            col = campaign_by_label.get(_norm_header(label)) if label else None
            if raw is None or (isinstance(raw, str) and not raw.strip()):
                continue
            their_ref = _cell_ref(SHEET_CAMPAIGN, f"B{r}")
            if col is None or col.path is None:
                result["not_adopted"].append({"cell": their_ref, "header": str(label), "value": raw,
                                              "reason": (f"no official field ({col.why_none})" if col else
                                                         "no row of this name in ISAAC's template")})
                continue
            if index.status(col)[0] != "mapped":
                result["not_adopted"].append({"cell": their_ref, "header": col.header, "value": raw,
                                              "reason": index.status(col)[1]})
                continue
            our_ref = f"B{CAMPAIGN_FIRST_ROW + list(CAMPAIGN_FIELDS).index(col)}"
            judge(SHEET_CAMPAIGN, our_ref, col, raw, their_ref, None)

        # Configurations: Config ID and Technique
        wcf = wb_theirs[SHEET_CONFIG]
        cfg_header = {_norm_header(c.value): j for j, c in enumerate(wcf[1], start=1) if c.value}
        configs: dict[str, dict] = {}
        if "config id" in cfg_header:
            for r in range(2, wcf.max_row + 1):
                cid = wcf.cell(row=r, column=cfg_header["config id"]).value
                if cid is None or not str(cid).strip():
                    continue
                tech = wcf.cell(row=r, column=cfg_header["technique"]).value if "technique" in cfg_header else None
                configs[str(cid).strip()] = {"technique": tech,
                                             "ref": _cell_ref(SHEET_CONFIG, f"{_letter(cfg_header.get('technique', 2))}{r}")}
        tech_col = next(c for c in CONFIG_COLUMNS if c.header == "Technique")
        wcfg = wb[SHEET_CONFIG]
        for i, (cid, cfg) in enumerate(configs.items(), start=2):
            wcfg.cell(row=i, column=1, value=cid)
            cfg["our_row"] = i

        # File List: their columns -> ours
        mapped: dict[int, Column] = {}
        seen_cols: dict[str, int] = {}
        for j, h in enumerate(theirs["header"], start=1):
            if j == 1 or h is None or not str(h).strip():
                continue
            nh = _norm_header(h)
            target = next((c for c in FILE_COLUMNS if _norm_header(c.header) == nh), None)
            via_alias = False
            if target is None and nh in THEIR_FILE_HEADER_ALIASES:
                target = next(c for c in FILE_COLUMNS if c.header == THEIR_FILE_HEADER_ALIASES[nh])
                via_alias = True
            if target is None or target.header == "File Name":
                result["their_columns_not_read"].append({"column": _letter(j), "header": str(h).strip(),
                                                         "reason": "no column of this name in ISAAC's template"})
                continue
            if target.header in seen_cols:
                result["their_columns_not_read"].append({"column": _letter(j), "header": str(h).strip(),
                                                         "reason": f"column {_letter(seen_cols[target.header])} already maps to {target.header!r}"})
                continue
            if target.path is None and target.header != "Config ID":
                result["their_columns_not_read"].append({"column": _letter(j), "header": str(h).strip(),
                                                         "reason": f"no official field ({target.why_none})"})
                continue
            seen_cols[target.header] = j
            mapped[j] = target
            result["their_columns_mapped"].append({"column": _letter(j), "header": str(h).strip(), "to": target.header,
                                                   **({"by_alias": True} if via_alias else {})})

        # rows -> runs
        live = set(our_row)
        rows_of_run: dict[str, list[int]] = {}
        runs_of_row: dict[int, list[str]] = {}
        for r, name in theirs["names"]:
            ids = _match_runs(bundle, name)
            runs_of_row[r] = ids
            for rid in ids:
                rows_of_run.setdefault(rid, []).append(r)
        # a row that reaches several runs (a scan beside its ``_dir`` calibration file, or a folder of runs) is
        # said with its count: one value of theirs goes to each of them
        result["rows_reaching_several_runs"] = [{"row": _cell_ref(SHEET_FILES, f"A{r}"), "name": name, "runs": len(runs_of_row[r]),
                                                 "run_ids": runs_of_row[r]} for r, name in theirs["names"] if len(runs_of_row[r]) > 1]
        several = {rid for rid, rs in rows_of_run.items() if len(rs) > 1}
        result["runs_on_several_rows"] = [{"run_id": rid, "rows": [_cell_ref(SHEET_FILES, f"A{r}") for r in rows_of_run[rid]],
                                           "reason": "more than one of your File List rows names this run, so neither is adopted for it"}
                                          for rid in sorted(several)]
        for r, name in theirs["names"]:
            ids = runs_of_row[r]
            here = [rid for rid in ids if rid in live and rid not in several]
            outside = [rid for rid in ids if rid not in live]
            if outside:
                result["runs_outside_live_groups"].append({"row": _cell_ref(SHEET_FILES, f"A{r}"), "name": name, "run_ids": outside})
            if not ids:
                result["unmatched_rows"].append({"row": _cell_ref(SHEET_FILES, f"A{r}"), "name": name})
                continue
            if here:
                result["rows_matched"] += 1
            for rid in here:
                orow = our_row[rid]
                for j, col in mapped.items():
                    raw = wf_theirs.cell(row=r, column=j).value
                    if raw is None or (isinstance(raw, str) and not raw.strip()):
                        continue
                    their_ref = _cell_ref(SHEET_FILES, f"{_letter(j)}{r}")
                    if col.header == "Config ID":
                        cid = str(raw).strip()
                        cfg = configs.get(cid)
                        our_ref = f"{_letter(file_col['Config ID'])}{orow}"
                        if cfg is None:
                            disagree(("invalid", SHEET_FILES, "Config ID", repr(raw)), kind="invalid", sheet=SHEET_FILES,
                                     header="Technique", path=tech_col.path, their_value=raw, allowed=None, scope="run",
                                     reason=f"your Configurations sheet has no Config ID {cid!r}",
                                     _their_cells=[their_ref], _runs=[rid], _answer_cells=[_cell_ref(SHEET_FILES, our_ref)])
                            continue
                        tech = cfg["technique"]
                        if tech is not None and str(tech).strip():
                            value, why = _parse(tech, rules.get(tech_col.path))
                            bad, good = ([], []) if why else ad.compare(value, rid, tech_col.path)
                            if why or bad:
                                ad.expose(good)
                                kind = "file_reading" if bad else ("not_allowed" if (rules.get(tech_col.path) or {}).get("allowed") else "invalid")
                                key = (kind, SHEET_CONFIG, "Technique", cfg["ref"], tuple(sorted({repr(b["value"]) for b in bad})))
                                disagree(key, kind=kind, sheet=SHEET_CONFIG, header="Technique", path=tech_col.path, their_value=tech,
                                         reason=("a file reading says otherwise" if bad else "not an allowed value" if kind == "not_allowed"
                                                 else why.partition("; allowed: ")[0]),
                                         _readings=[f"differs: {_shown_reading(b)}" for b in bad]
                                         + [f"agrees: {_shown_reading(b)}" for b in good], scope="run",
                                         allowed=(rules.get(tech_col.path) or {}).get("allowed") if kind == "not_allowed" else None,
                                         _their_cells=[cfg["ref"], their_ref], _runs=[rid], _answer_cells=[_cell_ref(SHEET_FILES, our_ref)])
                                continue
                            ours = _cell_ref(SHEET_CONFIG, f"B{cfg['our_row']}")
                            if wcfg[f"B{cfg['our_row']}"].value is None:
                                wcfg[f"B{cfg['our_row']}"] = value
                                wcfg[f"B{cfg['our_row']}"].comment = Comment(f"From your sheet: {their_file} {cfg['ref']}", "ISAAC")
                                adopted_cells.append((ours, cfg["ref"]))
                        wb[SHEET_FILES][our_ref] = cid
                        wb[SHEET_FILES][our_ref].comment = Comment(f"From your sheet: {their_file} {their_ref}", "ISAAC")
                        adopted_cells.append((_cell_ref(SHEET_FILES, our_ref), their_ref))
                        result["cells_adopted"] += 1
                        continue
                    if index.status(col)[0] != "mapped":
                        result["not_adopted"].append({"cell": their_ref, "header": col.header, "value": raw, "run_id": rid,
                                                      "reason": index.status(col)[1]})
                        continue
                    judge(SHEET_FILES, f"{_letter(file_col[col.header])}{orow}", col, raw, their_ref, rid, their_row=r)
        matched_runs = {rid for rids in runs_of_row.values() for rid in rids if rid in live and rid not in several}
        result["runs_matched"] = len(matched_runs)
        result["unmatched_runs"] = [{"run_id": r["run_id"], "file_name": r["file_name"]} for r in rows if r["run_id"] not in matched_runs]

        # Disagreements: one row per disagreement; the answer cell stays EMPTY and says why
        order = {"file_reading": 0, "hint": 1, "name_only": 2, "not_allowed": 3, "invalid": 4}
        entries = sorted(groups.values(), key=lambda g: order[g["kind"]])
        wd = wb.create_sheet(SHEET_DISAGREE, index=wb.sheetnames.index(SHEET_SUGGEST))
        from openpyxl.styles import Alignment, Font

        for j, h in enumerate(DISAGREE_HEADERS, start=1):
            wd.cell(row=1, column=j, value=h).font = Font(bold=True)
            wd.column_dimensions[_letter(j)].width = 28 if h != "File reading(s)" else 60
        wd[f"{DISAGREE_RULING_COLUMN}1"].comment = Comment(
            "Type the value to record for this row, exactly as the field takes it (a number for a number field; one of "
            "the allowed values for a list). Leave it blank and nothing is recorded for it. ISAAC never picks a side.", "ISAAC")
        wd[f"{DISAGREE_CONC_COLUMN}1"].comment = Comment(
            "Only for an Electrolyte row whose value also states a concentration: the concentration in mol/L, a number "
            "(e.g. 0.01 for 0.01 M). It is recorded as Electrolyte Concentration (M) for exactly the runs of that row; the "
            "name goes in 'Your ruling'. Leave it blank on every other row.", "ISAAC")
        what = {"file_reading": "your value contradicts a file reading",
                "hint": "your value names another one than a hint (the folder's or the acquisition system's name; a hint, "
                        "not a file reading)",
                "name_only": "", "not_allowed": "", "invalid": "cannot be adopted as written: "}
        allowed_once: dict[str, list] = {}
        for i, g in enumerate(entries, start=1):
            g["id"] = f"D{i}"
            row = 1 + i
            g["row"] = row
            g["unread_hints"] = [{"column": c, "header": h, "value": v, "allowed_value": a, "rows": n}
                                 for (c, h, v, a), n in g["unread_hints"].items()]
            said = what[g["kind"]] if g["kind"] in ("file_reading", "hint") else what[g["kind"]] + g["reason"] + (
                " (each field's allowed values are listed once, below)" if g["kind"] == "not_allowed" else "")
            if g["unread_hints"]:
                said += ". " + "; ".join(_unread_hint_text(h, len(g["runs"])) for h in g["unread_hints"])
            values = (g["id"], said, g["header"], g["their_value"], _refs_shown(g["their_cells"]), "\n".join(g["readings"]) or None,
                      _refs_shown(g["runs"], noun="runs") if g["runs"] else "every File List run", None,
                      _refs_shown(g["answer_cells"]), None)
            for j, v in enumerate(values, start=1):
                c = wd.cell(row=row, column=j, value=v if not isinstance(v, str) or not v.startswith("=") else "'" + v)
                c.alignment = Alignment(wrap_text=True, vertical="top")
            wd[f"{DISAGREE_RULING_COLUMN}{row}"].number_format = "@" if not (rules.get(g["path"] or "") or {}).get("numeric") else "General"
            if g.get("allowed"):
                allowed_once.setdefault(g["header"], list(g["allowed"]))
            for ref in g["answer_cells"]:
                sheet, _, cell = ref.partition("!")
                ws = wb[sheet.strip("'")]
                ws[cell].fill = flag
                ws[cell].comment = Comment(
                    f"Not adopted ({g['id']} on the {SHEET_DISAGREE} sheet): your {their_file} {_refs_shown(g['their_cells'], 2)} "
                    f"has {_shown(g['their_value'])}; " + ("; ".join(g["readings"]) if g["readings"] else g["reason"])
                    + ". Rule on it there; this cell stays empty.", "ISAAC")
        if allowed_once:
            r = len(entries) + 3
            wd.cell(row=r, column=1, value="Allowed values (each field's list, once)").font = Font(bold=True)
            for k, (header, allowed) in enumerate(allowed_once.items(), start=1):
                wd.cell(row=r + k, column=1, value=header)
                wd.cell(row=r + k, column=2, value=", ".join(map(str, allowed)))
        if not entries:
            wd.cell(row=2, column=1, value="(no disagreements: every value of your sheet that ISAAC read agrees with the files or the files state nothing)")
        wd.freeze_panes = "A2"

        # _isaac: where each adopted cell came from, and what each Disagreements row's ruling reaches
        r = marker.max_row + 2
        marker.cell(row=r, column=1, value="[adopted]")
        r += 1
        for j, h in enumerate(("cell", "their_file", "their_cell"), start=1):
            marker.cell(row=r, column=j, value=h)
        for ours, their_ref in adopted_cells:
            r += 1
            for j, v in enumerate((ours, their_file, their_ref), start=1):
                marker.cell(row=r, column=j, value=v)
        r += 2
        marker.cell(row=r, column=1, value="[disagreements]")
        r += 1
        heads = ("id", "row", "kind", "header", "target_field_path", "scope", "targets", "their_file", "their_cells")
        for j, h in enumerate(heads, start=1):
            marker.cell(row=r, column=j, value=h)
        for g in entries:
            r += 1
            targets = ";".join(g["runs"]) if g["scope"] == "run" else "experiment"
            for j, v in enumerate((g["id"], g["row"], g["kind"], g["header"], g["path"], g["scope"], targets, their_file,
                                   ", ".join(g["their_cells"])), start=1):
                marker.cell(row=r, column=j, value=v)
        r += 2
        marker.cell(row=r, column=1, value="[meta]")  # read into the same meta as the first block
        marker.cell(row=r + 1, column=1, value="adopted_from")
        marker.cell(row=r + 1, column=2, value=their_file)

        result["disagreements"] = [{k: g[k] for k in ("id", "kind", "header", "path", "their_value", "their_cells", "readings",
                                                      "runs", "answer_cells", "reason", "scope", "unread_hints")}
                                   | {"ruling_cell": _cell_ref(SHEET_DISAGREE, f"{DISAGREE_RULING_COLUMN}{g['row']}")}
                                   | ({"concentration_ruling_cell": _cell_ref(SHEET_DISAGREE, f"{DISAGREE_CONC_COLUMN}{g['row']}")}
                                      if g["kind"] == "name_only" else {})
                                   for g in entries]
        result["allowed_values"] = allowed_once
        # L3: a reaction (adopted, or to be ruled on its row) makes drive and catalysis required for its runs
        reaction_runs = [rid for rid, p in adopted_paths if p in REACTION_PATHS and rid] + [
            rid for g in entries if g["path"] in REACTION_PATHS for rid in g["runs"]]
        result["reaction_needs"] = reaction_lacking(bundle, reaction_runs, adopted_paths) if reaction_runs else None
        result["ai_readings_of_their_sheet_not_compared"] = ad.sheet_readings_skipped
        result["_exposed"] = ad.exposed
        # a not-allowed value that is ONLY off its field's list: a word the vocabulary lacks (1.10.0)
        result["_off_list"] = [{"path": g["path"], "value": g["their_value"].strip(), "header": g["header"],
                                "cells": list(g["their_cells"]), "run_ids": list(g["runs"])}
                               for g in entries if g["kind"] == "not_allowed" and isinstance(g["their_value"], str)
                               and B.off_list_only(g["path"], g["their_value"].strip())]
        return result

    exported = export_template(bundle_dir, out, suggestions=suggestions, force=force, _fill=fill,
                               _command="bundle template adopt",
                               _keep_off_list={"by": by, "orcid": orcid, "file": their_file})
    exported["adopted"] = _as_text_dates(exported["adopted"])  # their cells as read: a date cell is a datetime
    exported["statement"] = (f"Answer cells hold your sheet's value only where no file reading contradicts it; each cites "
                             f"its cell in {their_file}. Nothing that contradicts a file is in an answer cell: rule on each "
                             f"row of the {SHEET_DISAGREE} sheet (Your ruling), then `template import`.")
    return exported


def render_adopt_summary(result: Mapping) -> str:
    """Compact: rows matched, cells adopted, EVERY disagreement named (run, field, their value, the
    file's reading), not-allowed values with each field's allowed list once, unmatched rows."""
    a = result["adopted"]
    lines = [f"Wrote {result['written']} from {a['their_file']} (bundle {result['bundle_id']}).",
             f"Rows: {a['rows_matched']} of your File List rows matched {a['runs_matched']} run(s); "
             f"{len(a['unmatched_rows'])} row(s) matched no run; {len(a['unmatched_runs'])} run(s) are on none of your rows.",
             f"Cells adopted: {a['cells_adopted']} (each cites your cell)."]
    mapped = [f"{m['header']} -> {m['to']}" for m in a["their_columns_mapped"] if m.get("by_alias")]
    if mapped:
        lines.append("Columns read under ISAAC's name: " + "; ".join(mapped) + ".")
    if a["their_columns_not_read"]:
        lines.append(f"Your columns not read ({len(a['their_columns_not_read'])}): "
                     + ", ".join(f"{c['column']} {c['header']!r}" for c in a["their_columns_not_read"]) + ".")
    file_ds = [d for d in a["disagreements"] if d["kind"] == "file_reading"]
    hint_ds = [d for d in a["disagreements"] if d["kind"] == "hint"]
    other = [d for d in a["disagreements"] if d["kind"] not in ("file_reading", "hint")]
    lines.append("")
    lines.append(f"Disagreements with a file reading ({len(file_ds)}; not adopted, each awaits your ruling; every reading "
                 "is listed, the agreeing ones too):")
    for d in file_ds:
        runs = _refs_shown(d["runs"], noun="runs") if d["runs"] else "every run"
        lines.append(f"  {d['id']} {runs} {d['header']}: yours {_shown(d['their_value'])} ({_refs_shown(d['their_cells'], 2)}); "
                     f"files: {'; '.join(d['readings'])}")
    if not file_ds:
        lines.append("  (none)")
    if hint_ds:
        lines.append("")
        lines.append(f"Disagreements with a hint ({len(hint_ds)}; not adopted, each awaits your ruling; a hint is not a file "
                     "reading, so say both and let the scientist rule):")
        for d in hint_ds:
            lines.append(f"  {d['id']} {d['header']}: yours {_shown(d['their_value'])} ({_refs_shown(d['their_cells'], 2)}); "
                         f"{'; '.join(d['readings'])}")
    if other:
        lines.append("")
        lines.append(f"Not adopted as written ({len(other)}; each awaits your ruling):")
        for d in other:
            reach = f"for {len(d['runs'])} run(s)" if d["runs"] else "for every File List run"
            lines.append(f"  {d['id']} {d['header']} {_shown(d['their_value'])} {reach}, from {len(d['their_cells'])} of your "
                         f"cell(s) ({_refs_shown(d['their_cells'], 2)}): {d['reason']}")
            for h in d.get("unread_hints") or ():
                lines.append(f"     hint: {_unread_hint_text(h, len(d['runs']))}")
        for header, allowed in a.get("allowed_values", {}).items():
            lines.append(f"  allowed for {header}: {', '.join(map(str, allowed))}")
    if a.get("vocabulary_requests"):
        lines.append("")
        lines.append("Kept as vocabulary requests (not decisions, not values: each is a word its field's list lacks; "
                     "`isaac bundle propose-terms` can draft it):")
        lines += _request_lines(a["vocabulary_requests"], True)
    for title, key, shown in (("Your rows that match no run", "unmatched_rows", lambda x: f"{x['row']} {x['name']}"),
                              ("Your rows that reach more than one run (each value goes to every one)", "rows_reaching_several_runs",
                               lambda x: f"{x['row']} {x['name']} -> {x['runs']} runs"),
                              ("Runs on none of your rows", "unmatched_runs", lambda x: x["file_name"]),
                              ("Runs on several of your rows (not adopted)", "runs_on_several_rows", lambda x: f"{x['run_id']} ({', '.join(x['rows'])})")):
        items = a[key]
        if items:
            more = f" and {len(items) - 10} more (--json lists all)" if len(items) > 10 else ""
            lines.append("")
            lines.append(f"{title} ({len(items)}): " + "; ".join(shown(x) for x in items[:10]) + more)
    if a["not_adopted"]:
        by_header: dict[str, list] = {}
        for x in a["not_adopted"]:
            by_header.setdefault(x["header"], []).append(x)
        lines.append(f"Cells with nowhere to go ({len(a['not_adopted'])}; not copied): " + "; ".join(
            f"{h} {len(xs)} cell(s) ({xs[0]['reason']})" for h, xs in by_header.items()) + ".")
    if a.get("reaction_needs"):
        lines.append(_reaction_needs_line(a["reaction_needs"], "(adopted, or once ruled on its Disagreements row)"))
    if a.get("ai_readings_of_their_sheet_not_compared"):
        lines.append(f"{a['ai_readings_of_their_sheet_not_compared']} AI reading(s) of your sheet itself were not compared: they are your sheet, not the files.")
    lines.append("")
    lines.append(result["statement"])
    lines.append(f"Then: isaac bundle template import {result.get('bundle_dir') or '<bundle_dir>'} {result['written']} --by \"<your name>\" "
                 "(a preview; add --apply to record).")
    return "\n".join(lines)
